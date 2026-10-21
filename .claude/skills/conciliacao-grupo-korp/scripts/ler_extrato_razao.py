import openpyxl,datetime,collections,itertools,pickle,re
def s(v): return v.strip() if isinstance(v,str) else v
wa=openpyxl.load_workbook('a.xlsx',data_only=True).active
ext=[]
for r in wa.iter_rows(values_only=True):
    if isinstance(r[0],datetime.datetime) and r[20] is not None:
        v=r[20]; ext.append(dict(d=r[0].date(),desc=s(r[6]),inf=s(r[11]),ger=s(r[15]),sv=round(v if s(r[19])=='+' else -v,2),saldo=r[22]))
wb=openpyxl.load_workbook('b.xlsx',data_only=True).active
raz=[];tot={};mes=None
for r in wb.iter_rows(values_only=True):
    if isinstance(r[0],str) and r[0].startswith('Mês:'): mes=r[0].replace('Mês:','').strip()
    elif isinstance(r[0],datetime.datetime) and r[1]:
        raz.append(dict(d=r[0].date(),n=s(r[1]),obs=s(r[3]).replace('VLR. DOCUMENTO Nro.: ',''),deb=r[6] or 0,cred=r[7] or 0,cc=s(r[10]),sa=r[12],sv=round((r[6] or 0)-(r[7] or 0),2)))
    elif any(isinstance(x,str) and x.strip().startswith('Total :') for x in r):
        tot[mes]=[x for x in r if isinstance(x,(int,float))]
print(len(ext),len(raz)); print(tot)
print('ext range',ext[0]['d'],ext[-1]['d'],'raz range',raz[0]['d'],raz[-1]['d'])
pickle.dump((ext,raz,tot),open('data.pkl','wb'))
