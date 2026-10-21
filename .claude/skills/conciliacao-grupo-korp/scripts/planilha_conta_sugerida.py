import re,itertools,collections,datetime,xlrd,openpyxl,sys
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
src=open('p.py').read().split("print(len(ext)")[0]
src=src.replace("print('A other:'","pass #").replace("print('B other:'","pass #")
exec(src)
for e in ext: e['sv']=round(e['v'] if e['dc']=='+' else -e['v'],2)
for r in raz: r['sv']=round(r['deb']-r['cred'],2)
# ---- plano
ws=xlrd.open_workbook('plano.xls').sheet_by_index(0)
rows=[]
for i in range(1,ws.nrows):
    r=ws.row_values(i)
    red=[str(x) for x in r[1:10] if x not in ('','-')]
    rows.append((red[0] if red else '',r[10],r[13].strip(),r[12]))
codes=[c for _,c,_,_ in rows]
plano={red:(c,n,dc) for red,c,n,dc in rows if not any(o.startswith(c+'.') for o in codes)}
# ---- casar razão x extrato
pool=collections.defaultdict(list)
for r in raz: pool[(r['d'],r['sv'])].append(r)
for r in raz: r['ext']=None
unmatched=[]
for e in ext:
    k=(e['d'],e['sv'])
    if pool[k]:
        r=pool[k].pop(0); r['ext']=e; e['lin']=[r]
    else: unmatched.append(e)
resto=[r for v in pool.values() for r in v]
falt=[]
byd=collections.defaultdict(list)
for e in unmatched: byd[e['d']].append(e)
for d,es in byd.items():
    cand=[r for r in resto if r['d']==d]
    if len(es)==1 and cand and round(sum(x['sv'] for x in cand),2)==es[0]['sv']:
        es[0]['lin']=cand
        for r in cand: r['ext']=es[0]; resto.remove(r)
        continue
    for e in es:
        cand=[r for r in resto if r['d']==d]
        ach=None
        if len(cand)<=18:
            for n in range(1,len(cand)+1):
                for c in itertools.combinations(cand,n):
                    if round(sum(x['sv'] for x in c),2)==e['sv']: ach=c;break
                if ach:break
        if ach:
            e['lin']=list(ach)
            for r in ach: r['ext']=e; resto.remove(r)
        else: falt.append(e)
assert not resto
# ---- regras de conta
R=[
(r'EMPRESTIMO PARA KORP','3086','Média','Empréstimo entre Matriz e Filial (mútuo intercompany). Confirmar conta com a contrapartida na outra empresa.'),
(r'SEGURO DE VIDA','3047','Média','Seguro de vida em grupo dos colaboradores (Benefícios). Alternativa: despesa 3.1.3.2.1.17.8.7 (3006).'),
(r'SAL[ÁA]RIO','2521','Alta','Pagamento de salários (2521 Salários a Pagar).'),
(r'SETRAB|SEGURANCA E M','3107','Média','Assessoria em segurança/medicina do trabalho. Alternativa: 2637.'),
(r'LIGGA|INTERNET','3059','Média','Internet/link dedicado.'),
(r'NF VIVO TELEFONE|TELEFONE FIXO','2657','Alta','Telefone/comunicação.'),
(r'TARIFA|CESTA DE RELACIONAMENTO','2698','Alta','Tarifa bancária.'),
(r'TED TRANSFERENCIA','2872','Média','Transferência entre Matriz e Filial. Confirmar se o escritório usa 2872 (Transferência de recursos) ou mútuo (3086/2861).'),
(r'DCK SERVICOS.*EMPREST|EMPRESTIMO PARA.*DCK|DCK SERVICOS DE T\.I.*RECEBIMENTO','3085','Média','Empréstimo/mútuo recebido da DCK. Alternativa: 2861 (DCK/KORP Intercompany).'),
(r'KORP INFORMATICA.*EMPREST','3086','Média','Empréstimo entre Korp Filial e Filial (mútuo). Confirmar conta (3086 ativo / 2861 passivo).'),
(r'LOTE CNAB|MENSALIDADE','2436','Alta','Baixa de duplicata/mensalidade de cliente.'),
(r'IOF CART','2698','Média','IOF do cartão Sicredi (não há conta específica de IOF sobre compras).'),
(r'ICMS COMP','2511','Alta','Recolhimento de ICMS (competência 08/26).'),
(r'PIS \d+ COMP|PIS A RECOLHER','2512','Alta','Recolhimento de PIS (competência 08/26).'),
(r'COFINS \d+ COMP','2513','Alta','Recolhimento de COFINS (competência 08/26).'),
(r'ISS COMP','2516','Alta','Recolhimento de ISS (competência 08/26).'),
(r'MINISTERIO DA FAZENDA.*\d+/U|\d+/U .*MINISTERIO DA FAZENDA','2822','Média','DARF ligado à NF do fornecedor (nº/U): retenção PIS/COFINS/CSLL. Confirmar o tributo.'),
(r'DARF UNIFICADO','2531','Baixa','DARF unificado: separar INSS (2531) e IRRF (2533) conforme a guia; o extrato traz o valor total.'),
(r'FGTS','2530','Alta','Recolhimento de FGTS.'),
(r'RESCIS','2523','Alta','Rescisão (padrão da folha Korp: 2523 Indenizações Trabalhistas).'),
(r'F[ÉE]RIAS','2524','Alta','Pagamento de férias (2524 Férias a Pagar).'),
(r'SINDICA','2529','Alta','Contribuição sindical a recolher.'),
(r'\bVT\b|VALE TRANSPORTE|METROCARD','3050','Alta','Vale-transporte de funcionários.'),
(r'TICKET SERVICOS','3046','Média','Vale-refeição/alimentação (Benefícios - Refeições). Alternativa: 3043 PAT.'),
(r'UNIMED','3044','Alta','Plano de saúde dos colaboradores.'),
(r'SESI','3107','Média','Exames ocupacionais. Se for admissional/demissional/periódico, usar 3038/3039/3069.'),
(r'MAPA HDS|GUPY','3113','Alta','Recrutamento e seleção.'),
(r'VR BENEFICIOS','3112','Média','Taxa/benefício VR. Alternativa: 3046.'),
(r'PAYTRACK','2687','Baixa','Cartão Paytrack (despesas de viagem/combustível). Conferir prestação de contas; recarga pode ser adiantamento (2464).'),
(r'HOTEL|APART','2587','Média','Hospedagem (despesa comercial). Se for administrativa: 2647 Viagens.'),
(r'UBER','2680','Alta','Transporte por aplicativo.'),
(r'MAXMILHAS|MM TURISMO|PASSAGENS|CLICKBUS','2647','Média','Passagens/viagens. Se for comercial: 3034 ou 2587.'),
(r'NOMACHINE|CURSOR|OPENAI|OPENROUTER|FIGMA|ANTHROPIC|OBSIDIAN|MANUS AI|SOFTWARE P/','2859','Alta','Assinatura/licença de software.'),
(r'ORACLE|INGRAM|ACBR|DG TECH|IANDEV|ADVANT|LEADS2B|ROBSON DEL|MANUTEN[ÇC][ÃA]O DE SOFTWARE','2859','Média','Licenciamento/manutenção de software de terceiros. Alternativa (custo de serviço): 2839.'),
(r'CSP DESENVOLVEDOR','3139','Alta','Custo de serviço terceiro - Desenvolvedor.'),
(r'CSP QUALIDADE','3138','Alta','Custo de serviço terceiro - Qualidade.'),
(r'CSP CONSULTOR','3065','Alta','Custo de serviço terceiro - Consultores.'),
(r'CSP PRESTADOR','3067','Média','Custo de serviço terceiro (PJ).'),
(r'CONSULTORIA COMERCIAL','3141','Alta','Custo de serviço terceiro - PJ Comercial.'),
(r'TRANSF\. DIGITAL','3067','Baixa','Categoria RSV Transf. Digital: confirmar natureza (custo de serviço PJ ou repasse).'),
(r'INSTITUTO SYNAPSE','3067','Baixa','Mesmo fornecedor/categoria RSV Transf. Digital: confirmar natureza.'),
(r'CONDOM','2869','Alta','Condomínio.'),
(r'ENERGIA|COPEL|REDUZZA','2688','Alta','Energia elétrica.'),
(r'TELEFONE M[ÓO]VEL|TELEFONE VIVO|MONITORA DADOS','2657','Alta','Telefone/comunicação.'),
(r'VIVO IP DEDICADO|BANCA LARGA|INTERNET','3059','Média','Internet/link dedicado. Telefonia: 2657.'),
(r'TRANSPORTE, FRETE|MOTOBOY|GUSTAVO RIZZARDI','2658','Alta','Frete/motoboy.'),
(r'ADVOCAC|ADVOGADOS','2629','Alta','Honorários advocatícios.'),
(r'CONTABILIDADE','2628','Alta','Honorários contábeis.'),
(r'COPA/COZINHA|FABESUL','2838','Alta','Copa e cozinha.'),
(r'COPIADORA|TECNICOPIAS','2664','Média','Cópias/material de expediente.'),
(r'IX TECNOLOGIA|PUBLICIDADE E PROPAGANDA|EQTOON','3060','Média','Publicidade e propaganda. Administrativa: 2685.'),
(r'ARKE|JANAINA|ADRIANO FELIPPETO|BRINDES','3133','Média','Brindes/presentes para clientes. Se for propaganda: 3060.'),
(r'ASSB COMERCIO','2656','Baixa','Compra de doces: confraternização/presentes? Confirmar.'),
(r'LOCABOOK','2637','Baixa','Locação de equipamentos de informática: confirmar (2637 serviços de terceiros ou conta de locação).'),
(r'JMTEC|SAIPH|INFORM[ÁA]TICA','2661','Baixa','Informática (infra/peças). Confirmar.'),
(r'MERCADO LIVRE','2649','Baixa','Equipamento (cat. gerencial EQ). Bem de pequeno valor; se imobilizar: 2477.'),
(r'MERCADO PAGO.*DEVOLU','2649','Média','Devolução de compra (estorno): creditar a mesma conta da despesa original.'),
(r'R L M FERNANDES|DAEDALUZ|ERIC HENRIQUE|LEANDRO RODRIGUES|C S BARBOSA|TASS COMERCIAL|ADRIANA APARECIDA','2649','Baixa','Compra AC (sem categoria gerencial no extrato): bem de pequeno valor/material. Confirmar.'),
]
def sugere(r_obs,e):
    if e['sv']>0 and (e['ger'] or '').startswith('RSV'):
        return '2436','Alta','Recebimento de cliente (baixa de duplicata/mensalidade).'
    if r_obs:
        o=r_obs.upper()
        for pat,cod,conf,obs in R:
            if re.search(pat,o) and not re.search(r'^(TARIFA|CESTA)',o): return cod,conf,obs
    t=' '.join(x for x in [r_obs or '',e['desc'] or '',e['inf'] or '',e['ger'] or ''] if x).upper()
    for pat,cod,conf,obs in R:
        if re.search(pat,t): return cod,conf,obs
    return None,'Sem regra','Revisar manualmente'
def linha(r,e,cod,conf,obs):
    if cod:
        c,n,dc=plano[cod]
    else: c=n=dc=''
    return c,n,dc
AREA={'CC320':'ADM geral','CC337':'Desenvolvimento (DES)','CC345':'Suporte','CC379':'Comercial','CC506':'Qualidade','CC644':'Consultoria / CS'}
CSP={'3139','3140','3138','3065'}
def porcc(cod,cc,obs,conf):
    area=AREA.get(cc)
    if not cod: return cod,obs,conf,area or ('CC sem de-para' if cc else 'sem CC')
    if cc in ('CC337','CC345','CC506','CC644') and cod in CSP|{'3067'}:
        novo={'CC337':'3139','CC345':'3140','CC506':'3138','CC644':'3065'}[cc]
        if novo!=cod: obs=f'CC {cc} ({area}) manda para a conta de custo do grupo; categoria do extrato sugeria outra. '+obs; cod=novo
    if cc=='CC379':
        m={'2647':'3034','2657':'2592','3059':'2870','2664':'2584'}
        if cod in m: obs=f'CC379 Comercial: conta comercial. '+obs; cod=m[cod]
    if cc=='CC320':
        m={'2587':'2647','3060':'2685','2592':'2657','2870':'3059'}
        if cod in m: obs=f'CC320 ADM: conta administrativa. '+obs; cod=m[cod]
    if cc and cc not in AREA:
        obs=obs+' | CC sem de-para de área: validar a conta pelo centro de custo.'
        if conf=='Alta': conf='Média'
    return cod,obs,conf,area or ('CC sem de-para' if cc else 'sem CC')
# ---- planilha
wb=openpyxl.Workbook(); w=wb.active; w.title='Conta sugerida'
H=['N. Lanc (Razão)','Data','Valor razão (R$)','D/C no banco','Histórico (razão)','Histórico (extrato)','Informação (extrato)','Categoria gerencial (extrato)','Valor do mov. no extrato (R$)','Confere com extrato?','CC (razão)','Área do CC','Cód. reduzido sugerido','Classificação','Nome da conta sugerida','Confiança','Observação']
w.append(H)
out=[]
for r in sorted(raz,key=lambda x:(x['d'],int(x['n']))):
    e=r['ext']; cod,conf,obs=sugere(r['obs'],e); cod,obs,conf,area=porcc(cod,r['cc'],obs,conf)
    c,n,_=linha(r,e,cod,conf,obs)
    out.append(r)
    rateio=len(e['lin'])>1
    ok='Sim' if abs(sum(x['sv'] for x in e['lin'])-e['sv'])<0.005 else 'NÃO'
    if rateio: ok+=f' (rateio de {len(e["lin"])} lançamentos)'
    w.append([int(r['n']),r['d'],abs(r['sv']),'Entrada (D)' if r['sv']>0 else 'Saída (C)',r['obs'].replace('VLR. DOCUMENTO Nro.: ',''),e['desc'],e['inf'],e['ger'],abs(e['sv']),ok,r['cc'],area,int(cod) if cod else '',c,n,conf,obs])
# faltantes
w2=wb.create_sheet('Faltam no razão')
w2.append(['N. Lanc (Razão)','Data','Valor extrato (R$)','D/C no banco','Histórico (extrato)','Informação (extrato)','Categoria gerencial (extrato)','Cód. reduzido sugerido','Classificação','Nome da conta sugerida','Confiança','Observação'])
for e in sorted(falt,key=lambda x:x['d']):
    cod,conf,obs=sugere('',e); c,n,_=linha(None,e,cod,conf,obs)
    w2.append(['NÃO LANÇADO',e['d'],abs(e['sv']),'Entrada (D)' if e['sv']>0 else 'Saída (C)',e['desc'],e['inf'],e['ger'],int(cod) if cod else '',c,n,conf,obs])
# resumo
w3=wb.create_sheet('Resumo')
INI=round(ext[0]['saldo']-ext[0]['sv'],2)
sx=round(sum(e['sv'] for e in ext),2); sr=round(sum(r['sv'] for r in raz),2)
for row in [['Conferência extrato × razão - conta 2871 Banco Sicredi (Korp Matriz) - 09/2026'],[],
 ['Saldo inicial (igual nos dois)',INI],['Saldo final extrato',round(INI+sx,2)],['Saldo final razão',round(INI+sr,2)],['Diferença',round(sx-sr,2)],[],
 ['Movimentos do extrato',len(ext)],['Lançamentos do razão',len(raz)],['Movimentos do extrato SEM lançamento no razão',len(falt)],['Valor líquido que falta no razão',round(sum(e['sv'] for e in falt),2)],[],
 ['Legenda de confiança'],['Alta','Histórico/categoria identifica a conta sem dúvida'],['Média','Conta provável; ver observação (alternativas)'],['Baixa','Precisa de confirmação (natureza do gasto ou rateio da guia)'],[],
 ['Fonte do plano de contas','Plano de Contas - Atualizado.xls (Drive) - código reduzido = coluna "Código" da Korp (analíticas)'],
 ['Atenção','A coluna "Cód. reduzido" é o número que vai no TXT/lançamento da Korp. Sugestão automática por histórico: validar antes de lançar.']]:
    w3.append(row)
# estilo
hf=PatternFill('solid',fgColor='1F3864')
thin=Side(style='thin',color='BBBBBB')
for sh,wd in [(w,[11,11,14,12,52,44,26,34,16,22,9,18,12,16,42,10,70]),(w2,[13,11,14,12,44,24,34,12,16,42,10,70])]:
    for c in sh[1]:
        c.font=Font(bold=True,color='FFFFFF'); c.fill=hf; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    for i,x in enumerate(wd,1): sh.column_dimensions[get_column_letter(i)].width=x
    sh.freeze_panes='C2'; sh.auto_filter.ref=sh.dimensions; sh.row_dimensions[1].height=42
    for row in sh.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
        row[1].number_format='dd/mm/yyyy'; row[2].number_format='#,##0.00'
    colconf=[c for c in sh[1] if c.value=='Confiança'][0].column
    for row in sh.iter_rows(min_row=2):
        c=row[colconf-1]
        c.fill=PatternFill('solid',fgColor={'Alta':'C6EFCE','Média':'FFEB9C','Baixa':'F8CBAD'}.get(c.value,'D9D9D9'))
for row in w.iter_rows(min_row=2): row[8].number_format='#,##0.00'
for row in w2.iter_rows(min_row=2):
    for c in row[:4]: c.fill=PatternFill('solid',fgColor='FFF2CC')
w3.column_dimensions['A'].width=52; w3.column_dimensions['B'].width=90
w3['A1'].font=Font(bold=True,size=13)
for r_ in range(3,12):
    if isinstance(w3.cell(r_,2).value,float): w3.cell(r_,2).number_format='#,##0.00'
wb.move_sheet('Resumo',offset=-2)
wb.save('Conciliacao_Sicredi_Matriz_09-2026_conta_sugerida.xlsx')
sem=[(r['n'],r['obs']) for r in raz if sugere(r['obs'],r['ext'])[0] is None]
print('sem regra',sem); print('faltam',len(falt),'linhas',len(out))
for r in out[:0]: pass
print(collections.Counter(sugere(r['obs'],r['ext'])[1] for r in raz))
