"""Concilia extrato bancário Korp x razão da conta do banco (exportações .xlsx do Korp).

Uso: python3 -I conciliar.py EXTRATO.xlsx RAZAO.xlsx
Os .xlsx do Korp costumam ter cores inválidas no styles.xml (openpyxl falha);
este script corrige isso numa cópia antes de ler.
"""
import sys, re, zipfile, io, datetime, collections, itertools
import openpyxl

def abrir(path):
    zin = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == 'xl/styles.xml':
                data = re.sub(rb'rgb="[^"]*"', b'rgb="FF000000"', data)
            zout.writestr(it, data)
    buf.seek(0)
    return openpyxl.load_workbook(buf, data_only=True).active

def s(v): return v.strip() if isinstance(v, str) else v

def ler_extrato(ws):
    out = []
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], datetime.datetime):
            v = r[20]
            out.append(dict(d=r[0].date(), desc=s(r[6]), inf=s(r[11]),
                            sv=round(v if s(r[19]) == '+' else -v, 2)))
    return out

def ler_razao(ws):
    out = []
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], datetime.datetime):
            out.append(dict(d=r[0].date(), n=s(r[1]), obs=s(r[3]), cc=s(r[10]),
                            sv=round((r[6] or 0) - (r[7] or 0), 2)))
    return out

def main(fe, fr):
    ext, raz = ler_extrato(abrir(fe)), ler_razao(abrir(fr))
    print(f'Extrato: {len(ext)} mov. | Razão: {len(raz)} lanç.')
    print('Soma extrato:', round(sum(e['sv'] for e in ext), 2),
          '| Soma razão:', round(sum(r['sv'] for r in raz), 2))
    # 1) casa data + valor exato
    pool = collections.defaultdict(list)
    for r in raz: pool[(r['d'], r['sv'])].append(r)
    sem_raz = []
    for e in ext:
        k = (e['d'], e['sv'])
        if pool[k]: pool[k].pop()
        else: sem_raz.append(e)
    sobra = [r for v in pool.values() for r in v]
    # 2) rateios: um mov. do extrato = soma de várias linhas do razão (mesma data)
    faltam = []
    for e in sem_raz:
        cand = [r for r in sobra if r['d'] == e['d']]
        achou = None
        for n in range(2, min(len(cand), 16) + 1):
            for comb in itertools.combinations(cand, n):
                if round(sum(r['sv'] for r in comb), 2) == e['sv']:
                    achou = comb; break
            if achou: break
        if achou:
            print(f"RATEIO OK  {e['d']} {e['sv']:>12} {e['desc']} = {len(achou)} linhas")
            for r in achou: sobra.remove(r)
        else:
            faltam.append(e)
    print('\n== FALTAM NO RAZÃO (estão no extrato) ==')
    for e in faltam: print(e['d'], f"{e['sv']:>12}", e['desc'], e['inf'])
    print('\n== SOBRAM NO RAZÃO (não estão no extrato) ==')
    for r in sobra: print(r['d'], r['n'], f"{r['sv']:>12}", r['obs'], r['cc'])
    print('\nDiferença explicada (extrato sem razão):', round(sum(e['sv'] for e in faltam), 2),
          '| razão sem extrato:', round(sum(r['sv'] for r in sobra), 2))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
