"""Planilhas padrão da folha Korp.

python planilhas.py correcoes correcoes.json  saida.xlsx
python planilhas.py fechamento fechamento.json saida.xlsx

correcoes.json: {"titulo": "...", "linhas": [[valor, "CC Domínio", "Rubrica", "Aba Domínio",
                 "Débito atual", "Crédito atual", "Débito sugerido", "Crédito sugerido", "Sim/Não (mês anterior)", "Nota"], ...]}
fechamento.json: {"titulo": "...", "meses": {"05-2026": {"nota": "...", "2521": [folha, txt, correção, "composição", outros_meses], ...}}}
  - um arquivo por mês (padrão): passar só o mês em "meses"
  - outros_meses (opcional): valores de competência de outro mês que estão no TXT deste mês (+) ou que faltam nele (-)
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

F = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k})
thin = Side(style='thin', color='BFBFBF'); B = Border(top=thin, bottom=thin, left=thin, right=thin)
HEAD = PatternFill('solid', fgColor='D9E1F2'); EDIT = PatternFill('solid', fgColor='FFF2CC'); ASK = PatternFill('solid', fgColor='FCE4D6')
MOEDA = '#,##0.00;-#,##0.00;-'
CONTAS = [('2521', 'Salários a Pagar'), ('2523', 'Indenizações Trabalhistas (rescisões)'), ('2524', 'Férias a Pagar'),
          ('2526', 'Estágios a pagar'), ('2530', 'FGTS a Recolher'), ('2531', 'INSS a Recolher'),
          ('2533', 'IRRF a Recolher s/ Folha'), ('3137', 'Empréstimo a Funcionário (consignado)'),
          ('2463', 'Adiantamento de Férias (ativo)')]
MESES_ANO = ['01', '02', '03', '04', '05', '06', '07', '08', '09', '10', '11', '12']

def cab(ws, row, headers):
    for c, h in enumerate(headers, 1):
        x = ws.cell(row, c, h); x.font = F(bold=True); x.fill = HEAD; x.border = B
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def correcoes(cfg, dst):
    wb = Workbook(); ws = wb.active; ws.title = cfg.get('aba', 'Correções')
    ws['A1'] = cfg['titulo']; ws['A1'].font = F(bold=True, size=12)
    ws['A2'] = 'Preencha apenas as colunas amarelas ("Corrigido?" e "Observação"). Células laranja = aguardando confirmação.'
    ws['A2'].font = F(italic=True, color='666666')
    H = ['Valor', 'CC Domínio', 'Rubrica', 'Aba Domínio', 'Débito atual', 'Crédito atual', 'Débito sugerido',
         'Crédito sugerido', 'Já ocorreu no mês anterior?', 'Nota', 'Corrigido?', 'Observação']
    hr = 4; cab(ws, hr, H)
    for i, r in enumerate(cfg['linhas'], hr + 1):
        for c, v in enumerate(list(r) + ['', ''], 1):
            x = ws.cell(i, c, v); x.font = F(); x.border = B
            x.alignment = Alignment(horizontal='left' if c in (2, 3, 10, 12) else 'center', vertical='center', wrap_text=c == 10)
            if c in (11, 12): x.fill = EDIT
            if v in ('Confirmar', '?'): x.fill = ASK
        ws.cell(i, 1).number_format = '#,##0.00'; ws.cell(i, 1).alignment = Alignment(horizontal='right', vertical='center')
    last = hr + len(cfg['linhas']); t = last + 1
    ws.cell(t, 1, f'=SUM(A{hr+1}:A{last})').font = F(bold=True); ws.cell(t, 1).number_format = '#,##0.00'
    ws.cell(t, 2, 'Total').font = F(bold=True)
    ws.cell(t, 10, 'Corrigidos:').font = F(bold=True); ws.cell(t, 10).alignment = Alignment(horizontal='right')
    ws.cell(t, 11, f'=COUNTIF(K{hr+1}:K{last},"Sim")&" de "&ROWS(K{hr+1}:K{last})').font = F(bold=True)
    dv = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'K{hr+1}:K{last}')
    for c, w in zip('ABCDEFGHIJKL', [11, 28, 36, 11, 11, 11, 13, 13, 12, 46, 11, 34]): ws.column_dimensions[c].width = w
    ws.freeze_panes = ws.cell(hr + 1, 1)
    wb.calculation.fullCalcOnLoad = True; wb.save(dst)

def fechamento(cfg, dst):
    wb = Workbook(); rs = wb.active; rs.title = 'Resumo'
    meses = cfg['meses']; ano = cfg.get('ano', '2026')
    cols = [m for m in sorted(meses, key=lambda k: k[3:] + k[:2])]
    rs['A1'] = cfg['titulo']; rs['A1'].font = F(bold=True, size=12)
    rs['A2'] = 'Valores = movimento do mês (crédito - débito). Confere = folha do mês + competência de outros meses = TXT corrigido.'
    rs['A2'].font = F(italic=True, color='666666')
    cab(rs, 4, ['Conta', 'Descrição'] + cols)
    n = len(CONTAS)
    for i, (ct, nm) in enumerate(CONTAS, 5):
        rs.cell(i, 1, ct); rs.cell(i, 2, nm)
        for j, m in enumerate(cols, 3):
            if m in meses: rs.cell(i, j, f"='{m}'!F{i}")
            rs.cell(i, j).number_format = MOEDA
        for c in range(1, len(cols) + 3): rs.cell(i, c).font = F(); rs.cell(i, c).border = B
    r = 5 + n; rs.cell(r, 2, 'Confere com a folha?').font = F(bold=True)
    for j, m in enumerate(cols, 3):
        if m in meses: rs.cell(r, j, f"=IF(COUNTIF('{m}'!H5:H{4+n},\"Não\")=0,\"Sim\",\"Não\")").font = F(bold=True)
        rs.cell(r, j).alignment = Alignment(horizontal='center')
    rs.column_dimensions['A'].width = 8; rs.column_dimensions['B'].width = 38
    for j in range(3, 3 + len(cols)): rs.column_dimensions[rs.cell(4, j).column_letter].width = 13
    rs.freeze_panes = 'C5'
    H = ['Conta', 'Descrição', 'Valor pela folha (extrato do mês)', 'Valor no TXT do Domínio', 'Correções de conta',
         'TXT corrigido (lançado na Korp)', 'Competência de outros meses no TXT', 'Confere?', 'Composição']
    for m, d in meses.items():
        ws = wb.create_sheet(m)
        ws['A1'] = f"{cfg['titulo']} - {m.replace('-', '/')}"; ws['A1'].font = F(bold=True, size=12)
        ws['A2'] = d.get('nota', ''); ws['A2'].font = F(italic=True, color='666666')
        cab(ws, 4, H)
        for i, (ct, nm) in enumerate(CONTAS, 5):
            v = list(d.get(ct, [0, 0, 0, ''])) + [0]
            fo, tx, aj, ex, ou = v[:5]
            vals = [ct, nm, fo, tx, aj, f'=D{i}+E{i}', ou, f'=IF(ABS(C{i}+G{i}-F{i})<1,"Sim","Não")', ex]
            for c, v in enumerate(vals, 1):
                x = ws.cell(i, c, v); x.font = F(); x.border = B
                x.alignment = Alignment(horizontal='left' if c in (2, 9) else 'center', vertical='center', wrap_text=c == 9)
                if c in (3, 4, 5, 6, 7): x.number_format = MOEDA; x.alignment = Alignment(horizontal='right', vertical='center')
            ws.row_dimensions[i].height = 42 if ex else 16
        t = 6 + n
        ws.cell(t, 1, 'Salários + Indenizações + Férias + Estágios a pagar').font = F(bold=True)
        ws.cell(t + 1, 2, 'Total').font = F(bold=True)
        for c, col in ((3, 'C'), (4, 'D'), (6, 'F'), (7, 'G')):
            x = ws.cell(t + 1, c, f'={col}5+{col}6+{col}7+{col}8'); x.font = F(bold=True); x.number_format = '#,##0.00'
        for c, w in zip('ABCDEFGHI', [8, 34, 16, 15, 13, 16, 16, 10, 90]): ws.column_dimensions[c].width = w
        ws.freeze_panes = 'C5'
    wb.calculation.fullCalcOnLoad = True; wb.save(dst)

if __name__ == '__main__':
    tipo, cfg, dst = sys.argv[1:4]
    {'correcoes': correcoes, 'fechamento': fechamento}[tipo](json.load(open(cfg, encoding='utf-8')), dst)
    print('ok ->', dst)
