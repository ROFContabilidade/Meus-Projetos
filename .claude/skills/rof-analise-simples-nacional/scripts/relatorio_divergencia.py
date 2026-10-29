"""Relatório INTERNO (não enviar ao cliente): faturamento declarado × recebimentos bancários × compras.

Uso:
  python relatorio_divergencia.py --apuracao apuracao.json --recebimentos recebimentos.json --compras compras.json \\
      --empresa "NOME (empresa NN)" --saida RELATORIO.xlsx [--resumo divergencia.json]

Fontes: apuracao.json (ler_apuracao.py), recebimentos.json (ler_razao.py), compras.json (ler_compras.py)."""
import argparse, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference

ap = argparse.ArgumentParser()
for k in ['apuracao', 'recebimentos', 'compras', 'empresa', 'saida']: ap.add_argument('--' + k, required=True)
ap.add_argument('--resumo')
a = ap.parse_args()
apur = json.load(open(a.apuracao)); rz = json.load(open(a.recebimentos)); comp = json.load(open(a.compras))
f_ano = {m['competencia']: m['rpa'] for m in apur}
hist = {}
for m in apur:
    for mes, v in m.get('historico_12m', []): hist[mes] = v
hist.update(f_ano)
ordem = lambda mm: mm[3:] + mm[:2]
serie = sorted(hist.items(), key=lambda x: ordem(x[0]))
ano = apur[-1]['competencia'][3:]; ant = str(int(ano) - 1)
f_ant = [v for mm, v in serie if mm.endswith(ant)]
contas = list(rz['bancos'].items())  # [(código, nome)]
def cli(mm, cod=None):
    return sum(v for k, v in rz['meses'].get(mm, {}).items() if k.endswith('|clientes') and (cod is None or k.split('|')[0] == cod))
def excl(mm): return sum(v for k, v in rz['meses'].get(mm, {}).items() if not k.endswith('|clientes'))
MR = sorted([mm for mm in rz['meses'] if mm in f_ano], key=ordem)  # meses com razão e apuração
MA = sorted(f_ano, key=ordem)
merc = lambda mm: comp.get(mm, {}).get('mercadorias', 0)
tot_rec = sum(cli(mm) for mm in MR); tot_decl_r = sum(f_ano[mm] for mm in MR)
tot_comp = sum(merc(mm) for mm in MA); tot_decl = sum(f_ano.values())
ind = dict(recebido=tot_rec, declarado_periodo_razao=tot_decl_r, vezes=tot_rec / tot_decl_r if tot_decl_r else None,
           compras=tot_comp, declarado=tot_decl, compras_pct=tot_comp / tot_decl if tot_decl else None,
           media_ant=sum(f_ant) / len(f_ant) if f_ant else None, media_ano=tot_decl / len(MA),
           excluido=sum(excl(mm) for mm in MR), lancamentos=sum(rz.get('quantidade', {}).get(mm, 0) for mm in MR))

AZ = '1F4E78'; H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor=AZ); B = Font(bold=True)
RED = PatternFill('solid', fgColor='F8CBAD'); TOT = PatternFill('solid', fgColor='DEEAF6'); M = '#,##0.00'; P = '0.0%'; X = '0.0"x"'
thin = Side(style='thin', color='BFBFBF'); BR = Border(left=thin, right=thin, top=thin, bottom=thin)
def title(ws, t, s=None):
    ws['A1'] = t; ws['A1'].font = Font(bold=True, size=13, color=AZ)
    if s: ws['A2'] = s; ws['A2'].font = Font(italic=True, color='595959')
def hdr(ws, r, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = H; c.fill = HF; c.border = BR; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
def put(ws, r, c, v, fmt=None, bold=False, fill=None, wrap=False):
    x = ws.cell(r, c, v); x.border = BR
    if fmt: x.number_format = fmt
    if bold: x.font = B
    if fill: x.fill = fill
    if wrap: x.alignment = Alignment(wrap_text=True, vertical='top')
    return x
def W(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v
br = lambda v: f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
per_r = f'{MR[0]} a {MR[-1]}' if MR else '–'; per_a = f'{MA[0]} a {MA[-1]}'
wb = Workbook()
# 1 resumo
r = wb.active; r.title = '1-Resumo'
title(r, f'{a.empresa} – Faturamento declarado × entradas', 'Relatório interno ROF Contabilidade – uso da equipe, não enviar ao cliente nesta forma')
hdr(r, 4, ['Indicador', 'Valor'])
rows = [(f'Faturamento declarado – {ant} (meses disponíveis: {len(f_ant)})', sum(f_ant), M), (f'Faturamento declarado – {per_a}', tot_decl, M),
        (f'Média mensal declarada {ant}', ind['media_ant'], M), (f'Média mensal declarada {ano}', ind['media_ano'], M),
        ('Variação da média mensal declarada', (ind['media_ano'] / ind['media_ant'] - 1) if ind['media_ant'] else None, P),
        (f'Recebimentos de clientes – todos os bancos – {per_r}', tot_rec, M), (f'Faturamento declarado – {per_r}', tot_decl_r, M),
        ('Recebimentos de clientes ÷ faturamento declarado', '=B10/B11', X),
        (f'Compras de mercadorias (NF de entrada) – {per_a}', tot_comp, M), ('Compras de mercadorias ÷ faturamento declarado', '=B13/B6', P)]
for i, (t, v, fm) in enumerate(rows, 5): put(r, i, 1, t); put(r, i, 2, v, fm)
alertas = ['O que chama a atenção:']
if ind['vezes'] and ind['vezes'] > 1.2:
    alertas.append(f'• Recebimentos de clientes nos bancos ({per_r}) somam R$ {br(tot_rec)}, {ind["vezes"]:.1f} vezes o faturamento declarado no mesmo período. Já sem transferências próprias/sócio e aplicações (R$ {br(ind["excluido"])} excluídos).')
    r['B12'].fill = RED; r['B12'].font = B
if ind['compras_pct'] and ind['compras_pct'] > 0.8:
    alertas.append(f'• Compras de mercadorias = {ind["compras_pct"]:.0%} do faturamento declarado. A LC 123 prevê exclusão de ofício quando as aquisições passam de 80% dos ingressos (art. 29, X), salvo aumento de estoque justificado.')
    r['B14'].fill = RED; r['B14'].font = B
if ind['media_ant'] and ind['media_ano'] / ind['media_ant'] < 0.7:
    alertas.append(f'• A média mensal declarada caiu de R$ {br(ind["media_ant"])} ({ant}) para R$ {br(ind["media_ano"])} ({ano}).'); r['B9'].fill = RED
if len(alertas) == 1: alertas.append('• Sem divergência relevante pelos critérios da skill (recebido até 1,2x o declarado; compras até 80%).')
alertas.append('• Esses dados chegam ao fisco por cruzamento automático: NF-e de entrada, cartão/PIX informados pelas instituições de pagamento, e-Financeira e DIMP.')
for i, t in enumerate(alertas, 17):
    c = r.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical='top'); r.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    r.row_dimensions[i].height = 15 if i == 17 else 44
    if i == 17: c.font = B
W(r, dict(A=78, B=20))
# 2 mês a mês
m = wb.create_sheet(f'2-Mês a mês {ano}')
title(m, f'Mês a mês {ano}: declarado × recebimentos de clientes × compras', 'Recebimentos: Razão do Domínio (débito em banco com contrapartida 504), sem transferências próprias/sócio e sem aplicações')
nb = len(contas); cT = 3 + nb
hdr(m, 4, ['Competência', 'Faturamento declarado'] + [f'{nome} (conta {cod})' for cod, nome in contas] +
   ['Total recebido de clientes', 'Nº lançamentos', 'Excluídos (próprias/aplicações)', 'Compras de mercadorias', 'Recebido ÷ declarado', 'Compras ÷ declarado'])
L = lambda k: m.cell(1, k).column_letter
for i, mm in enumerate(MA, 5):
    put(m, i, 1, mm); put(m, i, 2, f_ano[mm], M)
    if mm in rz['meses']:
        for j, (cod, _) in enumerate(contas, 3): put(m, i, j, cli(mm, cod), M)
        put(m, i, cT, f'=SUM(C{i}:{L(cT-1)}{i})', M, True); put(m, i, cT + 1, rz.get('quantidade', {}).get(mm, 0), '#,##0')
        put(m, i, cT + 2, excl(mm), M); put(m, i, cT + 4, f'={L(cT)}{i}/B{i}', X, fill=RED)
    else:
        for j in range(3, cT + 3): put(m, i, j, '–')
        put(m, i, cT + 4, '–')
    put(m, i, cT + 3, merc(mm), M); put(m, i, cT + 5, f'={L(cT+3)}{i}/B{i}', P, fill=RED)
rt = 5 + len(MA); put(m, rt, 1, 'Total', bold=True, fill=TOT)
for j in range(2, cT + 4): put(m, rt, j, f'=SUM({L(j)}5:{L(j)}{rt-1})', M, True, TOT)
put(m, rt, cT + 4, f'={L(cT)}{rt}/SUMPRODUCT(B5:B{rt-1},--({L(cT)}5:{L(cT)}{rt-1}<>"–"))', X, True, RED)
put(m, rt, cT + 5, f'={L(cT+3)}{rt}/B{rt}', P, True, RED)
m.cell(rt + 1, 1, f'Razão disponível: {per_r}. Meses sem razão aparecem com "–".')
m.column_dimensions['A'].width = 13
for j in range(2, cT + 6): m.column_dimensions[L(j)].width = 16
ch = BarChart(); ch.type = 'col'; ch.title = 'Declarado × recebido de clientes × compras'
for col in (2, cT, cT + 3): ch.add_data(Reference(m, min_col=col, min_row=4, max_row=rt - 1), titles_from_data=True)
ch.set_categories(Reference(m, min_col=1, min_row=5, max_row=rt - 1))
for s_, col in zip(ch.series, ['2A78D6', 'EB6834', '1BAF7A']): s_.graphicalProperties.solidFill = col; s_.graphicalProperties.line.solidFill = col
ch.y_axis.numFmt = '#,##0'; ch.height = 9; ch.width = 26; ch.legend.position = 'b'; ch.x_axis.delete = False; ch.y_axis.delete = False
m.add_chart(ch, f'A{rt+3}')
# 3 histórico
h = wb.create_sheet('3-Faturamento declarado')
title(h, 'Faturamento declarado no PGDAS por mês', 'Fonte: aba Anexo dos relatórios do Domínio')
hdr(h, 4, ['Competência', 'Faturamento declarado'])
for i, (mm, v) in enumerate(serie, 5): put(h, i, 1, mm); put(h, i, 2, v, M)
W(h, dict(A=14, B=22))
lc = LineChart(); lc.title = 'Faturamento declarado por mês'; lc.add_data(Reference(h, min_col=2, min_row=4, max_row=4 + len(serie)), titles_from_data=True)
lc.set_categories(Reference(h, min_col=1, min_row=5, max_row=4 + len(serie))); lc.series[0].graphicalProperties.line.solidFill = '2A78D6'
lc.y_axis.numFmt = '#,##0'; lc.legend = None; lc.height = 8; lc.width = 22; lc.x_axis.delete = False; lc.y_axis.delete = False
h.add_chart(lc, 'D4')
# 4 riscos
b = wb.create_sheet('4-Riscos e base legal')
title(b, 'Riscos fiscais e base legal (conferir texto em Legislacao-Simples-Nacional/)')
hdr(b, 3, ['Risco', 'O que diz a lei', 'Base legal'])
leg = [('Presunção de omissão de receita', 'O Simples aplica as presunções de omissão de receita da legislação do IR e das contribuições.', 'LC 123/2006, art. 34'),
       ('Depósitos sem origem comprovada', 'Créditos em conta sem origem comprovada presumem-se receita omitida.', 'Lei 9.430/1996, art. 42 (fora da pasta – conferir texto)'),
       ('Compras acima de 80% dos ingressos', 'Exclusão de ofício quando as aquisições para comercialização superam 80% dos ingressos, salvo aumento de estoque justificado.', 'LC 123/2006, art. 29, X'),
       ('Movimentação financeira não identificada', 'Exclusão quando falta livro-caixa ou não é possível identificar a movimentação financeira, inclusive bancária.', 'LC 123/2006, art. 29, VIII'),
       ('Documento fiscal', 'Obrigação de emitir documento fiscal; exclusão se houver descumprimento reiterado.', 'LC 123/2006, art. 26, I, e art. 29, XI'),
       ('Efeitos da exclusão de ofício', 'Vale desde o mês da infração; impede voltar ao Simples por 3 anos (10 anos se houver fraude).', 'LC 123/2006, art. 29, §§ 1º e 2º'),
       ('Exclusão por débitos', 'Débito sem exigibilidade suspensa impede o Simples; permanece se regularizar em até 30 dias da ciência.', 'LC 123/2006, art. 17, V, e art. 31, § 2º')]
for i, row in enumerate(leg, 4):
    for j, v in enumerate(row, 1): put(b, i, j, v, wrap=True)
    b.row_dimensions[i].height = 48
W(b, dict(A=32, B=70, C=34))
# 5 próximos passos
dd = wb.create_sheet('5-Próximos passos')
title(dd, 'Próximos passos sugeridos (para a equipe)')
for i, t in enumerate(['Conversar com o cliente: entender a diferença (vendas sem nota? outro CNPJ/CPF? mudança de atividade?).',
                       'Levantar todas as vendas (PDV/caixa, maquininhas, PIX) e conciliar com as notas emitidas.',
                       'Havendo receita não declarada: avaliar retificação dos PGDAS com denúncia espontânea (CTN, art. 138) antes de qualquer procedimento fiscal.',
                       'Conferir a conta 504 (Clientes): recebimentos baixando clientes que não foram faturados geram saldo credor.',
                       'Só refazer o estudo da reforma (2027) depois de ter uma base de faturamento confiável.'], 3):
    c = dd.cell(i, 1, f'{i-2}. {t}'); c.alignment = Alignment(wrap_text=True, vertical='top'); dd.row_dimensions[i].height = 30
W(dd, dict(A=120))
wb.calculation.fullCalcOnLoad = True; wb.save(a.saida)
if a.resumo: json.dump(ind, open(a.resumo, 'w'), ensure_ascii=False, indent=1)
print('ok', a.saida, {k: round(v, 2) if isinstance(v, float) else v for k, v in ind.items()})
