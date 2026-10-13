"""Planilha de conferência de NFS-e (prestadas ou tomadas) com apontamento de inconsistências.

python planilha_nfse.py prestado notas.json empresa.json saida.xlsx [--competencia AAAA-MM]
python planilha_nfse.py tomado   notas.json empresa.json saida.xlsx [--competencia AAAA-MM]

notas.json      lista gerada por ler_nfse_xml.py (ou montada à mão a partir do PDF), podendo ter, por nota:
                  "apontamentos": ["texto", ...]   inconsistências que você (Claude) encontrou lendo a nota
                                  (prefixo "ERRO: " ou "ALERTA: " define a gravidade; sem prefixo = APONTAMENTO)
                  (prestado) "anexo": anexo decidido para a atividade da nota (senão usa a tabela pelo item da LC 116)
                  (tomado) "enquadramento": chave de tabelas.RETENCOES (senão usa o sugerido pelo item, com alerta),
                  "prestador_simples"/"prestador_mei"/"prestador_anexo_iv": bool (senão lê do regime/texto da nota),
                  "cessao_mao_obra": bool, "base_inss": valor
                  "valor_zero_ok": true quando a nota tem mesmo valor R$ 0,00 (ex.: pedido de crédito em cartão de benefícios)
empresa.json    saída de consulta_cnpj.py --json, opcionalmente com "conclusao_anexo": "anexo de cada atividade e o porquê"
                  e "apontamentos_gerais": ["ALERTA: ...", ...] (cruzamento com relatórios da prefeitura/portal etc.)
"""
import json, os, re, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tabelas import retencoes_esperadas, RETENCOES, anexo_item, ITENS_RETENCAO_ISS

F = lambda **k: Font(**{'name': 'Arial', 'size': 9, **k})
thin = Side(style='thin', color='BFBFBF'); B = Border(top=thin, bottom=thin, left=thin, right=thin)
HEAD = PatternFill('solid', fgColor='D9E1F2'); ERRO = PatternFill('solid', fgColor='F8CBAD')
ALERTA = PatternFill('solid', fgColor='FFF2CC'); OK = PatternFill('solid', fgColor='E2EFDA')
MOEDA = '#,##0.00;-#,##0.00;-'
TOL = 0.02


def v(x):
    return float(x or 0)


def doc(d):
    d = re.sub(r'\D', '', d or '')
    if len(d) == 14: return f'{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'
    if len(d) == 11: return f'{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}'
    return d


def data(s):
    return (s or '')[:10]


def cancelada(n):
    return 'cancel' in str(n.get('situacao', '')).lower()


def total_ret(n):
    return round(sum(v(n.get(k)) for k in ('irrf', 'pis', 'cofins', 'csll', 'inss'))
                 + (v(n.get('valor_iss')) if str(n.get('iss_retido', '')).startswith('Sim') else 0), 2)


def apontamentos(n):
    """Apontamentos da leitura; prefixo "ERRO: " ou "ALERTA: " define a gravidade."""
    out = []
    for t in n.get('apontamentos', []):
        g = next((x for x in ('ERRO', 'ALERTA') if t.startswith(x + ': ')), None)
        out.append((g, t[len(g) + 2:]) if g else ('APONTAMENTO', t))
    return out


def checks_comuns(n, emp, comp, lado):
    """lado = 'prestado' ou 'tomado'. Retorna lista de (gravidade, texto)."""
    a = []
    if cancelada(n):
        a.append(('ALERTA', 'Nota CANCELADA - não entra no faturamento/apuração; conferir se o cancelamento está no portal'))
    if n.get('substituida'):
        a.append(('ALERTA', f"Nota substitui outra ({n['substituida']}) - conferir se a substituída não foi lançada"))
    if comp and n.get('competencia') and n['competencia'][:7] != comp:
        a.append(('ERRO', f"Competência {n['competencia'][:7]} diferente do mês em análise {comp}"))
    if comp and n.get('data_emissao') and n['data_emissao'][:7] != comp and not (n.get('competencia') or '').startswith(comp):
        a.append(('ALERTA', f"Emitida em {data(n['data_emissao'])}, fora do mês {comp}"))
    vs = v(n.get('valor_servico'))
    if not vs and n.get('valor_zero_ok'):
        a.append(('APONTAMENTO', 'Nota com valor do serviço R$ 0,00 (documento só informativo, ex.: crédito de benefícios) - confirmado na leitura'))
    elif not vs:
        a.append(('ERRO', 'Valor do serviço não encontrado/zerado'))
    base, al, iss = n.get('base_iss'), n.get('aliquota_iss'), n.get('valor_iss')
    if base is not None and al is not None and iss is not None and abs(v(base) * v(al) / 100 - v(iss)) > TOL:
        a.append(('ERRO', f'ISS destacado {v(iss):.2f} ≠ base {v(base):.2f} × {v(al)}% = {v(base) * v(al) / 100:.2f}'))
    if base is not None and vs and abs(v(base) - (vs - v(n.get('deducoes')) - v(n.get('desconto_incond')))) > TOL:
        a.append(('ALERTA', f'Base do ISS {v(base):.2f} diferente de valor - deduções - desc. incondicionado'))
    if n.get('valor_liquido') is not None and vs:
        calc = round(vs - v(n.get('desconto_incond')) - total_ret(n), 2)
        if abs(calc - v(n['valor_liquido'])) > TOL:
            a.append(('ERRO', f"Valor líquido da nota {v(n['valor_liquido']):.2f} ≠ calculado {calc:.2f} (valor - desc. incond. - retenções)"))
    if str(n.get('iss_retido', '')).startswith('Sim') and len(re.sub(r'\D', '', n.get('tomador_doc') or '')) == 11:
        a.append(('ERRO', 'ISS marcado como retido com tomador pessoa física (CPF) - PF não é substituto tributário'))
    lp, mt = (n.get('municipio_prestacao') or '').strip(), (n.get('tomador_municipio') or '').strip()
    if lado == 'prestado' and lp and mt and not lp.isdigit() and not mt.isdigit() and lp.lower() != mt.lower():
        a.append(('ALERTA', f'Local da prestação ({lp}) diferente do município do tomador ({mt}) - confirmar onde o serviço foi feito'))
    if not (n.get('descricao') or '').strip():
        a.append(('ERRO', 'Nota sem discriminação do serviço'))
    elif len(n['descricao'].strip()) < 15:
        a.append(('ALERTA', f"Discriminação genérica/curta: \"{n['descricao'].strip()}\""))
    if not n.get('item_lc116') and not n.get('atividade'):
        a.append(('ALERTA', 'Item da LC 116 / código de tributação não identificado'))
    return a


def prestado(notas, emp, comp):
    cnpj = re.sub(r'\D', '', emp.get('cnpj', ''))
    simples, mei = emp.get('simples'), emp.get('mei')
    linhas, incons = [], []
    for n in notas:
        a = checks_comuns(n, emp, comp, 'prestado')
        it, an, desc_it, _ = anexo_item(n.get('item_lc116'))
        n['_item'] = it; n['_anexo'] = n.get('anexo') or an; n['_anexo_desc'] = desc_it
        if simples and not mei:
            if 'IMPEDITIVO' in n['_anexo']:
                a.append(('ERRO', f'Atividade {it} ({desc_it}) é IMPEDITIVA ou duvidosa no Simples - conferir'))
            elif n['_anexo'] == '?' or ('conferir' in desc_it and not n.get('anexo')):
                a.append(('ALERTA', f'Anexo da atividade {it or "(sem item)"} não definido pela tabela: {desc_it} - decidir e preencher "anexo" na nota'))
            if n['_anexo'] == 'IV':
                a.append(('ALERTA', f'Atividade {it} no Anexo IV: CPP fora do DAS (INSS patronal à parte) e sujeita a retenção de INSS 11% se houver cessão/empreitada'))
        if cnpj and re.sub(r'\D', '', n.get('prestador_cnpj') or '') not in ('', cnpj):
            a.append(('ERRO', f"Prestador da nota ({doc(n['prestador_cnpj'])}) não é a empresa {doc(cnpj)}"))
        if not re.sub(r'\D', '', n.get('tomador_doc') or ''):
            a.append(('ALERTA', 'Tomador sem CPF/CNPJ'))
        fed = sum(v(n.get(k)) for k in ('irrf', 'pis', 'cofins', 'csll'))
        if simples and fed:
            a.append(('ERRO', f'Empresa do Simples com retenção federal destacada (IRRF/PIS/COFINS/CSLL = {fed:.2f}) - optante não sofre (IN SRF 765/2007 e IN SRF 459/2004 art. 3º); só órgão público federal retém (IN RFB 1.234/2012)'))
        if simples and v(n.get('inss')):
            a.append(('ALERTA', f"INSS retido {v(n['inss']):.2f}: no Simples só cabe para Anexo IV (cessão de mão de obra/empreitada) - conferir"))
        al = n.get('aliquota_iss')
        if simples and not mei and al is not None and v(al) and not (2 <= v(al) <= 5):
            a.append(('ERRO', f'Alíquota de ISS {v(al)}% fora do intervalo do Simples (2% a 5%)'))
        if simples and str(n.get('prestador_regime_sn', '')).startswith('Não'):
            a.append(('ERRO', 'Nota emitida como NÃO optante do Simples, mas o CNPJ é optante - corrigir o cadastro no emissor'))
        if simples is False and 'Simples' in str(n.get('prestador_regime_sn', '')):
            a.append(('ERRO', 'Nota emitida como optante do Simples, mas o CNPJ NÃO é optante na Receita'))
        if simples and not mei and 'simples' not in (n.get('info_complementar') or '').lower() and n.get('layout') != 'NFS-e Nacional':
            a.append(('ALERTA', 'Sem a frase "Documento emitido por ME ou EPP optante pelo Simples Nacional" nas informações complementares'))
        a += apontamentos(n)
        linhas.append((n, a))
    # sequência
    nums = sorted({int(re.sub(r'\D', '', n['numero'])) for n in notas if re.sub(r'\D', '', n.get('numero') or '')})
    dup = [k for k, c in Counter(n.get('numero') for n in notas).items() if c > 1]
    faltam = [x for x in range(nums[0], nums[-1] + 1) if x not in nums] if nums else []
    return linhas, {'nums': nums, 'faltam': faltam, 'dup': dup}


def tomado(notas, emp, comp):
    cnpj = re.sub(r'\D', '', emp.get('cnpj', ''))
    linhas = []
    vistos = Counter((re.sub(r'\D', '', n.get('prestador_cnpj') or ''), n.get('numero')) for n in notas)
    for n in notas:
        a = checks_comuns(n, emp, comp, 'tomado')
        if cnpj and re.sub(r'\D', '', n.get('tomador_doc') or '') != cnpj:
            a.append(('ERRO', f"Tomador da nota ({doc(n.get('tomador_doc'))}) não é a empresa {doc(cnpj)}"))
        if vistos[(re.sub(r'\D', '', n.get('prestador_cnpj') or ''), n.get('numero'))] > 1:
            a.append(('ERRO', 'Nota em duplicidade (mesmo prestador e número)'))
        it, _, desc_it, enq_sug = anexo_item(n.get('item_lc116'))
        n['_item'] = it; n['_item_desc'] = desc_it
        texto = ' '.join(str(n.get(k) or '') for k in ('info_complementar', 'descricao')).lower()
        reg = str(n.get('prestador_regime_sn') or '')
        if 'prestador_mei' not in n:
            n['prestador_mei'] = reg == 'MEI' or 'microempreendedor individual' in texto
        if 'prestador_simples' not in n:
            n['prestador_simples'] = 'Simples' in reg or 'simples nacional' in texto
            if ('Simples' in reg and 'não optante' in texto) or (reg.startswith('Não') and 'optante pelo simples' in texto):
                a.append(('ALERTA', 'Regime do prestador contraditório na nota (campo x texto) - conferir se é optante do Simples'))
        if not reg and 'simples' not in texto and not n.get('prestador_mei'):
            a.append(('ALERTA', 'Nota não informa se o prestador é do Simples - considerado NÃO optante'))
        if n.get('prestador_simples') and 'prestador_anexo_iv' not in n and enq_sug in ('construcao_empreitada', 'limpeza_conservacao', 'seguranca_vigilancia'):
            n['prestador_anexo_iv'] = True
            a.append(('ALERTA', f'Prestador do Simples com atividade de Anexo IV ({it}) - sofre INSS 11% se houver cessão/empreitada'))
        if not n.get('enquadramento'):
            if enq_sug:
                n['enquadramento'] = enq_sug
                a.append(('ALERTA', f'Enquadramento "{enq_sug}" sugerido pelo item {it} ({desc_it}) - confirmar pela discriminação'))
            else:
                a.append(('ERRO', 'Atividade sem item da LC 116 reconhecido - enquadrar pela discriminação (campo "enquadramento")'))
        if it in ITENS_RETENCAO_ISS and not str(n.get('iss_retido', '')).startswith('Sim') and not n.get('prestador_mei'):
            a.append(('ALERTA', f'Item {it} tem retenção de ISS pelo tomador (LC 116 art. 6º §2º II) e a nota não marca ISS retido - conferir local da prestação e lei do município'))
        if not n.get('enquadramento'):
            esp = None
        else:
            esp = retencoes_esperadas(n)
            dest = {'irrf': v(n.get('irrf')), 'csrf': round(v(n.get('pis')) + v(n.get('cofins')) + v(n.get('csll')), 2), 'inss': v(n.get('inss'))}
            for k, nome in (('irrf', 'IRRF'), ('csrf', 'CSRF (PIS+COFINS+CSLL)'), ('inss', 'INSS')):
                if abs(dest[k] - esp[k]) > TOL:
                    if esp[k] and not dest[k]:
                        a.append(('ERRO', f'{nome}: deveria reter {esp[k]:.2f} e a nota não destaca - a empresa (tomadora) é responsável pelo recolhimento'))
                    elif dest[k] and not esp[k]:
                        a.append(('ALERTA', f'{nome}: nota destaca {dest[k]:.2f}, mas não se esperava retenção ({esp["obs"]})'))
                    else:
                        a.append(('ERRO', f'{nome}: destacado {dest[k]:.2f} ≠ esperado {esp[k]:.2f}'))
            n['_dest'] = dest
        n['_esp'] = esp
        a += apontamentos(n)
        linhas.append((n, a))
    return linhas


def cab(ws, row, hs, larg):
    for c, (h, w) in enumerate(zip(hs, larg), 1):
        x = ws.cell(row, c, h); x.font = F(bold=True); x.fill = HEAD; x.border = B
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[x.column_letter].width = w
    ws.row_dimensions[row].height = 30


def put(ws, r, vals, money=(), fill=None):
    for c, val in enumerate(vals, 1):
        x = ws.cell(r, c, val); x.font = F(); x.border = B
        x.alignment = Alignment(vertical='top', wrap_text=isinstance(val, str) and len(val) > 18)
        if c in money: x.number_format = MOEDA
        if fill: x.fill = fill


def gravidade_fill(a):
    g = {x[0] for x in a}
    return ERRO if 'ERRO' in g else ALERTA if g else OK


def titulo(ws, txt, sub):
    ws['A1'] = txt; ws['A1'].font = F(bold=True, size=12)
    ws['A2'] = sub; ws['A2'].font = F(italic=True, color='666666')


def aba_empresa(wb, emp, fat=None):
    ws = wb.create_sheet('Empresa e CNAEs')
    titulo(ws, f"{emp.get('razao_social')} - CNPJ {doc(emp.get('cnpj'))}", f"Consulta: {emp.get('fonte', '')}")
    info = [('Situação cadastral', emp.get('situacao')), ('Município', emp.get('municipio')),
            ('Natureza jurídica', emp.get('natureza_juridica')), ('Porte', emp.get('porte')),
            ('Optante do Simples', 'Sim' if emp.get('simples') else 'Não'), ('Data opção Simples', emp.get('data_opcao_simples')),
            ('Data exclusão Simples', emp.get('data_exclusao_simples')), ('MEI', 'Sim' if emp.get('mei') else 'Não'),
            ('Início atividade', emp.get('abertura'))]
    for i, (k, val) in enumerate(info, 4):
        ws.cell(i, 1, k).font = F(bold=True); ws.cell(i, 2, val).font = F()
    r = 4 + len(info) + 1
    cab(ws, r, ['Tipo', 'CNAE', 'Descrição', 'Anexo sugerido', 'Observação'], [12, 12, 60, 16, 70])
    for i, c in enumerate(emp.get('cnaes', []), r + 1):
        put(ws, i, [c['tipo'], c['cnae'], c['descricao'], c['anexo_sugerido'], c['obs_anexo']])
    r = r + len(emp.get('cnaes', [])) + 2
    if fat:
        cab(ws, r, ['Anexo', 'Item LC 116', 'Atividade', 'Faturamento do mês', ''], [12, 12, 60, 16, 70])
        for i, ((an, it, d), val) in enumerate(sorted(fat.items()), r + 1):
            put(ws, i, [an, it, d, round(val, 2), ''], money=(4,))
        r += len(fat) + 2
    for k, rot in (('conclusao_anexo', 'CONCLUSÃO - ANEXO DE TRIBUTAÇÃO'),):
        if emp.get(k) not in (None, ''):
            ws.cell(r, 1, rot).font = F(bold=True); x = ws.cell(r, 2, emp[k]); x.font = F(bold=k == 'conclusao_anexo')
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); x.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[r].height = 15 * (1 + len(str(emp[k])) // 120); r += 1
    ws.cell(r + 1, 1, 'O anexo é da ATIVIDADE da nota (item da LC 116). "V (Fator R)" = Anexo V, ou III se o Fator R for >= 28%. Referência: art. 18 LC 123/2006 e Anexo VI da Res. CGSN 140/2018.').font = F(italic=True, color='666666')


def aba_incons(wb, linhas, extra=()):
    ws = wb.create_sheet('Inconsistências', 1)
    titulo(ws, 'Inconsistências e apontamentos', 'ERRO = corrigir/esclarecer antes de apurar | ALERTA = conferir | APONTAMENTO = observação da leitura')
    cab(ws, 4, ['Nota', 'Data', 'Participante', 'Gravidade', 'Descrição', 'Resolvido?', 'Observação'], [10, 11, 34, 14, 90, 11, 30])
    r = 5
    for g, t in extra:
        put(ws, r, ['Geral', '', '', g, t, '', ''], fill=ERRO if g == 'ERRO' else ALERTA); r += 1
    for n, a in linhas:
        for g, t in sorted(a, key=lambda x: ('ERRO', 'ALERTA', 'APONTAMENTO').index(x[0])):
            put(ws, r, [n.get('numero'), data(n.get('data_emissao')), n.get('_part', ''), g, t, '', ''],
                fill=ERRO if g == 'ERRO' else ALERTA if g == 'ALERTA' else None); r += 1
    if r == 5:
        put(ws, r, ['', '', '', 'OK', 'Nenhuma inconsistência encontrada', '', ''], fill=OK)
    ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:G{max(r - 1, 5)}'


def gerar(tipo, notas, emp, dst, comp=None):
    wb = Workbook(); ws = wb.active
    notas = sorted(notas, key=lambda n: (re.sub(r'\D', '', n.get('prestador_cnpj') or '') if tipo == 'tomado' else '',
                                         int(re.sub(r'\D', '', n.get('numero') or '0') or 0)))
    extra = [(t.split(': ', 1)[0], t.split(': ', 1)[1]) if t.split(': ', 1)[0] in ('ERRO', 'ALERTA') else ('APONTAMENTO', t)
             for t in emp.get('apontamentos_gerais', [])]
    if tipo == 'prestado':
        ws.title = 'Notas Prestadas'
        linhas, seq = prestado(notas, emp, comp)
        titulo(ws, f"NFS-e PRESTADAS - {emp.get('razao_social')} - {comp or ''}",
               f"Sequência {seq['nums'][0] if seq['nums'] else '-'} a {seq['nums'][-1] if seq['nums'] else '-'} | "
               f"Faltando: {', '.join(map(str, seq['faltam'])) or 'nenhuma'} | Duplicadas: {', '.join(map(str, seq['dup'])) or 'nenhuma'}")
        if seq['faltam']:
            extra.append(('ERRO', f"Quebra de sequência - números não recebidos: {', '.join(map(str, seq['faltam']))} (pedir as notas ou confirmar cancelamento)"))
        if seq['dup']:
            extra.append(('ERRO', f"Números repetidos: {', '.join(map(str, seq['dup']))}"))
        H = ['Nº', 'Emissão', 'Compet.', 'Situação', 'CPF/CNPJ tomador', 'Tomador', 'Município do tomador', 'Local da prestação',
             'Município de incidência do ISS', 'Item LC 116 / cód.', 'Atividade prestada',
             'Anexo Simples (atividade)', 'Discriminação', 'Informações complementares', 'Valor serviço', 'Base ISS', 'Alíq. ISS %', 'ISS', 'ISS retido?',
             'IRRF', 'PIS', 'COFINS', 'CSLL', 'INSS', 'Total retido', 'Valor líquido', 'Status']
        cab(ws, 4, H, [7, 10, 9, 10, 17, 30, 16, 16, 16, 10, 30, 14, 45, 40, 12, 12, 7, 10, 10, 9, 9, 9, 9, 9, 10, 12, 10])
        for i, (n, a) in enumerate(linhas, 5):
            n['_part'] = n.get('tomador_nome')
            put(ws, i, [n.get('numero'), data(n.get('data_emissao')), (n.get('competencia') or '')[:7], 'Cancelada' if cancelada(n) else 'Normal',
                        doc(n.get('tomador_doc')), n.get('tomador_nome'), n.get('tomador_municipio'), n.get('municipio_prestacao'),
                        n.get('municipio_incidencia'), n.get('_item') or n.get('item_lc116'),
                        n.get('atividade') or n.get('_anexo_desc'), n['_anexo'] if emp.get('simples') else 'Não optante', n.get('descricao'),
                        n.get('info_complementar'), n.get('valor_servico'), n.get('base_iss'), n.get('aliquota_iss'), n.get('valor_iss'),
                        n.get('iss_retido'), n.get('irrf'), n.get('pis'), n.get('cofins'), n.get('csll'), n.get('inss'),
                        total_ret(n), n.get('valor_liquido'), 'ERRO' if gravidade_fill(a) is ERRO else 'Conferir' if a else 'OK'],
                money=(15, 16, 18, 20, 21, 22, 23, 24, 25, 26))
            ws.cell(i, 27).fill = gravidade_fill(a)
        fim = 4 + len(linhas)
        fat = {}
        for n, _ in linhas:
            if not cancelada(n):
                k = (n['_anexo'], n['_item'], n['_anexo_desc']); fat[k] = fat.get(k, 0) + v(n.get('valor_servico'))
        aba_empresa(wb, emp, fat)
    else:
        ws.title = 'Notas Tomadas'
        linhas = tomado(notas, emp, comp)
        titulo(ws, f"NFS-e TOMADAS - {emp.get('razao_social')} - {comp or ''}",
               'Retenção esperada calculada pelo enquadramento da atividade e pelo regime do prestador (aba Regras de retenção)')
        H = ['Nº', 'Emissão', 'Compet.', 'Situação', 'CNPJ prestador', 'Prestador', 'Local da prestação', 'Município de incidência do ISS',
             'Regime prestador', 'Item LC 116 / cód.',
             'Atividade', 'Discriminação', 'Informações complementares', 'Valor serviço', 'Alíq. ISS %', 'ISS', 'ISS retido?',
             'IRRF nota', 'PIS+COFINS+CSLL nota', 'INSS nota', 'Enquadramento', 'Alíquotas esperadas',
             'IRRF esperado', 'CSRF esperada', 'INSS esperado', 'Valor líquido', 'Obs. retenção', 'Status']
        cab(ws, 4, H, [7, 10, 9, 10, 17, 30, 16, 16, 13, 10, 28, 45, 35, 12, 7, 10, 10, 10, 11, 10, 20, 22, 10, 10, 10, 12, 45, 10])
        for i, (n, a) in enumerate(linhas, 5):
            n['_part'] = n.get('prestador_nome'); e = n.get('_esp') or {}
            reg = n.get('prestador_regime_sn') or ''
            if 'prestador_mei' in n or 'prestador_simples' in n:
                reg = 'MEI' if n.get('prestador_mei') else ('Simples' + (' Anexo IV' if n.get('prestador_anexo_iv') else '')) if n.get('prestador_simples') else 'Não optante'
            put(ws, i, [n.get('numero'), data(n.get('data_emissao')), (n.get('competencia') or '')[:7], 'Cancelada' if cancelada(n) else 'Normal',
                        doc(n.get('prestador_cnpj')), n.get('prestador_nome'), n.get('municipio_prestacao'), n.get('municipio_incidencia'),
                        reg, n.get('_item') or n.get('item_lc116'),
                        n.get('atividade') or n.get('_item_desc'),
                        n.get('descricao'), n.get('info_complementar'), n.get('valor_servico'), n.get('aliquota_iss'), n.get('valor_iss'),
                        n.get('iss_retido'), n.get('irrf'), round(v(n.get('pis')) + v(n.get('cofins')) + v(n.get('csll')), 2), n.get('inss'),
                        n.get('enquadramento'), e.get('aliquotas'), e.get('irrf'), e.get('csrf'), e.get('inss'), n.get('valor_liquido'),
                        e.get('obs'), 'ERRO' if gravidade_fill(a) is ERRO else 'Conferir' if a else 'OK'],
                money=(14, 16, 18, 19, 20, 23, 24, 25, 26))
            ws.cell(i, 28).fill = gravidade_fill(a)
        fim = 4 + len(linhas)
        wr = wb.create_sheet('Regras de retenção')
        titulo(wr, 'Enquadramentos usados', 'Dispensa: IRRF e CSRF <= R$ 10,00; INSS < R$ 10,00. Simples: sem IRRF/CSRF; INSS só Anexo IV. MEI: sem retenção.')
        cab(wr, 4, ['Enquadramento', 'IRRF %', 'CSRF %', 'INSS %', 'Base legal / observação'], [30, 8, 8, 8, 100])
        for i, (k, (a1, a2, a3, d)) in enumerate(RETENCOES.items(), 5):
            put(wr, i, [k, a1, a2, a3, d])
        if emp.get('cnaes'):
            aba_empresa(wb, emp)
    # totais
    t = fim + 1
    ws.cell(t, 1, 'TOTAL (sem canceladas)').font = F(bold=True)
    ativos = [n for n, _ in linhas if not cancelada(n)]
    cols = {'prestado': {15: 'valor_servico', 18: 'valor_iss', 20: 'irrf', 21: 'pis', 22: 'cofins', 23: 'csll', 24: 'inss', 26: 'valor_liquido'},
            'tomado': {14: 'valor_servico', 16: 'valor_iss', 18: 'irrf', 20: 'inss', 26: 'valor_liquido'}}[tipo]
    for c, k in cols.items():
        x = ws.cell(t, c, round(sum(v(n.get(k)) for n in ativos), 2)); x.font = F(bold=True); x.number_format = MOEDA
    if tipo == 'tomado':
        for c, k in ((23, 'irrf'), (24, 'csrf'), (25, 'inss')):
            x = ws.cell(t, c, round(sum(v((n.get('_esp') or {}).get(k)) for n in ativos), 2)); x.font = F(bold=True); x.number_format = MOEDA
    ws.freeze_panes = 'B5'; ws.auto_filter.ref = f'A4:{ws.cell(4, ws.max_column).column_letter}{max(fim, 5)}'
    aba_incons(wb, linhas, extra)
    wb.save(dst)
    ne = sum(1 for _, a in linhas for g, _ in a if g == 'ERRO') + sum(1 for g, _ in extra if g == 'ERRO')
    na = sum(1 for _, a in linhas for g, _ in a if g != 'ERRO')
    print(f'{len(linhas)} nota(s) | {ne} erro(s) | {na} alerta(s)/apontamento(s) -> {dst}')


if __name__ == '__main__':
    a = sys.argv[1:]
    opt = lambda k: a[a.index(k) + 1] if k in a else None
    comp = opt('--competencia')
    tipo, fn, fe, dst = a[:4]
    gerar(tipo, json.load(open(fn, encoding='utf-8')), json.load(open(fe, encoding='utf-8')), dst, comp)
