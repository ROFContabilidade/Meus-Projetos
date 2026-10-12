"""Relatório das NFS-e de serviços tomados do mês (Relacao_Servicos_Tomados_Cosmetici_MM-AAAA.xlsx) e a base de dados do mês
(dados/nfse_tomadas_MM-AAAA.json). Modelo pedido pela usuária - simples e detalhado: 1ª aba Todas as notas, 2ª Com retenção
(com o total de cada retenção e a guia), 3ª Sem retenção.
Valores e retenções EXATAMENTE como estão no XML da NFS-e (padrão nacional) - nada é recalculado."""
import datetime as dt
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import planilha_entradas as pe

F = pe.F
H = pe.HEAD
RET = pe.RET
TOT = pe.TOT
M = pe.MOEDA

# guia de cada retenção (informativo, para conferência do recolhimento)
GUIAS = [('ISS retido', 'iss_retido', 'Guia do ISS retido do município de incidência (tomador recolhe)', 'conforme o calendário do município'),
         ('IRRF', 'irrf', 'DARF 1708 - IRRF sobre serviços de PJ', 'até o dia 20 do mês seguinte (último dia útil do 2º decêndio)'),
         ('PIS/COFINS/CSLL (CSRF)', 'csll', 'DARF 5952 - PIS/COFINS/CSLL retidos (4,65%)', 'até o dia 20 do mês seguinte (último dia útil do 2º decêndio)'),
         ('INSS retido', 'inss', 'DCTFWeb / DARF previdenciário - retenção de 11% (cessão de mão de obra)', 'até o dia 20 do mês seguinte')]

COLS = [('Emissão', 11, 'data', 'd'), ('Nº NFS-e', 10, 'numero', ''), ('Prestador', 38, 'emitente', ''), ('CNPJ', 19, 'cnpj_fmt', ''),
        ('Serviço (resumo da descrição do XML)', 50, 'descricao', ''), ('Simples Nacional', 9, 'simples', ''), ('Valor do serviço', 13, 'vNF', 'm'),
        ('ISS da nota', 11, 'iss', 'm'), ('ISS retido', 11, 'iss_retido', 'm'), ('IRRF retido', 11, 'irrf', 'm'),
        ('PIS/COFINS/CSLL retidos', 12, 'csll', 'm'), ('INSS retido', 11, 'inss', 'm'), ('Total retido', 12, 'total_ret', 'm'),
        ('Líquido a pagar', 13, 'liquido', 'm'), ('Retenção', 22, 'tem_ret', ''), ('Acumulador', 11, 'acumulador', ''),
        ('Situação', 10, 'situacao', '')]


def linhas(a, xml_nfse, rel_nfse=None, forn=None):
    import cruzamento_entradas as ce
    import lancamentos_contabeis as lc
    from creditos import Credito
    rows = pe.nfse(a.cfg, xml_nfse)
    por = {(x['prest_cnpj'], str(int(x['numero'])) if str(x['numero']).isdigit() else str(x['numero'])): x for x in xml_nfse or []}
    conf = {(ln[3], str(ln[0])): ln[24] for ln in ce.nfse(xml_nfse, rel_nfse)} if rel_nfse else {}
    cred = Credito(a, forn or {}, [], a.cfg.get('natureza_entradas', {})).analisar()
    lanc, pend = lc.montar(a, cred, {}, {}, xml_nfse, forn or {})
    conta = {}
    for x in lanc:
        if x['doc'] == 'NFS-e':
            d = next(ln for _, b in x['blocos'] for ln in b if ln[0])
            conta[(x['cnpj'], str(x['numero']))] = f'D {d[0]} / C {d[1] or "506"}'
    for x in pend:
        if x['doc'] == 'NFS-e':
            conta[(x['cnpj'], str(x['numero']))] = 'definir: ' + x['motivo']
    for r in rows:
        k = (r['cnpj'], str(r['numero']))
        x = por.get(k) or {}
        comp = x.get('competencia')
        r['competencia'] = dt.date.fromisoformat(str(comp)[:10]) if comp else None
        r['data'] = dt.date.fromisoformat(str(r['data'])[:10]) if r['data'] and not isinstance(r['data'], dt.date) else r['data']
        d = ' '.join(str(x.get('descricao') or r.get('servico') or '').replace('|', ' ').split())
        r['descricao'] = d if len(d) <= 120 else d[:117] + '...'
        c = r['cnpj'] or ''
        r['cnpj_fmt'] = f'{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}' if len(c) == 14 else c
        r['conta'] = conta.get(k, '')
        r['conferencia'] = conf.get(k, 'sem relatório ABRASF' if not rel_nfse else 'não está no relatório')
    return rows


def gerar(a, xml_nfse, destino, rel_nfse=None, forn=None, base_json=None):
    rows = linhas(a, xml_nfse, rel_nfse, forn)
    com = [r for r in rows if r['tem_ret'].startswith('SIM')]
    sem = [r for r in rows if not r['tem_ret'].startswith('SIM')]
    emp = a.cfg['empresa']['nome']
    mm = f'{a.mes[5:]}/{a.mes[:4]}'
    wb = Workbook()
    for r in rows:
        r['tem_ret'] = r['tem_ret'].replace('SIM - ', 'Sim - ')
    cor = lambda ln: RET if ln['tem_ret'].startswith('Sim') else None
    pe._aba(wb, 'Todas as notas', f'{emp} - serviços tomados {mm}: todas as NFS-e, por data. Amarelo = nota com retenção. '
            'Valores exatamente como no XML.', COLS, rows, cor, primeira=True,
            contar=lambda i, u: [('Notas com retenção', f'=COUNTIF(O{i}:O{u},"Sim*")'), ('Notas sem retenção', f'=COUNTIF(O{i}:O{u},"Não")')])
    pe._aba(wb, 'Com retenção', f'NFS-e com retenção - {mm}. A Cosmetici recolhe o valor retido e paga ao prestador o líquido.',
            COLS, [r for r in rows if r['tem_ret'].startswith('Sim')], lambda ln: RET,
            contar=lambda i, u: [(f'{rot}: {guia}', f'=SUM({col}{i}:{col}{u})') for rot, col, guia in
                                 (('ISS retido', 'I', 'guia do ISS retido do município'), ('IRRF', 'J', 'DARF 1708 - até o dia 20 do mês seguinte'),
                                  ('PIS/COFINS/CSLL', 'K', 'DARF 5952 - até o dia 20 do mês seguinte'), ('INSS', 'L', 'DCTFWeb - até o dia 20'))])
    pe._aba(wb, 'Sem retenção', f'NFS-e sem retenção - {mm}. O valor integral é pago ao prestador.', COLS,
            [r for r in rows if not r['tem_ret'].startswith('Sim')])
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    if base_json:
        with open(base_json, 'w', encoding='utf-8') as fh:
            json.dump({'mes': a.mes, 'gerado_de': 'XML das NFS-e (padrão nacional) da pasta Doc\'s/Recebidas',
                       'notas': [{k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in x.items() if k != 'tipo'} for x in rows]},
                      fh, ensure_ascii=False, indent=1)
    return rows, sem, com


def auditar_xml(planilha, caminhos_xml):
    """Conferência obrigatória (regra da usuária): relê cada XML de NFS-e de forma INDEPENDENTE (ElementTree direto, sem o leitor
    que monta a planilha) e compara com a aba "Todas as notas", nota a nota e campo a campo. Devolve (notas, campos, diferenças)."""
    import xml.etree.ElementTree as ET
    from openpyxl import load_workbook
    num = lambda t: float(t) if t else 0.0
    xml = {}
    for f in caminhos_xml:
        root = ET.parse(f).getroot()
        tags = {}
        for e in root.iter():
            tags.setdefault(e.tag.split('}')[-1], e.text)
        emit = next((e for e in root.iter() if e.tag.endswith('}emit')), None)
        cnpj = next((c.text for c in emit if c.tag.endswith('CNPJ')), '') if emit is not None else ''
        if not tags.get('nNFSe'):
            continue
        ret_iss = tags.get('tpRetISSQN') in ('2', '3')
        xml[(cnpj, str(int(tags['nNFSe'])))] = {
            'Valor do serviço': num(tags.get('vServ')), 'ISS da nota': num(tags.get('vISSQN')),
            'ISS retido': num(tags.get('vISSQN')) if ret_iss else 0.0, 'IRRF retido': num(tags.get('vRetIRRF')),
            'PIS/COFINS/CSLL retidos': num(tags.get('vRetCSLL')), 'INSS retido': num(tags.get('vRetCP')),
            'Total retido': num(tags.get('vTotalRet')), 'Líquido a pagar': num(tags.get('vLiq')), '_ret': ret_iss or num(tags.get('vTotalRet')) > 0}
    ws = load_workbook(planilha)['Todas as notas']
    h = [c.value for c in ws[4]]
    dif, n, campos, vistos = [], 0, 0, set()
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not isinstance(r[1], int):
            continue
        k = (str(r[h.index('CNPJ')]).replace('.', '').replace('/', '').replace('-', ''), str(r[1]))
        vistos.add(k)
        n += 1
        x = xml.get(k)
        if not x:
            dif.append(f'NFS-e {k[1]} ({k[0]}): não há XML na pasta')
            continue
        for c, v in x.items():
            if c.startswith('_'):
                continue
            campos += 1
            if abs((r[h.index(c)] or 0.0) - v) > 0.001:
                dif.append(f'NFS-e {k[1]} {c}: planilha {r[h.index(c)]} x XML {v}')
        campos += 1
        if (r[h.index('Retenção')] != 'Não') != x['_ret']:
            dif.append(f'NFS-e {k[1]}: marcação de retenção diferente do XML')
    for k in set(xml) - vistos:
        dif.append(f'XML da NFS-e {k[1]} ({k[0]}) não está na planilha')
    return n, campos, dif
