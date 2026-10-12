"""Lançamentos contábeis das entradas (NF-e, CT-e, NFS-e) no leiaute TXT do Domínio da skill rof-contabilidade-cosmetici:
    |0000|CNPJ|  /  |6000|X||||  /  |6100|DD/MM/AAAA|DÉBITO|CRÉDITO|VALOR||HISTÓRICO||||

As contas vêm do que o próprio Domínio já faz: o relatório "Lançamentos gerados pela Escrita" (Lançamentos.xls) cruzado nota a nota
com o Acompanhamento de Entradas deu a conta de cada acumulador (config `acumuladores_entrada[cfop].conta` e
`contas_servicos_por_acumulador`) e a conta de cada fornecedor (fornecedores.json `contas_dominio`). Valores exatamente do XML.
Documento sem conta comprovada NÃO entra no TXT: vai para a lista de pendências (nunca inventar conta)."""
import collections
import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

R = lambda x: round((x or 0.0) + 0.0, 2)

# histórico padrão do Domínio (Lançamentos gerados pela Escrita) por tipo de lançamento
HIST = {
    'compra': 'COMPRAS DE MERCADORIAS NESTA DATA', 'uso': 'COMPRAS PARA USO E CONSUMO', 'devolucao': 'DEVOLUÇÃO VENDA NESTA DATA',
    'industrializacao': 'ENTRADA DE MERC. PARA INDUSTRIALIZAÇÃO', 'amostra': 'AMOSTRA GRATIS', 'bonificacao': 'BONIFICAÇÕES RECEBIDAS',
    'remessa': 'REMESSA ', 'terceiros': 'INDUSTRIALIZAÇÃO EFETUADA POR TERCEIROS', 'frete': 'SERVIÇO DE TRANSPORTE ADQUIRIDO',
    'servico': 'SERVIÇO TOMADO NESTA DATA', 'recolher': 'VALOR A RECOLHER',
}
TIPO_ACUM = {'100': 'compra', '101': 'compra', '192': 'compra', '104': 'uso', '105': 'uso', '107': 'uso', '113': 'devolucao',
             '133': 'devolucao', '134': 'devolucao', '144': 'industrializacao', '149': 'industrializacao', '164': 'industrializacao',
             '165': 'industrializacao', '198': 'industrializacao', '147': 'amostra', '148': 'amostra', '110': 'bonificacao',
             '189': 'bonificacao', '155': 'remessa', '199': 'terceiros', '1000': 'frete', '1001': 'frete'}
TIPO_HIST = {v.strip(): k for k, v in HIST.items()}


def fmt_cnpj(c):
    c = re.sub(r'\D', '', c or '')
    return f'{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}' if len(c) == 14 else (f'{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}' if len(c) == 11 else c)


def historico(tipo, numero, nome, cnpj):
    nome = ' '.join(str(nome or '').upper().split())
    return f'{HIST[tipo]} {numero} {nome} {nome} {fmt_cnpj(cnpj)}'.replace('|', '/')


# ------------------------------------------------------------------ aprender com o Domínio

def ler_escrita(caminho):
    """Lançamentos.xls ("Lançamentos gerados pela Escrita") -> [{origem, num, data, deb, cred, valor, hist, cnpj, tipo}]."""
    import python_calamine as pc
    wb = pc.CalamineWorkbook.from_path(caminho)
    cod = lambda x: str(int(x)) if isinstance(x, float) else str(x).strip()
    out = []
    for r in wb.get_sheet_by_name(wb.sheet_names[0]).to_python():
        if not str(r[0]).strip().endswith(' EF'):
            continue
        h = str(r[27]).strip()
        m = re.search(r'(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}|\d{3}\.\d{3}\.\d{3}-\d{2})\s*$', h)
        t = re.match(r'[^\d]*', h).group(0).strip()
        out.append({'origem': str(r[0]).strip(), 'num': int(r[3]), 'data': str(r[7])[:10], 'deb': cod(r[12]), 'cred': cod(r[15]),
                    'valor': float(r[20] or 0), 'hist': h, 'cnpj': re.sub(r'\D', '', m.group(1)) if m else '',
                    'tipo': TIPO_HIST.get(t, t)})
    return out


def aprender(lancs, forn):
    """Grava em fornecedores.json, por CNPJ, a conta que o Domínio usou em cada tipo de lançamento: contas_dominio[tipo][conta] = n."""
    n = 0
    for l in lancs:
        if not l['cnpj'] or l['tipo'] in ('recolher',):
            continue
        conta = l['deb'] if l['deb'] not in ('0', '506', '504', '664') else None
        if not conta:
            continue
        nome = re.sub(r'^[^\d]*\d+\s', '', l['hist']).rsplit(' ', 1)[0].strip()
        meio = len(nome) // 2
        if nome[:meio].strip() == nome[meio:].strip():
            nome = nome[:meio].strip()
        f = forn.setdefault(l['cnpj'], {'nome': nome[:60]})
        c = f.setdefault('contas_dominio', {}).setdefault(l['tipo'], {})
        c[conta] = c.get(conta, 0) + 1
        n += 1
    return n


def _conta_forn(forn, cnpj, tipo):
    c = ((forn.get(cnpj) or {}).get('contas_dominio') or {}).get(tipo) or {}
    if not c:
        return None, ''
    conta, n = max(c.items(), key=lambda x: x[1])
    return conta, ('' if len(c) == 1 else f'fornecedor já usou {", ".join(f"{k} ({v}x)" for k, v in sorted(c.items()))} - conferir produto')


# ------------------------------------------------------------------ montar os lançamentos

def montar(a, cred, cte_xml, cte_sieg, nfse, forn):
    cfg = a.cfg
    acum = cfg.get('acumuladores_entrada', {})
    serv_ac = cfg.get('acumuladores_servicos', {})
    serv_conta = cfg.get('contas_servicos_por_acumulador', {})
    lanc, pend = [], []

    # NF-e: um lançamento por nota x CFOP de entrada (valor contábil dos itens do XML)
    grupos = collections.OrderedDict()
    for l in cred.linhas:
        n = l['nota']
        g = grupos.setdefault((n['chave'], l['cfop_ent']), {'n': n, 'vc': 0.0})
        g['vc'] += l['vc']
    for (ch, cf), g in grupos.items():
        n = g['n']
        ac = acum.get(cf) or {}
        base = {'doc': 'NF-e', 'numero': n['numero'], 'data': n['dhemi'], 'nome': n['emit_nome'], 'cnpj': n['emit_cnpj'], 'cfop': cf,
                'acumulador': ac.get('codigo', ''), 'valor': R(g['vc']), 'chave': ch}
        conta = ac.get('conta') if isinstance(ac.get('conta'), dict) else None
        tipo = TIPO_ACUM.get(ac.get('codigo', ''))
        if not conta or not tipo:
            pend.append({**base, 'motivo': f'CFOP {cf} sem acumulador/conta no histórico do Domínio'})
            continue
        deb, cr, obs = conta['debito'], conta['credito'], 'conta do acumulador'
        cf_forn, aviso = _conta_forn(forn, n['emit_cnpj'], tipo)
        if cf_forn and tipo in ('compra', 'uso'):
            deb, obs = cf_forn, 'conta do fornecedor no Domínio' + (f' ({aviso})' if aviso else '')
        elif tipo in ('compra', 'uso'):
            obs = f'fornecedor sem histórico: conta padrão do acumulador ({deb}) - conferir'
            if tipo == 'compra' and re.search(r'EMBALAG|ROTUL|ADESIV|ETIQUET|PLASTIC|PACK', (n['emit_nome'] or '').upper()):
                obs += '; embalagem/rótulo: no Domínio esses fornecedores vão em 600 INSUMOS'
        if tipo == 'devolucao' and 'COSMETICON' in (n['emit_nome'] or '').upper():
            cr = '664'
        elif '(' in cr:
            cr = cr.split()[0]
        lanc.append({**base, 'blocos': [('X', [(deb, cr, R(g['vc']), historico(tipo, n['numero'], n['emit_nome'], n['emit_cnpj']))])], 'obs': obs})

    # CT-e em que a Cosmetici é tomadora (XML; sem XML = relatório SIEG); cancelado não entra
    xs, rs = dict(cte_xml or {}), dict(cte_sieg or {})
    for ch in sorted(set(xs) | set(rs)):
        x, r = xs.get(ch), rs.get(ch)
        b = x or r
        if b['toma_cnpj'] != a.cnpj or (r and not r['autorizada']):
            continue
        cf = {'5': '1', '6': '2'}.get(b['cfop'][:1], b['cfop'][:1]) + b['cfop'][1:]
        ac = acum.get(cf) or {}
        conta = ac.get('conta') if isinstance(ac.get('conta'), dict) else None
        base = {'doc': 'CT-e', 'numero': b['numero'], 'data': b['data'], 'nome': b['emit_nome'], 'cnpj': b['emit_cnpj'], 'cfop': cf,
                'acumulador': ac.get('codigo', ''), 'valor': R(b['vprest']), 'chave': ch}
        if not conta:
            pend.append({**base, 'motivo': f'CFOP {cf} sem conta'})
            continue
        lanc.append({**base, 'blocos': [('X', [(conta['debito'], '506', R(b['vprest']), historico('frete', b['numero'], b['emit_nome'], b['emit_cnpj']))])],
                     'obs': 'CT-e' + ('' if x else ' sem XML na pasta (dados do relatório SIEG)')})

    # NFS-e tomadas: conta pelo acumulador do prestador; senão pelo histórico do prestador no Domínio
    for s in nfse or []:
        ac = serv_ac.get(s['prest_cnpj']) or ''
        if '/' in ac:
            ac = '819' if abs(s['valor'] - 83.99) < 0.01 else ('820' if abs(s['valor'] - 195.97) < 0.01 else ac)
        num = int(s['numero']) if str(s['numero']).isdigit() else s['numero']
        data = dt.date.fromisoformat(str(s['emissao'])[:10]) if s.get('emissao') else None
        base = {'doc': 'NFS-e', 'numero': num, 'data': data, 'nome': s['prest_nome'], 'cnpj': s['prest_cnpj'],
                'cfop': '1933' if (s.get('prest_uf') or 'PR') == 'PR' else '2933', 'acumulador': ac, 'valor': R(s['valor']),
                'chave': s.get('chave', '')}
        deb = (serv_conta.get(ac) or {}).get('debito')
        cr = (serv_conta.get(ac) or {}).get('credito', '506')
        if not deb:
            deb, _ = _conta_forn(forn, s['prest_cnpj'], 'servico')
            cr = '506'
        if cr in ('0', ''):   # no Domínio a NFS-e com ISS retido é multipartida (crédito "0"): o fornecedor é sempre 506
            cr = '506'
        if not deb:
            pend.append({**base, 'motivo': 'prestador sem acumulador e sem histórico no Domínio'})
            continue
        h = historico('servico', num, s['prest_nome'], s['prest_cnpj'])
        iss_ret, irrf, outras = R(s['iss_retido']), R(s['irrf']), R(s['csll']) + R(s['inss'])
        if outras:
            pend.append({**base, 'motivo': f'retenção de PIS/COFINS/CSLL ou INSS ({outras:,.2f}) sem conta no histórico do Domínio'})
            continue
        if iss_ret:   # como o Domínio faz: D despesa (total) / C fornecedor (líquido) / C 183 ISS retido - lote tipo D
            blocos = [('D', [(deb, '', R(s['valor']), h), ('', cr, R(s['valor'] - iss_ret), h), ('', '183', iss_ret, HIST['recolher'])])]
        else:
            blocos = [('X', [(deb, cr, R(s['valor']), h)])]
        if irrf:      # D fornecedor / C 178 IRRF a recolher
            blocos.append(('X', [('506', '178', irrf, HIST['recolher'])]))
        lanc.append({**base, 'blocos': blocos, 'obs': 'conta do acumulador' if ac else 'conta do prestador no Domínio'})
    key = lambda x: (x['data'] or dt.date.min, x['doc'], str(x['numero']))
    return sorted(lanc, key=key), sorted(pend, key=key)


def gerar_txt(cnpj, lanc, destino):
    out = [f'|0000|{re.sub(r"[^0-9]", "", cnpj)}|']
    for x in lanc:
        d = x['data'].strftime('%d/%m/%Y') if x['data'] else ''
        for tipo, linhas in x['blocos']:
            out.append(f'|6000|{tipo}||||')
            for deb, cr, v, h in linhas:
                out.append(f'|6100|{d}|{deb}|{cr}|{v:.2f}'.replace('.', ',') + f'||{h}||||')
    with open(destino, 'w', encoding='cp1252', errors='replace', newline='\r\n') as fh:
        fh.write('\n'.join(out) + '\n')
    return len(out)


def gerar_planilha(lanc, pend, destino, titulo):
    wb = Workbook()
    ws = wb.active; ws.title = 'Lançamentos'
    H = PatternFill('solid', fgColor='1F3864'); AM = PatternFill('solid', fgColor='FFE699'); VM = PatternFill('solid', fgColor='F8CBAD')
    F = lambda **k: Font(**{'name': 'Arial', 'size': 9, **k})
    def cab(ws, t, sub, cols, larg):
        ws['A1'] = t; ws['A1'].font = F(bold=True, size=12); ws['A2'] = sub; ws['A2'].font = F(italic=True, color='595959')
        for i, (c, w) in enumerate(zip(cols, larg), 1):
            x = ws.cell(4, i, c); x.font = F(bold=True, color='FFFFFF'); x.fill = H; ws.column_dimensions[L(i)].width = w
            x.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = 'A5'
    cab(ws, titulo, 'Rascunho para conferência - valores do XML; contas do histórico do Domínio. Amarelo = conferir a conta.',
        ['Data', 'Doc', 'Número', 'Fornecedor / prestador', 'CNPJ', 'CFOP', 'Acum.', 'Débito', 'Crédito', 'Valor', 'Histórico', 'Observação', 'Chave'],
        [10, 6, 9, 32, 16, 6, 6, 7, 7, 11, 60, 44, 46])
    r = 5
    for x in lanc:
        for deb, cr, v, h in [ln for _, b in x['blocos'] for ln in b]:
            vals = [x['data'], x['doc'], x['numero'], x['nome'], x['cnpj'], x['cfop'], x['acumulador'], deb, cr, v, h, x['obs'], x['chave']]
            for i, v2 in enumerate(vals, 1):
                c = ws.cell(r, i, v2); c.font = F()
                if i == 10: c.number_format = '#,##0.00'
                if i == 1: c.number_format = 'DD/MM/YYYY'
                if 'conferir' in x['obs']: c.fill = AM
            r += 1
    ws.cell(r, 9, 'Total débitos').font = F(bold=True)
    c = ws.cell(r, 10, f'=SUMPRODUCT((H5:H{r - 1}<>"")*J5:J{r - 1})'); c.font = F(bold=True); c.number_format = '#,##0.00'
    ws2 = wb.create_sheet('Pendências')
    cab(ws2, 'Documentos que NÃO entraram no TXT', 'Sem conta comprovada no histórico do Domínio - definir com a contadora.',
        ['Data', 'Doc', 'Número', 'Fornecedor / prestador', 'CNPJ', 'CFOP', 'Acum.', 'Valor', 'Motivo', 'Chave'],
        [10, 6, 9, 34, 16, 6, 6, 11, 60, 46])
    for j, x in enumerate(pend, 5):
        for i, v in enumerate([x['data'], x['doc'], x['numero'], x['nome'], x['cnpj'], x['cfop'], x['acumulador'], x['valor'], x['motivo'], x['chave']], 1):
            c = ws2.cell(j, i, v); c.font = F(); c.fill = VM
            if i == 8: c.number_format = '#,##0.00'
            if i == 1: c.number_format = 'DD/MM/YYYY'
    wb.save(destino)
