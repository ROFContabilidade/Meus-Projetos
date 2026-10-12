import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
def br(v): return f"{abs(v):.2f}".replace('.',',')
HDR=PatternFill('solid',fgColor='305496'); AM=PatternFill('solid',fgColor='FFF2CC')
def head(ws,cols):
    ws.append(cols)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=HDR
    ws.freeze_panes='A2'
def txt(path,linhas):
    L=['|0000|33039048000162|']
    for data,d,c,v,h in linhas: L+=['|6000|X||||',f"|6100|{data}|{d}|{c}|{br(v)}||{h.upper().replace('|','/')[:200]}||||"]
    open(path,'w',encoding='cp1252',errors='replace',newline='\n').write('\n'.join(L)+'\n')
# extratos jul/ago
wb=Workbook(); wb.remove(wb.active); res=[]
for mes,mm in [('Julho','07'),('Agosto','08')]:
    rows=json.load(open(f'class_{mes}.json')); D=C=0; lin=[]
    for z in rows:
        if z['valor']<0: lin.append((z['data'],z['cp'],'8',z['valor'],z['hist'])); C+=-z['valor']
        else: lin.append((z['data'],'8',z['cp'],z['valor'],z['hist'])); D+=z['valor']
    txt(f'out/COSMETICI_{mes.lower()}2026_dominio.txt',lin)
    ws=wb.create_sheet(f'{mm}-2026'); head(ws,['Data','Descrição extrato Itaú','Valor','D','C','Histórico no TXT','Origem da classificação','Relatório cliente (ERP)','Observação / pendência'])
    for z in rows:
        d,c=(z['cp'],'8') if z['valor']<0 else ('8',z['cp'])
        ws.append([z['data'],z['desc'],z['valor'],d,c,z['hist'],z['how'],z['erp'],z['nota']])
        if z['how'].startswith('padrao') or z['nota'] or 'varias' in z['how']:
            for cc in ws[ws.max_row]: cc.fill=AM
    for col,w in zip('ABCDEFGHI',[11,30,12,6,6,42,22,55,70]): ws.column_dimensions[col].width=w
    res.append((mes,round(D,2),round(C,2),len(rows)))
ws=wb.create_sheet('Conciliação',0); head(ws,['Mês','Saldo inicial (C/C + aplicação)','Entradas (D 8)','Saídas (C 8)','Saldo final calculado','Saldo final extrato','Diferença','Qtde'])
ini={'Julho':-1254.93,'Agosto':8386.26}; fim={'Julho':8386.26,'Agosto':-3506.09}
for mes,D,C,n in res:
    calc=round(ini[mes]+D-C,2); ws.append([mes,ini[mes],D,C,calc,fim[mes],round(calc-fim[mes],2),n]); print(mes,calc,fim[mes])
ws.append([]); ws.append(['Aplicação automática não segregada (conta 8 = saldo total). Rendimentos em 432.'])
wb.save('out/COSMETICI_conferencia_extratos_07-08_2026.xlsx')
