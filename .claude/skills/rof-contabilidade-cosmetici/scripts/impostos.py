import sys; sys.path.insert(0,'../tools')
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from saida import head, txt
M=['01','02','03','04','05','06','07','08']
fim={'01':'31/01/2026','02':'28/02/2026','03':'31/03/2026','04':'30/04/2026','05':'31/05/2026','06':'30/06/2026','07':'31/07/2026','08':'31/08/2026'}
# fiscal (SPED EFD ICMS/IPI + EFD Contrib + Resumo Domínio) x contábil (Razão impostos)
F={ # mes: icms_deb_proprio, aj_deb(PR000080+PR000081), cred_notas, aj_cred(PR020212), recolher, st, ipi_deb, ipi_cred, ipi_rec, pis, cofins
'01':(15251.29,671.24,15067.09,671.24,184.20,10579.61,16617.84,6005.65,10612.19,2282.30,10629.10),
'02':(27172.88,3593.34,16099.76,3593.34,11073.12,11168.68,7910.40,8431.61,0,2060.55,9594.18),
'03':(36945.22,3960.79,7659.80,3960.79,29285.42,18843.53,10733.50,1278.91,8933.38,2659.55,12399.63),
'04':(56723.63,2280.72,5734.43,2280.72,50989.20,32374.98,35673.16,685.05,34988.11,5165.57,24242.76),
'05':(52742.89,498.06,5170.29,439.92,47630.74,49720.89,37121.28,923.65,36197.63,7395.67,34544.55),
'06':(25233.67,339.01,3183.13,339.01,22050.54,40567.90,17005.53,302.18,16703.35,3706.58,17297.17),
'07':(24479.59,2455.15,8489.06,1914.78,16530.90,11319.49,4892.15,1285.04,3607.11,2975.63,13787.78),
'08':(40635.74,1186.91,8040.87,1186.91,32594.87,54079.89,22001.67,593.26,21408.41,5562.12,25935.13)}
C={ # contábil: 172 créd (débitos), 30 créd notas, 30 via 527, 526 créd, 481 créd, 171 créd, 171 déb(29), pis, cofins
'01':(15251.29,15067.09,671.24,671.24,10579.61,16617.84,6005.65,2282.30,10629.10),
'02':(27172.88,16099.76,3593.34,3593.34,11168.68,7910.40,7910.40,2060.55,9594.18),
'03':(36945.22,7659.80,3960.79,3960.79,18843.53,10733.50,1800.12,2659.55,12399.63),
'04':(56723.63,5734.43,2280.72,2280.72,32374.98,35673.16,685.05,7636.50,35647.06),
'05':(52742.89,6417.81,439.92,498.06,49720.89,37121.28,923.65,7395.67,34544.55),
'06':(25233.67,3337.77,339.01,339.01,40567.90,17005.53,302.18,3706.58,17297.17),
'07':(24479.59,8489.06,1914.78,2455.15,11319.49,4892.15,1285.04,2975.63,13787.78),
'08':(40635.74,8040.87,1186.91,1186.91,54079.89,22001.67,593.26,5562.12,25935.13)}
r2=lambda x:round(x,2)
wb=Workbook(); ws=wb.active; ws.title='RESUMO'
head(ws,['Mês','Imposto','Fiscal (SPED / Resumo Domínio)','Contábil (Razão)','Diferença (contábil - fiscal)','Situação','O que fazer'])
AM=PatternFill('solid',fgColor='FFF2CC'); VD=PatternFill('solid',fgColor='E2EFDA')
corr=[]
for m in M:
    f=F[m]; c=C[m]
    rows=[]
    deb_f=f[0]+f[1]; deb_c=c[0]
    rows.append(('ICMS débitos (172)',deb_f,deb_c,'Ajuste de débito E111 (antecipado PR000080 / DIFAL PR000081) lançado na 526 e não na 172' if r2(deb_c-deb_f) else '', 'D 526 / C 172 (reclassificação, ver aba CORRECOES)' if r2(deb_c-deb_f) else ''))
    cred_f=f[2]+f[3]; cred_c=c[1]+c[2]
    rows.append(('ICMS créditos (30→172)',cred_f,cred_c,'Crédito de entradas no contábil maior que o do fiscal (notas alteradas/excluídas no fiscal após a integração)' if r2(cred_c-cred_f) else '','Refazer a integração das ENTRADAS do mês (fiscal→contábil) no Domínio' if r2(cred_c-cred_f) else ''))
    rows.append(('ICMS saldo a recolher (E116 cód.1015)',f[4],c[0]-c[1]-c[2],'',''))
    rows.append(('ICMS antecipado/DIFAL (526)',0.0,c[3],'No fiscal esse valor já está dentro da guia 1015 (débito e crédito na própria apuração); não há guia separada','D 526 / C 172'))
    rows.append(('ICMS ST (481)',f[5],c[4],'',''))
    rows.append(('IPI débitos (171)',f[6],c[5],'',''))
    rows.append(('IPI créditos (29→171)',f[7]+(521.21 if m=='03' else 0),c[6],'Inclui saldo credor de 02/2026 (521,21) usado em 03' if m=='03' else ('Crédito 8.431,61: 7.910,40 compensado e 521,21 ficou de saldo credor na 29' if m=='02' else ''),''))
    rows.append(('IPI a recolher',f[8],c[5]-c[6],'',''))
    rows.append(('PIS (179)',f[9],c[7],'Contábil com o valor da EFD original; a retificadora/fiscal está com 5.165,57' if r2(c[7]-f[9]) else '','D 179 / C 429' if r2(c[7]-f[9]) else ''))
    rows.append(('COFINS (180)',f[10],c[8],'Contábil com o valor da EFD original; a retificadora/fiscal está com 24.242,76' if r2(c[8]-f[10]) else '','D 180 / C 428' if r2(c[8]-f[10]) else ''))
    for nome,fv,cv,sit,acao in rows:
        d=r2(cv-fv)
        if nome.startswith('ICMS saldo') and d: sit='Diferença = soma das linhas acima'; acao='Some ao corrigir as linhas acima'
        if nome.startswith('ICMS antecipado') and not cv: continue
        ws.append([f'{m}/2026',nome,r2(fv),r2(cv),d,sit or 'OK',acao])
        for cc in ws[ws.max_row]: cc.fill=AM if d else VD
for q,(irf,irc,csf,csc) in {'1º TRI (31/03)':(7925.10,7925.10,7132.59,7132.59),'2º TRI (30/06)':(13972.85,13972.85,10785.34,10785.34)}.items():
    ws.append([q,'IRPJ (176) – base 52.834,02 / 79.891,39 + adicional',irf,irc,0,'OK','']); ws.append([q,'CSLL (177)',csf,csc,0,'OK',''])
ws.append(['3º TRI','IRPJ/CSLL','','','','Fecha em 30/09 – ainda não apurado','Lançar na apuração de setembro'])
for m,v in [('01',182.88),('05',55.50),('08',286.55)]: ws.append([f'{m}/2026','ISS retido (183)',v,v,0,'OK',''])
for m,v in [('04',100.18),('05',10.71),('08',36.70)]: ws.append([f'{m}/2026','IRRF (178)',v,v,0,'OK',''])
ws.append(['03/2026','IRRF (178)',52.93,-52.93,-105.86,'Fiscal tem IRRF 52,93 a recolher; no contábil o lote 22449 (04/03) DEBITOU a 178 contra a 31 IRRF a recuperar','Verificar a nota: se for retenção sobre serviço tomado → D 31 / C 178 105,86 (estorna e lança). Não incluído no TXT'])
for col,w in zip('ABCDEFG',[12,42,18,18,16,70,60]): ws.column_dimensions[col].width=w
# correções
ws2=wb.create_sheet('CORRECOES'); head(ws2,['Data','Débito','Crédito','Valor','Histórico','Motivo'])
for m in M:
    v=F[m][1]; corr.append((fim[m],'526','172',v,f'RECLASSIF. ICMS ANTECIPADO/DIFAL APURADO NA GUIA 1015 COMP. {m}/2026','Ajuste de débito E111 está dentro da apuração do ICMS (guia 1015)'))
corr.append(('30/04/2026','179','429',2470.93,'ESTORNO PIS A MAIOR CONF. EFD CONTRIBUICOES RETIFICADORA 04/2026','Contábil 7.636,50 x fiscal 5.165,57'))
corr.append(('30/04/2026','180','428',11404.30,'ESTORNO COFINS A MAIOR CONF. EFD CONTRIBUICOES RETIFICADORA 04/2026','Contábil 35.647,06 x fiscal 24.242,76'))
for x in corr: ws2.append([x[0],x[1],x[2],r2(x[3]),x[4],x[5]])
ws2.append([]); ws2.append(['PENDENTE (fora do TXT)','','',1247.52,'05/2026 crédito ICMS entradas a maior (contábil 6.417,81 x fiscal 5.170,29)','Refazer integração das entradas de maio no Domínio'])
ws2.append(['PENDENTE (fora do TXT)','','',154.64,'06/2026 crédito ICMS entradas a maior (contábil 3.337,77 x fiscal 3.183,13)','Refazer integração das entradas de junho no Domínio'])
for col,w in zip('ABCDEF',[22,8,8,12,75,60]): ws2.column_dimensions[col].width=w
ws3=wb.create_sheet('OBSERVACOES')
for t in ['IPI, ICMS-ST, IRPJ, CSLL, ISS retido e IRRF (exceto 03/2026) do contábil batem com o fiscal em todos os meses.',
 'PIS/COFINS batem em todos os meses, exceto 04/2026 (contábil ficou com a EFD-Contribuições original).',
 'O arquivo da EFD-Contribuições que está na pasta de 07/2026 tem os valores de agosto (PIS 5.562,12 / COFINS 25.935,13). O fiscal (Resumo Domínio) e o contábil estão com 2.975,63 / 13.787,78. Conferir se a EFD de julho transmitida é a correta.',
 'ICMS: o saldo da 172 fica menor que a guia porque o ajuste de débito (antecipado PR000080 e DIFAL PR000081) foi para a 526, enquanto o crédito PR020212 abateu a 172. Em jan, por ex., a 172 ficou devedora (487,04D) quando a guia é 184,20. O total 172+526 não muda com a correção, só a conta.',
 'A 526 já vinha com saldo anterior de 72.983,74C sem baixas: provavelmente a mesma situação nos anos anteriores. Vale revisar.',
 'Nenhuma conta de imposto tem baixa por pagamento em 2026 (só 3 SISPAG de ST em 04/2026, 1.723,36). Os saldos estão acumulando: 172 193.951,61 | 171 132.450,18 | 481 383.924,19 | PIS 34.278,90 | COFINS 159.834,60 | IRPJ 21.897,95 | CSLL 17.917,93.',
 'IBS/CBS 2026 (fase de teste): só informativo, não precisa lançar no contábil.']:
    ws3.append([t])
ws3.column_dimensions['A'].width=160
wb.save('out/COSMETICI_conferencia_impostos_fiscal_x_contabil_jan-ago2026.xlsx')
txt('out/COSMETICI_correcao_impostos_jan-ago2026.txt',[(a,b,c,v,h) for a,b,c,v,h,_ in corr])
print(sum(F[m][1] for m in M))
