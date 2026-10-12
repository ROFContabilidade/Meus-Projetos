"""Relação das notas de entrada do mês (Relacao_Entradas_Cosmetici_MM-AAAA.xlsx): todas juntas e depois separadas em
NF-e dentro do estado, NF-e fora do estado, CT-e (dentro e fora) e NFS-e tomadas com as retenções destacadas.
Valores exatamente como estão no XML (sem correção). CFOP de entrada e acumulador são sugestão de escrituração."""
import collections

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from creditos import Credito

F = lambda **k: Font(**{'name': 'Arial', 'size': 9, **k})
B = Border(bottom=Side(style='thin', color='D9D9D9'))
HEAD = PatternFill('solid', fgColor='1F3864')
CANC = PatternFill('solid', fgColor='F8CBAD')
RET = PatternFill('solid', fgColor='FFE699')
PEND = PatternFill('solid', fgColor='EDEDED')
TOT = PatternFill('solid', fgColor='DDEBF7')
MOEDA = '#,##0.00;-#,##0.00;-'
TP_RET_PC = {'0': 'PIS/COFINS/CSLL não retidos', '1': 'PIS/COFINS retidos', '2': 'PIS/COFINS não retidos',
             '3': 'PIS/COFINS/CSLL retidos', '4': 'PIS/COFINS retidos, CSLL não retida', '5': 'PIS retido',
             '6': 'COFINS retido', '7': 'PIS não retido; COFINS e CSLL retidos', '8': 'PIS/COFINS não retidos; CSLL retida',
             '9': 'COFINS não retido; PIS e CSLL retidos'}


def _acum(cfg, cfops):
    acum = cfg.get('acumuladores_entrada', {})
    out = []
    for cf in cfops:
        ac = acum.get(cf)
        t = f'{ac["codigo"]} - {ac["descricao"][:26]}' if ac else f'definir (CFOP {cf})'
        if t not in out:
            out.append(t)
    return ' / '.join(out)


def nfe(a, cred):
    cf_ent = {}
    for l in cred.linhas:
        cf_ent.setdefault(l['nota']['chave'], []).append(l['cfop_ent'])
    out = []
    docs = [(n, 'Autorizada') for n in a.entradas] + [(n, 'Cancelada') for n in a.canceladas if not n.get('propria')]
    for n, sit in docs:
        t = n['tot']
        ents = sorted(set(cf_ent.get(n['chave'], [])))
        out.append({
            'tipo': 'NF-e', 'data': n['dhemi'], 'numero': n['numero'], 'serie': n['serie'], 'emitente': n['emit_nome'],
            'cnpj': n['emit_cnpj'], 'uf': n.get('emit_uf') or '', 'situacao': sit,
            'cfop': '/'.join(sorted({i['cfop'] for i in n['itens']})), 'cfop_ent': '/'.join(ents),
            'acumulador': _acum(a.cfg, ents) if sit == 'Autorizada' else 'cancelada (situação 02)',
            **{k: t.get(k, 0.0) for k in ('vProd', 'vDesc', 'vFrete', 'vOutro', 'vBC', 'vICMS', 'vBCST', 'vST', 'vIPI', 'vPIS',
                                         'vCOFINS', 'vNF')},
            'chave': n['chave']})
    return sorted(out, key=lambda r: (r['data'], r['numero']))


def cte(a, xml_cte, sieg_cte):
    """CT-e em que a Cosmetici é tomadora: XML da pasta; os que não têm XML vêm do relatório SIEG (marcados)."""
    out = []
    xs, rs = dict(xml_cte or {}), dict(sieg_cte or {})
    for ch in set(xs) | set(rs):
        x, r = xs.get(ch), rs.get(ch)
        b = x or r
        if b['toma_cnpj'] != a.cnpj:
            continue
        canc = bool(r and not r['autorizada'])
        out.append({
            'tipo': 'CT-e', 'data': b['data'], 'numero': b['numero'], 'serie': b['serie'], 'emitente': b['emit_nome'],
            'cnpj': b['emit_cnpj'], 'uf': f'{b.get("uf_ini") or ""}-{b.get("uf_fim") or ""}', 'cfop': b['cfop'],
            'situacao': 'Cancelado' if canc else 'Autorizado', 'fonte': 'XML' if x else 'Relatório SIEG (sem XML na pasta)',
            'vNF': b['vprest'], 'vBC': b['vbc'], 'picms': b.get('picms'), 'vICMS': b['vicms'], 'cst': b.get('cst') or '',
            'nfe': ', '.join(f'{int(k[25:34])}' for k in b['chaves_nfe'] if len(k) == 44), 'chave': ch})
    return sorted(out, key=lambda r: (r['data'], r['numero']))


def nfse(cfg, xml_nfse):
    serv, por_nome = cfg.get('acumuladores_servicos', {}), cfg.get('acumuladores_servicos_por_nome', {})
    out = []
    for x in xml_nfse or []:
        ac = serv.get(x['prest_cnpj']) or ''
        if '/' in ac:
            ac = next((v for k, v in por_nome.items() if k.split()[0] in (x['prest_nome'] or '').upper()
                       and k.split()[-1] == f'{x["valor"]:.2f}'.replace('.', ',')), ac)
        ret = {'iss_retido': x['iss_retido'], 'irrf': x['irrf'], 'csll': x['csll'], 'inss': x['inss']}
        tem = any(v for v in ret.values()) or bool(x['total_ret'])
        quais = [rot for rot, k in (('ISS', 'iss_retido'), ('IRRF', 'irrf'), ('INSS', 'inss')) if ret[k]]
        if x['csll']:
            quais.append('PIS/COFINS/CSLL' if x.get('tp_ret_piscofins') == '3' else 'CSLL')
        out.append({
            'tipo': 'NFS-e', 'data': x['emissao'], 'numero': int(x['numero']) if str(x['numero']).isdigit() else x['numero'], 'serie': '',
            'emitente': x['prest_nome'], 'cnpj': x['prest_cnpj'], 'uf': x.get('prest_uf') or '', 'municipio': x.get('mun_incidencia') or '',
            'simples': 'Sim' if x['simples'] else 'Não', 'cod': x.get('cod_trib_nac') or '', 'servico': (x.get('servico') or '')[:60],
            'situacao': 'Autorizada' if x['cstat'] == '100' else f'Autorizada (cStat {x["cstat"]})',
            'vNF': x['valor'], 'base_iss': x['base_iss'], 'aliq_iss': x['aliq_iss'], 'iss': x['iss'], **ret,
            'pis': x['pis'], 'cofins': x['cofins'], 'tp_pc': TP_RET_PC.get(x.get('tp_ret_piscofins') or '', '-'),
            'total_ret': x['total_ret'], 'liquido': x['liquido'], 'tem_ret': 'SIM - ' + ', '.join(quais or ['ver total']) if tem else 'Não',
            'acumulador': ac or 'definir', 'chave': x.get('chave', '')})
    return sorted(out, key=lambda r: (str(r['data']), str(r['numero'])))


def _aba(wb, titulo, sub, cols, linhas, fill=None, primeira=False, contar=None):
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = titulo[:31]
    ws['A1'] = titulo; ws['A1'].font = F(bold=True, size=12)
    ws['A2'] = sub; ws['A2'].font = F(italic=True, color='595959')
    hr = 4
    for c, (tit, larg, _, _) in enumerate(cols, 1):
        x = ws.cell(hr, c, tit); x.font = F(bold=True, color='FFFFFF'); x.fill = HEAD
        x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[L(c)].width = larg
    ws.row_dimensions[hr].height = 30
    r = hr + 1
    for ln in linhas:
        fl = fill(ln) if fill else None
        for c, (_, _, k, tp) in enumerate(cols, 1):
            v = ln.get(k)
            x = ws.cell(r, c, v if v is not None else ('' if tp != 'm' else None)); x.font = F(); x.border = B
            if tp == 'm': x.number_format = MOEDA
            elif tp == 'd': x.number_format = 'DD/MM/YYYY'
            if fl: x.fill = fl
        r += 1
    ult = r - 1
    sit = next((c for c, col in enumerate(cols, 1) if col[2] == 'situacao'), None)
    ws.cell(r, 1, 'Total (sem canceladas)').font = F(bold=True)
    for c, (_, _, k, tp) in enumerate(cols, 1):
        ws.cell(r, c).fill = TOT
        if tp == 'm' and ult > hr:
            col, s = L(c), L(sit)
            f = (f'=SUMPRODUCT(({s}{hr + 1}:{s}{ult}<>"Cancelada")*({s}{hr + 1}:{s}{ult}<>"Cancelado")*{col}{hr + 1}:{col}{ult})'
                 if sit else f'=SUM({col}{hr + 1}:{col}{ult})')
            x = ws.cell(r, c, f); x.font = F(bold=True); x.number_format = MOEDA
    ws.cell(r, 3, f'=COUNTA(C{hr + 1}:C{ult})&" docs"').font = F(bold=True)
    if contar:
        rr = r + 2
        for rot, f in contar(hr + 1, ult):
            ws.cell(rr, 1, rot).font = F(bold=True); ws.cell(rr, 4, f).font = F(bold=True); rr += 1
    ws.freeze_panes = ws.cell(hr + 1, 4)
    ws.auto_filter.ref = f'A{hr}:{L(len(cols))}{max(ult, hr)}'
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1
    return ws


def _cor_sit(ln):
    if ln.get('situacao') in ('Cancelada', 'Cancelado'):
        return CANC
    if ln.get('tipo') == 'NFS-e' and str(ln.get('tem_ret', '')).startswith('SIM'):
        return RET
    if ln.get('fonte', 'XML') != 'XML':
        return PEND
    return None


COLS_NFE = [('Emissão', 10, 'data', 'd'), ('NF', 8, 'numero', ''), ('Série', 5, 'serie', ''), ('Fornecedor', 32, 'emitente', ''),
            ('CNPJ', 15, 'cnpj', ''), ('UF', 4, 'uf', ''), ('Situação', 10, 'situacao', ''), ('CFOP (XML)', 9, 'cfop', ''),
            ('CFOP entrada (sugestão)', 10, 'cfop_ent', ''), ('Produtos', 11, 'vProd', 'm'), ('Desconto', 9, 'vDesc', 'm'),
            ('Frete', 9, 'vFrete', 'm'), ('Outras desp.', 9, 'vOutro', 'm'), ('Base ICMS', 11, 'vBC', 'm'), ('ICMS', 10, 'vICMS', 'm'),
            ('Base ICMS-ST', 10, 'vBCST', 'm'), ('ICMS-ST', 9, 'vST', 'm'), ('IPI', 9, 'vIPI', 'm'), ('PIS', 8, 'vPIS', 'm'),
            ('COFINS', 9, 'vCOFINS', 'm'), ('Total da nota', 12, 'vNF', 'm'), ('Acumulador (sugestão)', 26, 'acumulador', ''),
            ('Natureza', 14, 'natureza', ''), ('Crédito ICMS', 10, 'cred_icms', 'm'), ('Crédito IPI', 9, 'cred_ipi', 'm'),
            ('DIFAL', 8, 'difal', 'm'), ('ICMS antecipado', 9, 'antecipado', 'm'), ('Conta contábil (Domínio)', 14, 'conta', ''),
            ('Conferência XML x relatório SIEG', 26, 'conferencia', ''), ('Análise', 50, 'analise', ''), ('Chave de acesso', 46, 'chave', '')]
COLS_CTE = [('Emissão', 10, 'data', 'd'), ('CT-e', 9, 'numero', ''), ('Série', 5, 'serie', ''), ('Transportadora', 32, 'emitente', ''),
            ('CNPJ', 15, 'cnpj', ''), ('Trajeto (UF)', 8, 'uf', ''), ('Situação', 10, 'situacao', ''), ('CFOP', 6, 'cfop', ''),
            ('CST', 5, 'cst', ''), ('Valor do frete', 11, 'vNF', 'm'), ('Base ICMS', 11, 'vBC', 'm'), ('% ICMS', 6, 'picms', ''),
            ('ICMS', 10, 'vICMS', 'm'), ('NF-e transportada', 22, 'nfe', ''), ('Fonte', 26, 'fonte', ''),
            ('Conta contábil (Domínio)', 14, 'conta', ''), ('Conferência XML x relatório SIEG', 30, 'conferencia', ''), ('Chave de acesso', 46, 'chave', '')]
COLS_NFSE = [('Emissão', 10, 'data', 'd'), ('NFS-e', 9, 'numero', ''), ('Prestador', 30, 'emitente', ''), ('CNPJ', 15, 'cnpj', ''),
             ('Situação', 10, 'situacao', ''), ('Simples', 7, 'simples', ''), ('Município incidência', 14, 'municipio', ''),
             ('Cód. serviço', 8, 'cod', ''), ('Serviço', 26, 'servico', ''), ('Valor', 11, 'vNF', 'm'), ('Base ISS', 10, 'base_iss', 'm'),
             ('% ISS', 5, 'aliq_iss', ''), ('ISS', 9, 'iss', 'm'), ('ISS retido', 9, 'iss_retido', 'm'), ('IRRF retido', 9, 'irrf', 'm'),
             ('PIS/COFINS/CSLL retidos (vRetCSLL)', 11, 'csll', 'm'), ('INSS retido', 9, 'inss', 'm'),
             ('Total retido (vTotalRet)', 10, 'total_ret', 'm'), ('Líquido', 11, 'liquido', 'm'), ('Tem retenção?', 18, 'tem_ret', ''),
             ('PIS (nota)', 8, 'pis', 'm'), ('COFINS (nota)', 8, 'cofins', 'm'), ('Retenção PIS/COFINS (XML)', 22, 'tp_pc', ''),
             ('Acumulador', 10, 'acumulador', ''), ('Conta contábil (Domínio)', 14, 'conta', ''),
             ('Conferência XML x relatório ABRASF', 30, 'conferencia', ''), ('Chave', 50, 'chave', '')]
COLS_TODAS = [('Emissão', 10, 'data', 'd'), ('Tipo', 6, 'tipo', ''), ('Número', 9, 'numero', ''), ('Emitente', 32, 'emitente', ''),
              ('CNPJ', 15, 'cnpj', ''), ('UF', 6, 'uf', ''), ('Situação', 10, 'situacao', ''), ('CFOP', 9, 'cfop', ''),
              ('Valor total', 12, 'vNF', 'm'), ('Base ICMS', 11, 'vBC', 'm'), ('ICMS', 10, 'vICMS', 'm'), ('ICMS-ST', 9, 'vST', 'm'),
              ('IPI', 9, 'vIPI', 'm'), ('ISS', 9, 'iss', 'm'), ('Retenções (NFS-e)', 10, 'total_ret', 'm'),
              ('Conta contábil (Domínio)', 14, 'conta', ''), ('Conferência', 26, 'conferencia', ''),
              ('Observação', 30, 'obs', ''), ('Chave', 46, 'chave', '')]


def _analise(a, cred, forn, xml_cte, sieg_cte, xml_nfse, sieg_prod, rel_nfse, l_nfe, l_cte, l_nfse):
    """Acrescenta a cada documento a análise feita XML por XML: natureza e créditos, conta contábil do Domínio e a conferência
    com os relatórios SIEG / ABRASF da pasta."""
    import cruzamento_entradas as ce
    import lancamentos_contabeis as lc
    lanc, pend = lc.montar(a, cred, xml_cte, sieg_cte, xml_nfse, forn)
    conta = collections.defaultdict(list)
    for x in lanc:
        d = next(ln for _, b in x['blocos'] for ln in b if ln[0])
        conta[(x['doc'], str(x['numero']), x['cnpj'])].append(f'D {d[0]} / C {d[1] or "506"}' + (' (conferir)' if 'conferir' in x['obs'] else ''))
    for x in pend:
        conta[(x['doc'], str(x['numero']), x['cnpj'])].append('sem conta: ' + x['motivo'])
    nat = collections.defaultdict(lambda: {'nat': set(), 'icms': 0.0, 'ipi': 0.0, 'difal': 0.0, 'ant': 0.0, 'mot': []})
    for l in cred.linhas:
        g = nat[l['nota']['chave']]
        g['nat'].add(l['natureza']); g['icms'] += l['cred_icms'] + l['cred_sn']; g['ipi'] += l['cred_ipi']
        g['difal'] += l['difal']; g['ant'] += l.get('antecipado', 0.0)
        for m in (l.get('motivo'), l.get('acao')):
            if m and m not in g['mot']:
                g['mot'].append(m)
    conf_nfe = {ln[19]: ln[18] for ln in ce.nfe(a, sieg_prod)} if sieg_prod else {}
    conf_cte = {ln[16]: ln[15] for ln in ce.cte(a, xml_cte, sieg_cte)}
    conf_nfse = {(ln[3], str(ln[0])): ln[24] for ln in ce.nfse(xml_nfse, rel_nfse)} if rel_nfse else {}
    for r in l_nfe:
        g = nat.get(r['chave'])
        if g:
            r.update(natureza=', '.join(sorted(g['nat'])), cred_icms=round(g['icms'], 2), cred_ipi=round(g['ipi'], 2),
                     difal=round(g['difal'], 2), antecipado=round(g['ant'], 2), analise='; '.join(g['mot'])[:250])
        r['conta'] = ' | '.join(conta.get(('NF-e', str(r['numero']), r['cnpj']), ['cancelada - sem lançamento'] if r['situacao'] == 'Cancelada' else []))
        r['conferencia'] = conf_nfe.get(r['chave'], 'sem relatório SIEG' if not sieg_prod else 'não está no relatório')
    for r in l_cte:
        r['conta'] = ' | '.join(conta.get(('CT-e', str(r['numero']), r['cnpj']), ['cancelado - sem lançamento'] if r['situacao'] == 'Cancelado' else []))
        r['conferencia'] = conf_cte.get(r['chave'], '')
    for r in l_nfse:
        r['conta'] = ' | '.join(conta.get(('NFS-e', str(r['numero']), r['cnpj']), []))
        r['conferencia'] = conf_nfse.get((r['cnpj'], str(r['numero'])), 'sem relatório ABRASF' if not rel_nfse else 'não está no relatório')


def gerar(a, forn, xml_cte, sieg_cte, xml_nfse, destino, sieg_prod=None, rel_nfse=None):
    cred = Credito(a, forn, [], a.cfg.get('natureza_entradas', {})).analisar()
    l_nfe, l_cte, l_nfse = nfe(a, cred), cte(a, xml_cte, sieg_cte), nfse(a.cfg, xml_nfse)
    _analise(a, cred, forn, xml_cte, sieg_cte, xml_nfse, sieg_prod, rel_nfse, l_nfe, l_cte, l_nfse)
    uf = a.cfg['empresa'].get('uf', 'PR')
    dentro = [r for r in l_nfe if r['uf'] == uf]
    fora = [r for r in l_nfe if r['uf'] != uf]
    for r in l_cte:
        r['obs'] = 'XML não está na pasta - dados do relatório SIEG' if r['fonte'] != 'XML' else ''
    for r in l_nfse:
        r['obs'] = r['tem_ret'] if r['tem_ret'] != 'Não' else ''
    todas = sorted(l_nfe + l_cte + l_nfse, key=lambda r: (str(r['data']), r['tipo'], str(r['numero'])))
    emp = a.cfg['empresa']
    mm = f'{a.mes[5:]}/{a.mes[:4]}'
    wb = Workbook()
    _aba(wb, 'Todas as entradas', f'{emp["nome"]} - {mm} - NF-e, CT-e e NFS-e na ordem da emissão. Valores como estão no XML. '
         'Vermelho = cancelada · amarelo = NFS-e com retenção · cinza = CT-e sem XML na pasta (dados do relatório SIEG).',
         COLS_TODAS, todas, _cor_sit, primeira=True,
         contar=lambda i, u: [(f'{t}', f'=COUNTIF(B{i}:B{u},"{t}")') for t in ('NF-e', 'CT-e', 'NFS-e')])
    _aba(wb, f'NF-e dentro do estado ({uf})', f'NF-e de entrada de fornecedores do {uf} - {mm}. Valores do XML; CFOP de entrada e acumulador = sugestão.',
         COLS_NFE, dentro, _cor_sit)
    _aba(wb, 'NF-e fora do estado', f'NF-e de entrada de fornecedores de outras UF - {mm}. Valores do XML; CFOP de entrada e acumulador = sugestão.',
         COLS_NFE, fora, _cor_sit,
         contar=lambda i, u: [(f'UF {x}', f'=COUNTIF(F{i}:F{u},"{x}")') for x in sorted({r["uf"] for r in fora})])
    _aba(wb, 'CT-e', f'CT-e em que a Cosmetici é tomadora, dentro e fora do estado - {mm}. Cinza = sem XML na pasta (dados do relatório SIEG); '
         'vermelho = cancelado.', COLS_CTE, l_cte, _cor_sit,
         contar=lambda i, u: [('Dentro do estado', f'=COUNTIF(F{i}:F{u},"{uf}-{uf}")'),
                              ('Com XML', f'=COUNTIF(O{i}:O{u},"XML")'), ('Sem XML, autorizados (pedir o XML)', f'=COUNTIFS(O{i}:O{u},"Relatório*",G{i}:G{u},"Autorizado")'),
                              ('Cancelados', f'=COUNTIF(G{i}:G{u},"Cancelado")')])
    _aba(wb, 'NFS-e tomadas', f'Serviços tomados - {mm}. Retenções exatamente como no XML (sem recálculo). Amarelo = nota com retenção.',
         COLS_NFSE, l_nfse, _cor_sit,
         contar=lambda i, u: [('Notas com retenção', f'=COUNTIF(T{i}:T{u},"SIM*")'), ('Notas sem retenção', f'=COUNTIF(T{i}:T{u},"Não")')])
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    return {'todas': todas, 'dentro': dentro, 'fora': fora, 'cte': l_cte, 'nfse': l_nfse}
