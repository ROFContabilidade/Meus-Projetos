"""Relação das notas de saída do mês (Relacao_Saidas_Cosmetici_MM-AAAA.xlsx): uma linha por nota, na ordem da data de saída
e do número, com tudo o que está destacado no XML (sem alterar nada), as canceladas e as inutilizadas destacadas e a sugestão
de importação para o Domínio."""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = lambda **k: Font(**{'name': 'Arial', 'size': 9, **k})
thin = Side(style='thin', color='D9D9D9')
B = Border(bottom=thin)
HEAD = PatternFill('solid', fgColor='1F3864')
CANC = PatternFill('solid', fgColor='F8CBAD')
INUT = PatternFill('solid', fgColor='D9D9D9')
TOT = PatternFill('solid', fgColor='DDEBF7')
MOEDA = '#,##0.00;-#,##0.00;-'

COLS = [  # (título, largura, chave, moeda)
    ('Data saída', 10, 'saida', False), ('Emissão', 10, 'emissao', False), ('NF', 7, 'numero', False), ('Série', 5, 'serie', False),
    ('Situação', 11, 'situacao', False), ('Cliente', 34, 'cliente', False), ('CNPJ/CPF', 16, 'doc', False), ('UF', 4, 'uf', False),
    ('CFOP', 10, 'cfop', False), ('Natureza da operação', 26, 'natop', False),
    ('Produtos', 11, 'vProd', True), ('Desconto', 9, 'vDesc', True), ('Frete', 9, 'vFrete', True), ('Outras desp.', 9, 'vOutro', True),
    ('Base ICMS', 11, 'vBC', True), ('ICMS', 10, 'vICMS', True), ('Base ICMS-ST', 11, 'vBCST', True), ('ICMS-ST', 9, 'vST', True),
    ('FCP-ST', 8, 'vFCPST', True), ('IPI', 9, 'vIPI', True), ('PIS', 9, 'vPIS', True), ('COFINS', 9, 'vCOFINS', True),
    ('DIFAL UF destino', 9, 'vICMSUFDest', True), ('Total da nota', 12, 'vNF', True),
    ('Acumulador Domínio', 22, 'acumulador', False), ('Importar no Domínio', 30, 'importar', False),
    ('Conferência XML x Faturamento x Registro (ERP)', 30, 'conferencia', False), ('Análise (informativa - segue o XML)', 60, 'analise', False),
    ('Chave de acesso', 46, 'chave', False),
]


def _acumulador(a, cfops):
    cods = []
    for cf in cfops:
        cod, ac = a.acum_por_cfop.get(cf, (None, None))
        txt = f'{cod} - {ac["descricao"][:28]}' if cod else f'sem acumulador (CFOP {cf})'
        if txt not in cods:
            cods.append(txt)
    return ' / '.join(cods)


def _analise(a):
    """Pontos da revisão XML por XML (revisao.revisar_saidas), por número de nota. Só informativo."""
    import re
    import revisao
    div, _ = revisao.revisar_saidas(a)
    por = {}
    for d in div:
        m = re.search(r'NF (\d+)', str(d[2]))
        if m:
            por.setdefault(int(m.group(1)), []).append(f'[{d[0]}] {d[1]}: {d[3]}')
    return por


def conferencia_erp(caminho):
    """Resultado por NF da planilha Cruzamento_Relatorios_ERP_x_XML (aba "Cruzamento XML x ERP", última coluna)."""
    import os
    from openpyxl import load_workbook
    if not caminho or not os.path.exists(caminho):
        return {}
    ws = load_workbook(caminho, read_only=True)['Cruzamento XML x ERP']
    out = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if isinstance(r[0], (int, float)):
            out[int(r[0])] = r[-1]
    return out


def linhas(a, conf=None):
    conf = conf or {}
    analise = _analise(a)
    out = []
    for n in a.saidas + [x for x in a.canceladas if x.get('propria')]:
        canc = n in a.canceladas
        cfops = sorted({i['cfop'] for i in n['itens']})
        t = n['tot']
        r = {'saida': n['dhsaient'] or n['dhemi'], 'emissao': n['dhemi'], 'numero': n['numero'], 'serie': n['serie'],
             'situacao': 'Cancelada' if canc else 'Autorizada', 'cliente': n['dest_nome'], 'doc': n['dest_cnpj'], 'uf': n['dest_uf'],
             'cfop': '/'.join(cfops), 'natop': n.get('natop') or '', 'chave': n['chave']}
        r.update({k: t.get(k, 0.0) for k in ('vProd', 'vDesc', 'vFrete', 'vOutro', 'vBC', 'vICMS', 'vBCST', 'vST', 'vFCPST', 'vIPI',
                                             'vPIS', 'vCOFINS', 'vICMSUFDest', 'vNF')})
        r['acumulador'] = _acumulador(a, cfops)
        r['importar'] = ('Sim - como CANCELADA (situação 02: valores zerados, só a chave)' if canc
                         else 'Sim - como está no XML')
        r['conferencia'] = conf.get(int(n['numero']), 'sem relatório do ERP' if not conf else 'não está no relatório')
        r['analise'] = ' | '.join(analise.get(int(n['numero']), []))[:400]
        out.append(r)
    for x in a.f.get('inut', []):
        for num in range(x['ini'], x['fim'] + 1):
            out.append({'saida': x.get('data'), 'emissao': x.get('data'), 'numero': num, 'serie': x['serie'], 'situacao': 'Inutilizada',
                        'cliente': f'Número inutilizado - {(x.get("just") or "")[:60]}', 'doc': '', 'uf': '', 'cfop': '', 'natop': '',
                        'chave': f'Protocolo {x.get("protocolo") or "-"}', 'acumulador': '',
                        'importar': 'Sim - como INUTILIZADA (situação 05, só o número)',
                        'conferencia': conf.get(num, ''), 'analise': ''})
    vistos, unicos = set(), []
    for r in out:
        k = (r['serie'], r['numero'], r['situacao'])
        if k not in vistos:
            vistos.add(k)
            unicos.append(r)
    unicos.sort(key=lambda r: (r['saida'] or r['emissao'] or a.fim, int(r['numero'])))
    return unicos


def gerar(a, destino, conf=None):
    rows = linhas(a, conf)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Saídas'
    emp = a.cfg['empresa']
    ws['A1'] = f'{emp["nome"]} - CNPJ {emp["cnpj"]} - Notas fiscais de saída {a.mes[5:]}/{a.mes[:4]}'
    ws['A1'].font = F(bold=True, size=12)
    ws['A2'] = ('Valores exatamente como estão no XML, na ordem da data de saída e do número. '
                'Vermelho = cancelada · cinza = inutilizada. Os totais somam só as autorizadas.')
    ws['A2'].font = F(italic=True, color='595959')
    hr = 4
    for c, (tit, larg, _, _) in enumerate(COLS, 1):
        cell = ws.cell(hr, c, tit)
        cell.font = F(bold=True, color='FFFFFF'); cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[get_column_letter(c)].width = larg
    ws.row_dimensions[hr].height = 30
    r = hr + 1
    for x in rows:
        for c, (_, _, k, moeda) in enumerate(COLS, 1):
            v = x.get(k, '' if not moeda else None)
            if moeda and x['situacao'] == 'Inutilizada':
                v = None
            cell = ws.cell(r, c, v)
            cell.font = F(); cell.border = B
            if moeda:
                cell.number_format = MOEDA
            elif k in ('saida', 'emissao'):
                cell.number_format = 'DD/MM/YYYY'
        r += 1
    ult = r - 1
    sit = f'$E${hr + 1}:$E${ult}'
    ws.cell(r, 1, 'Total autorizadas').font = F(bold=True)
    for c, (_, _, k, moeda) in enumerate(COLS, 1):
        ws.cell(r, c).fill = TOT
        if moeda:
            col = get_column_letter(c)
            cell = ws.cell(r, c, f'=SUMIF({sit},"Autorizada",{col}{hr + 1}:{col}{ult})')
            cell.font = F(bold=True); cell.number_format = MOEDA
    ws.cell(r, 3, f'=COUNTIF({sit},"Autorizada")').font = F(bold=True)
    rng = f'A{hr + 1}:{get_column_letter(len(COLS))}{ult}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$E{hr + 1}="Cancelada"'], fill=CANC))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$E{hr + 1}="Inutilizada"'], fill=INUT))
    r += 2
    resumo = [('Autorizadas', f'=COUNTIF({sit},"Autorizada")'), ('Canceladas', f'=COUNTIF({sit},"Cancelada")'),
              ('Inutilizadas', f'=COUNTIF({sit},"Inutilizada")'), ('Total de números no mês', f'=COUNTA(C{hr + 1}:C{ult})')]
    for rot, f in resumo:
        ws.cell(r, 1, rot).font = F(bold=True)
        ws.cell(r, 3, f).font = F(bold=True)
        r += 1
    ws.freeze_panes = ws.cell(hr + 1, 4)
    ws.auto_filter.ref = f'A{hr}:{get_column_letter(len(COLS))}{ult}'
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    return rows
