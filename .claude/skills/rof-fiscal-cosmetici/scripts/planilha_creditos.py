"""Planilha da análise de créditos das entradas e conferência XML x relatórios (Creditos_Entradas_Cosmetici_MM-AAAA.xlsx)."""
import collections

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from creditos import BASE_LEGAL
from planilha import ALTA, EDIT, F, MEDIA, MOEDA, OK, _bloco, _cab, _larguras, _linha, _titulo

COR_NAT = {'insumo': OK, 'devolucao': OK, 'confirmar': MEDIA, 'ativo': MEDIA}


def resumo(wb, c, cs, ce):
    a = c.a
    ws = wb.active; ws.title = 'Resumo'
    t = c.totais()
    _titulo(ws, f'COSMETICI - créditos das entradas e conferência {a.mes[5:]}/{a.mes[:4]}',
            'Indústria de cosméticos, Lucro Presumido (PIS/COFINS cumulativo: sem crédito de entradas). Base legal na última aba.')
    r = _bloco(ws, 4, 'Créditos de ICMS do mês')
    _cab(ws, r, ['Origem', 'Valor', 'Observação']); r += 1
    for rot, v, obs in [
        ('NF-e de fornecedor do regime normal (insumo/embalagem)', t.get('icms', 0), 'ICMS destacado - entra na apuração'),
        ('NF-e de fornecedor do Simples Nacional (vCredICMSSN)', t.get('sn', 0), 'CSOSN 101/201/900 com valor informado - entra na apuração'),
        ('CT-e com XML (frete tomado)', t.get('cte_xml', 0), 'Entra na apuração'),
        ('CT-e só no relatório SIEG (sem XML na pasta)', t.get('cte_pend', 0), 'Baixar os XML para poder creditar'),
        ('Simples com % informado mas valor zerado (potencial)', t.get('pot_sn', 0), 'NÃO creditar sem nota complementar do fornecedor'),
        ('Natureza a confirmar (ICMS destacado)', t.get('confirmar', 0), 'Confirmar se é insumo'),
    ]:
        _linha(ws, r, [rot, v, obs], moeda=(1,), fill=OK if 'entra' in obs.lower() else MEDIA); r += 1
    tot = round(t.get('icms', 0) + t.get('sn', 0) + t.get('cte_xml', 0), 2)
    _linha(ws, r, ['Total aproveitável agora', tot, 'NF-e normal + Simples + CT-e com XML'], moeda=(1,), bold=True); r += 1
    _linha(ws, r, ['Crédito de IPI (insumos de contribuinte do IPI)', t.get('ipi', 0), 'Simples não gera crédito de IPI'], moeda=(1,), bold=True); r += 1
    _linha(ws, r, ['DIFAL estimado (uso e consumo/ativo de outra UF)', t.get('difal', 0), 'Débito a recolher/apurar - conferir'], moeda=(1,), fill=MEDIA); r += 1

    r = _bloco(ws, r + 1, 'Entradas por natureza')
    _cab(ws, r, ['Natureza', 'Itens', 'Vlr contábil', 'ICMS destacado', 'Crédito ICMS', 'Crédito Simples', 'Crédito IPI']); r += 1
    g = collections.defaultdict(collections.Counter)
    for l in c.linhas:
        k = l['natureza']; x = g[k]
        x['n'] += 1; x['vc'] += l['vc']; x['icms'] += l['item']['vicms']; x['ci'] += l['cred_icms']; x['sn'] += l['cred_sn']; x['ipi'] += l['cred_ipi']
    nomes = {'insumo': 'Insumo / embalagem / rótulo', 'uso': 'Uso e consumo', 'ativo': 'Ativo imobilizado (possível)',
             'remessa': 'Remessa de terceiro (industrialização/conta e ordem)', 'outras': 'Outras saídas do fornecedor (amostra/brinde)',
             'financeiro': 'Nota financeira (antecipação)', 'devolucao': 'Devolução de venda', 'confirmar': 'A confirmar'}
    for k, x in sorted(g.items(), key=lambda kv: -kv[1]['vc']):
        _linha(ws, r, [nomes.get(k, k), x['n'], round(x['vc'], 2), round(x['icms'], 2), round(x['ci'], 2), round(x['sn'], 2), round(x['ipi'], 2)],
               moeda=(2, 3, 4, 5, 6), fill=COR_NAT.get(k)); ws.cell(r, 2).number_format = '0'; r += 1

    r = _bloco(ws, r + 1, 'Fornecedores do Simples Nacional')
    _cab(ws, r, ['Situação', 'Itens', 'Vlr contábil', 'Crédito']); r += 1
    sn = collections.defaultdict(collections.Counter)
    for l in c.linhas:
        if not l['simples']:
            continue
        if l['cred_sn']:
            k = '1. Com crédito informado (CSOSN 101/201) - aproveitar'
        elif l['pot_sn']:
            k = '2. % informado mas valor zerado - pedir correção'
        elif (l['item']['csosn'] or '') in ('101', '201', '900') and l['natureza'] in ('insumo', 'devolucao'):
            k = '3. CSOSN com crédito mas 0% - pedir correção'
        elif l['natureza'] not in ('insumo', 'devolucao'):
            k = '5. Não é insumo/revenda (uso, remessa, outras) - sem crédito'
        else:
            k = '4. CSOSN 102/103/300/400/500 - sem direito'
        sn[k]['n'] += 1; sn[k]['vc'] += l['vc']; sn[k]['c'] += l['cred_sn'] or l['pot_sn']
    for k in sorted(sn):
        _linha(ws, r, [k, sn[k]['n'], round(sn[k]['vc'], 2), round(sn[k]['c'], 2)], moeda=(2, 3)); ws.cell(r, 2).number_format = '0'; r += 1

    r = _bloco(ws, r + 1, 'Conferência XML x relatórios')
    ok_s = sum(1 for x in cs if x['obs'] == 'OK'); ok_e = sum(1 for x in ce if x['obs'] == 'OK')
    for rot, v in [('Notas de saída conferidas (XML x ERP x SIEG)', f'{ok_s} OK de {len(cs)}'),
                   ('Notas de entrada conferidas (XML x SIEG)', f'{ok_e} OK de {len(ce)}')]:
        _linha(ws, r, [rot, v]); r += 1
    _larguras(ws, [58, 16, 58, 16, 14, 16, 14])


def itens(wb, c):
    ws = wb.create_sheet('Entradas - crédito item a item')
    _titulo(ws, 'Cada item de NF-e de entrada com a natureza, o crédito e o motivo',
            'Verde = gera crédito; laranja = conferir. Colunas amarelas para a contadora (Confere?/Observação).')
    H = ['Emissão', 'NF', 'Fornecedor', 'CNPJ', 'UF', 'Regime', 'CFOP forn.', 'CFOP entrada', 'NCM', 'Produto', 'Vlr contábil',
         'CST/CSOSN', 'ICMS destacado', '% Cred. SN', 'Cred. SN (nota)', 'Natureza', 'Classificado por', 'Crédito ICMS', 'Crédito Simples',
         'Potencial Simples', 'Crédito IPI', 'DIFAL estimado', 'Motivo', 'Ação', 'Confere?', 'Observação']
    _cab(ws, 4, H)
    r = 5
    for l in sorted(c.linhas, key=lambda l: (l['nota']['dhemi'], l['nota']['emit_nome'] or '', l['nota']['numero'], l['item']['n_item'])):
        n, i = l['nota'], l['item']
        fill = OK if (l['cred_icms'] or l['cred_sn']) else (MEDIA if (l['acao'] or l['natureza'] == 'confirmar') else None)
        _linha(ws, r, [n['dhemi'], n['numero'], n['emit_nome'], n['emit_cnpj'], n['emit_uf'], l['regime'], i['cfop'], l['cfop_ent'], i['ncm'],
                       i['xprod'], l['vc'], i['csosn'] or i['cst_icms'], i['vicms'], i['pcredsn'] or None, i['vcredsn'] or None,
                       l['natureza_desc'], l['origem'], l['cred_icms'], l['cred_sn'], l['pot_sn'], l['cred_ipi'], l['difal'],
                       l['motivo'], l['acao'], '', ''], moeda=(10, 12, 14, 17, 18, 19, 20, 21), fill=fill)
        ws.cell(r, 25).fill = EDIT; ws.cell(r, 26).fill = EDIT
        r += 1
    _linha(ws, r, ['Total'] + [''] * 9 + [f'=SUM(K5:K{r-1})', '', f'=SUM(M5:M{r-1})', '', f'=SUM(O5:O{r-1})', '', '',
                                         f'=SUM(R5:R{r-1})', f'=SUM(S5:S{r-1})', f'=SUM(T5:T{r-1})', f'=SUM(U5:U{r-1})', f'=SUM(V5:V{r-1})', '', '', '', ''],
           moeda=(10, 12, 14, 17, 18, 19, 20, 21), bold=True)
    if r > 5:
        dv = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'Y5:Y{r-1}')
    _larguras(ws, [11, 8, 30, 15, 4, 15, 7, 8, 10, 34, 12, 8, 11, 7, 10, 34, 22, 11, 11, 11, 10, 10, 60, 40, 9, 26])
    ws.freeze_panes = 'D5'; ws.auto_filter.ref = f'A4:Z{r-1}'


def simples(wb, c):
    ws = wb.create_sheet('Simples Nacional')
    _titulo(ws, 'Notas de fornecedores do Simples Nacional - crédito pelo campo vCredICMSSN (LC 123/2006, art. 23)',
            'Só gera crédito: CSOSN 101/201/900, com valor informado na nota, em mercadoria para industrialização ou comercialização.')
    _cab(ws, 4, ['Emissão', 'NF', 'Fornecedor', 'CNPJ', 'CSOSN', 'Situação do CSOSN', 'Produto', 'Natureza', 'Vlr produto', '% crédito (pCredSN)',
                 'Crédito informado (vCredICMSSN)', 'Crédito aproveitável', 'Potencial (% x valor)', 'Motivo', 'Ação'])
    r = 5
    for l in c.linhas:
        if not l['simples']:
            continue
        n, i = l['nota'], l['item']
        cs = i['csosn'] or ''
        from creditos import CSOSN_DESCR
        _linha(ws, r, [n['dhemi'], n['numero'], n['emit_nome'], n['emit_cnpj'], cs, CSOSN_DESCR.get(cs, ''), i['xprod'], l['natureza_desc'],
                       i['vprod'], i['pcredsn'] or None, i['vcredsn'] or None, l['cred_sn'], l['pot_sn'], l['motivo'], l['acao']],
               moeda=(8, 10, 11, 12), fill=OK if l['cred_sn'] else (MEDIA if l['acao'] else None))
        r += 1
    _linha(ws, r, ['Total'] + [''] * 7 + [f'=SUM(I5:I{r-1})', '', f'=SUM(K5:K{r-1})', f'=SUM(L5:L{r-1})', f'=SUM(M5:M{r-1})', '', ''],
           moeda=(8, 10, 11, 12), bold=True)
    _larguras(ws, [11, 8, 30, 15, 7, 30, 34, 30, 12, 9, 13, 12, 12, 60, 50])
    ws.freeze_panes = 'D5'


def cte(wb, c):
    ws = wb.create_sheet('CT-e (frete)')
    _titulo(ws, 'CT-e em que a COSMETICI é tomadora - crédito de ICMS do frete')
    _cab(ws, 3, ['CT-e', 'Data', 'Transportadora', 'CFOP entrada', 'Valor', 'ICMS', 'Crédito', 'Fonte', 'Observação'])
    r = 4
    for x in c.cte:
        _linha(ws, r, [x['numero'], x['data'], x['emitente'], x['cfop'], x['valor'], x['icms'], x['credito'], x['fonte'], x['obs']],
               moeda=(4, 5, 6), fill=OK if x['fonte'] == 'XML' and x['credito'] else MEDIA)
        r += 1
    _linha(ws, r, ['Total', '', '', '', f'=SUM(E4:E{r-1})', f'=SUM(F4:F{r-1})', f'=SUM(G4:G{r-1})', '', ''], moeda=(4, 5, 6), bold=True)
    _larguras(ws, [11, 11, 36, 9, 11, 10, 10, 30, 70])


def conferencia(wb, cs, ce):
    ws = wb.create_sheet('Conferência saídas')
    _titulo(ws, 'Saídas: XML x relatório de faturamento do ERP (Cosmos) x relatório SIEG')
    _cab(ws, 3, ['NF', 'Emissão', 'Cliente', 'Vlr XML', 'Vlr ERP', 'Vlr SIEG', 'ICMS XML', 'ICMS ERP', 'IPI XML', 'IPI ERP', 'Resultado'])
    r = 4
    for x in cs:
        _linha(ws, r, [x['numero'], x['data'], x['cliente'], x['xml'], x['erp'], x['sieg'], x['icms_xml'], x['icms_erp'], x['ipi_xml'], x['ipi_erp'], x['obs']],
               moeda=tuple(range(3, 10)), fill=None if x['obs'] == 'OK' else (MEDIA if x['obs'] in ('cancelada', 'inutilizada') else ALTA))
        r += 1
    _linha(ws, r, ['Total', '', ''] + [f'=SUM({get_column_letter(k)}4:{get_column_letter(k)}{r-1})' for k in range(4, 11)] + [''],
           moeda=tuple(range(3, 10)), bold=True)
    _larguras(ws, [8, 11, 36, 12, 12, 12, 11, 11, 10, 10, 50])
    ws.freeze_panes = 'A4'

    ws = wb.create_sheet('Conferência entradas')
    _titulo(ws, 'Entradas: XML x relatório SIEG')
    _cab(ws, 3, ['NF', 'Emissão', 'Fornecedor', 'CNPJ', 'Vlr XML', 'Vlr SIEG', 'Resultado', 'Chave'])
    r = 4
    for x in ce:
        _linha(ws, r, [x['numero'], x['data'], x['fornecedor'], x['cnpj'], x['xml'], x['sieg'], x['obs'], x['chave']], moeda=(4, 5),
               fill=None if x['obs'] == 'OK' else ALTA)
        r += 1
    _larguras(ws, [9, 11, 36, 16, 12, 12, 40, 47])
    ws.freeze_panes = 'A4'


def base_legal(wb):
    ws = wb.create_sheet('Base legal')
    _titulo(ws, 'Fundamentação usada na análise')
    _cab(ws, 3, ['Tema', 'Base legal'])
    for k, (t, b) in enumerate(BASE_LEGAL, 4):
        _linha(ws, k, [t, b]); ws.cell(k, 2).alignment = Alignment(wrap_text=True, vertical='top')
    _larguras(ws, [40, 120])


def gerar(c, cs, ce, destino):
    wb = Workbook()
    resumo(wb, c, cs, ce)
    itens(wb, c)
    simples(wb, c)
    cte(wb, c)
    conferencia(wb, cs, ce)
    base_legal(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
