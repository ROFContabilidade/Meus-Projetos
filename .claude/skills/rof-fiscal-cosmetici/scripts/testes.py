"""Testes rápidos com dados sintéticos (não usam arquivos de cliente).  python testes.py"""
import copy
import datetime as dt
import json
import os
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import leitores  # noqa: E402
import txt_dominio  # noqa: E402
from apuracao import Apuracao  # noqa: E402

CFG = json.load(open(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..', 'Cosmetici', 'Fiscal', 'config', 'cosmetici.json'), encoding='utf-8'))
CNPJ = CFG['empresa']['cnpj']


def nfe(chave, num, emit, dest, uf_dest, itens, data='2026-08-10', tpnf='1', fin='1'):
    det = ''
    for k, (cfop, ncm, vprod, picms, vipi, vst, cst_pis) in enumerate(itens, 1):
        vicms = round(vprod * picms / 100, 2)
        det += (f'<det nItem="{k}"><prod><cProd>{k}</cProd><xProd>P{k}</xProd><NCM>{ncm}</NCM><CFOP>{cfop}</CFOP><qCom>1</qCom>'
                f'<vProd>{vprod}</vProd></prod><imposto><ICMS><ICMS00><orig>0</orig><CST>00</CST><vBC>{vprod}</vBC><pICMS>{picms}</pICMS>'
                f'<vICMS>{vicms}</vICMS><vICMSST>{vst}</vICMSST></ICMS00></ICMS><IPI><IPITrib><CST>50</CST><vBC>{vprod}</vBC><pIPI>10</pIPI>'
                f'<vIPI>{vipi}</vIPI></IPITrib></IPI><PIS><PISAliq><CST>{cst_pis}</CST><vBC>{vprod}</vBC><pPIS>2.2</pPIS><vPIS>0</vPIS></PISAliq></PIS>'
                f'</imposto></det>')
    return (f'<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe"><NFe><infNFe Id="NFe{chave}"><ide><mod>55</mod><serie>1</serie>'
            f'<nNF>{num}</nNF><dhEmi>{data}T10:00:00-03:00</dhEmi><tpNF>{tpnf}</tpNF><finNFe>{fin}</finNFe></ide>'
            f'<emit><CNPJ>{emit}</CNPJ><xNome>EMIT {emit[:4]}</xNome><enderEmit><UF>PR</UF></enderEmit></emit>'
            f'<dest><CNPJ>{dest}</CNPJ><xNome>DEST {dest[:4]}</xNome><enderDest><UF>{uf_dest}</UF></enderDest><indIEDest>1</indIEDest></dest>'
            f'{det}<total><ICMSTot><vNF>0</vNF></ICMSTot></total></infNFe></NFe>'
            f'<protNFe><infProt><cStat>100</cStat></infProt></protNFe></nfeProc>')


def cancelamento(chave):
    return (f'<procEventoNFe xmlns="http://www.portalfiscal.inf.br/nfe"><evento><infEvento><chNFe>{chave}</chNFe>'
            f'<dhEvento>2026-08-11T10:00:00-03:00</dhEvento><tpEvento>110111</tpEvento></infEvento></evento></procEventoNFe>')


def montar_zip(arquivos):
    d = tempfile.mkdtemp()
    p = os.path.join(d, 'm.zip')
    with zipfile.ZipFile(p, 'w') as z:
        for nome, txt in arquivos.items():
            z.writestr(nome, txt)
    return p


def teste_apuracao():
    ch = lambda n: f'412608{CNPJ}55001{n:09d}1000000{n % 10}'[:44].ljust(44, '0')
    forn = '11111111000191'
    arqs = {
        'v1.xml': nfe(ch(1), 1, CNPJ, '22222222000191', 'PR', [('5401', '33059000', 1000.0, 12, 143.0, 300.0, '02')]),
        'v2.xml': nfe(ch(2), 2, CNPJ, '33333333000191', 'SP', [('6124', '33059000', 2000.0, 12, 0, 0, '02')]),
        'v3.xml': nfe(ch(3), 3, CNPJ, '22222222000191', 'PR', [('5101', '33059000', 500.0, 19.5, 0, 0, '02')]),
        'c3.xml': cancelamento(ch(3)),
        'e1.xml': nfe(ch(9), 77, forn, CNPJ, 'PR', [('5101', '28470000', 400.0, 12, 0, 0, '01')]),
        'v5.xml': nfe(ch(5), 5, CNPJ, '22222222000191', 'SC', [('5102', '33059000', 100.0, 12, 0, 0, '02')]),
    }
    fontes = leitores.ler_xmls([montar_zip(arqs)])
    a = Apuracao(copy.deepcopy(CFG), '2026-08', fontes, {}).executar()
    assert len(a.saidas) == 3 and len(a.canceladas) == 1 and len(a.entradas) == 1, (len(a.saidas), len(a.canceladas), len(a.entradas))
    # faturamento = vendas (1000+143+300 + 100) + indus 2000 - ST 300 - IPI 143
    assert a.fat['liquido'] == 3100.0, a.fat
    # PIS: indus (2000-240)*0.65% + vendas (1443-120 + 100-12 - 143 - 300)*2.2%
    assert a.pc['base_industrializacao'] == 1760.0 and a.pc['base_vendas'] == 968.0, a.pc
    assert a.pc['pis'] == round(round(1760 * 0.0065, 2) + round(968 * 0.022, 2), 2)
    assert a.ic['debitos'] == 120 + 240 + 12 and a.ic['creditos_nfe'] == 48.0, a.ic
    assert a.st['recolher_pr'] == 300.0 and a.ip['recolher'] == 143.0
    tipos = {d[1] for d in a.div}
    assert 'CFOP x UF' in tipos, tipos                      # 5102 para SC
    assert 'Fornecedor sem histórico' in tipos, tipos
    assert any(d[1] == 'Numeração' and 'NF 4/' in d[2] for d in a.div), a.div   # número 4 faltando
    print('ok apuração')


def teste_irpj_trimestre():
    cfg = copy.deepcopy(CFG)
    a = Apuracao(cfg, '2026-06', {'nfe': {}, 'eventos': [], 'inut': [], 'cte': {}}, {})
    a.fat = {'liquido': 285690.16}
    r = a.irpj_csll()
    assert r['receita'] == 998642.37 and r['irpj'] == 13972.85 and r['csll'] == 10785.34, r   # DARFs do 2º tri/2026
    print('ok IRPJ/CSLL 2º trimestre = DARFs')


def teste_txt():
    cfg = copy.deepcopy(CFG)
    cfg['contas_dominio']['pis'].update(debito=111, credito=222)
    a = Apuracao(cfg, '2026-08', {'nfe': {}, 'eventos': [], 'inut': [], 'cte': {}}, {})
    a.pc = {'pis': 1234.5, 'cofins': 0}
    a.ip = {'recolher': 0}; a.ic = {'recolher': 0}; a.st = {'recolher_pr': 10.0}
    a.ir = {'fechamento': False, 'trimestre': '3º trimestre/2026'}
    d = tempfile.mkdtemp(); p = os.path.join(d, 't.txt')
    lanc, sem = txt_dominio.gerar(a, p)
    linhas = open(p, encoding='cp1252').read().splitlines()
    assert linhas == [f'|0000|{CNPJ}|', '|6000|X||||', '|6100|31/08/2026|111|222|1234,50||PROVISAO PIS S/ FATURAMENTO 08/2026||||'], linhas
    assert open(p, 'rb').read().count(b'\r\n') == 3 and sem == [('icms_st', 10.0)]
    print('ok TXT Domínio')


def teste_sieg_cte():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    cab = ['Chave', 'Numero', 'Serie', 'Dt_Emissao', 'CNPJ_CPF_Emitente', 'Rz_Emit', 'Uf_Inicio', 'Uf_Final', 'CNPJ_CPF_Tomador',
           'CNPJ_CPF_Outro_Tomador', 'CNPJ_CPF_Dest', 'CNPJ_CPF_Rem', 'CFOP', 'Valor_Transporte', 'CST', 'ICMS_Base_Calculo', 'ICMS_Aliq', 'ICMS', 'Chave_NFe', 'Status']
    ws.append(['Relatorio_CTe.xlsx']); ws.append(cab)
    ws.append(['4' * 44, '2079446', '1', '21/09/2026', '00428307002301', 'TRANSP', 'PR', 'SC', CNPJ, '', '1', CNPJ, '6352', 'R$ 88.00', '00', 'R$ 84.21', 'R$ 12.00', 'R$ 10.11', '', 'Autorizado o uso do CTe'])
    ws.append(['5' * 44, '2', '1', '21/09/2026', '00428307002301', 'TRANSP', 'PR', 'SC', '', '17666331000168', '1', CNPJ, '6353', 'R$ 50.00', '00', 'R$ 50.00', 'R$ 12.00', 'R$ 6.00', '', 'Cancelamento'])
    p = os.path.join(tempfile.mkdtemp(), 'Relatorio_CTe.xlsx'); wb.save(p)
    c = leitores.ler_sieg_cte(p)
    x = c['4' * 44]
    assert x['vicms'] == 10.11 and x['toma_cnpj'] == CNPJ and x['data'] == dt.date(2026, 9, 21) and x['autorizada']
    assert not c['5' * 44]['autorizada'] and c['5' * 44]['toma_cnpj'] == '17666331000168'
    print('ok relatório SIEG de CT-e')


def teste_simples():
    from creditos import Credito
    def item(cfop, ncm, vprod, csosn, pcred=0.0, vcred=0.0, vicms=0.0, cst=None, vipi=0.0):
        return {'n_item': 1, 'xprod': 'X', 'ncm': ncm, 'cfop': cfop, 'vprod': vprod, 'vdesc': 0.0, 'vfrete': 0.0, 'vseg': 0.0, 'voutro': 0.0,
                'vipi': vipi, 'vst': 0.0, 'vfcpst': 0.0, 'vipidevol': 0.0, 'cst_icms': cst or csosn, 'csosn': csosn, 'pcredsn': pcred,
                'vcredsn': vcred, 'vicms': vicms}
    def nota(crt, it, uf='PR', cnpj='99999999000191'):
        return {'numero': 1, 'serie': '1', 'dhemi': dt.date(2026, 9, 5), 'emit_cnpj': cnpj, 'emit_nome': 'F', 'emit_uf': uf, 'emit_crt': crt,
                'finnfe': '1', 'propria': False, 'chave': 'x', 'itens': [it]}
    class A:  # apuração mínima
        cnpj = CNPJ; cfg = CFG; ctes = []
    casos = [
        (nota('1', item('5101', '39191090', 1000, '101', 2.86, 28.60)), 'insumo', 28.60, 0, 0),       # rótulo, Simples com crédito
        (nota('1', item('5101', '48211000', 1000, '102')), 'insumo', 0, 0, 0),                        # Simples sem permissão
        (nota('1', item('5102', '29221100', 1000, '101', 3.0, 0.0)), 'insumo', 0, 30.0, 0),           # % sem valor -> potencial
        (nota('1', item('5102', '64039190', 100, '102')), 'uso', 0, 0, 0),                             # EPI
        (nota('1', item('5901', '39233090', 500, '400')), 'remessa', 0, 0, 0),                         # insumo do cliente
        (nota('3', item('6101', '33029019', 1000, None, vicms=120, cst='00', vipi=50), 'SP'), 'insumo', 120, 0, 50),
        (nota('3', item('6102', '84439923', 1000, None, vicms=120, cst='00'), 'SP'), 'uso', 0, 0, 0),  # uso e consumo -> DIFAL
    ]
    A.entradas = [c[0] for c in casos]
    c = Credito(A, {}).analisar()
    for (n, nat, cred, pot, ipi), l in zip(casos, c.linhas):
        assert l['natureza'] == nat, (n['itens'][0]['ncm'], l['natureza'])
        assert round(l['cred_icms'] + l['cred_sn'], 2) == cred and l['pot_sn'] == pot and l['cred_ipi'] == ipi, l
    assert c.linhas[-1]['difal'] > 0
    print('ok crédito Simples/insumo/uso e consumo')


def teste_revisao():
    import revisao
    ch = lambda n: f'412609{CNPJ}55001{n:09d}2000000{n % 10}'[:44].ljust(44, '0')
    cli, gula = '44444444000191', '55555555000191'
    cpl = lambda x, txt: x.replace('</infNFe>', f'<infAdic><infCpl>{txt}</infCpl></infAdic></infNFe>')
    cst41 = lambda x: x.replace('<CST>00</CST>', '<CST>41</CST>').replace('<pICMS>12</pICMS><vICMS>72.0</vICMS>', '<pICMS>0</pICMS><vICMS>0</vICMS>')
    final = lambda x: x.replace('<indIEDest>1</indIEDest>', '<IE>9084324601</IE><indIEDest>9</indIEDest>')
    arqs = {
        # industrialização do cliente + frasco em 5124 (CST 41) + o mesmo frasco no retorno 5902
        'i1.xml': cpl(nfe(ch(1), 101, CNPJ, cli, 'PR', [('5124', '28470000', 1000.0, 0, 0, 0, '02')], data='2026-09-05'), 'INDUSTRIALIZACAO COM INSUMOS DA NFE 555'),
        'i2.xml': cst41(nfe(ch(2), 102, CNPJ, cli, 'PR', [('5124', '39233090', 600.0, 12, 0, 0, '49')], data='2026-09-05')),
        'i3.xml': cpl(cst41(nfe(ch(3), 103, CNPJ, cli, 'PR', [('5902', '39233090', 600.0, 12, 0, 0, '49')], data='2026-09-05')),
                      'RETORNO SIMBOLICO DOS INSUMOS ENVIADOS NA NFE 555; INDUSTRIALIZACAO EFETUADA NA NFE 104'),
        # 6124 sem remessa nem retorno
        'i4.xml': nfe(ch(4), 104, CNPJ, '66666666000191', 'SC', [('6124', '33059000', 2000.0, 12, 0, 0, '02')], data='2026-09-06'),
        # cliente com IE: primeira como contribuinte com ST, depois como consumidor final
        'g1.xml': nfe(ch(5), 105, CNPJ, gula, 'PR', [('5401', '33059000', 1000.0, 12, 143.0, 300.0, '02')], data='2026-09-03'),
        'g2.xml': final(nfe(ch(6), 106, CNPJ, gula, 'PR', [('5102', '33059000', 1000.0, 19.5, 143.0, 0, '02')], data='2026-09-10')),
        'g3.xml': final(nfe(ch(7), 107, CNPJ, gula, 'SP', [('6102', '33059000', 100.0, 12, 0, 0, '02')], data='2026-09-11')),
    }
    fontes = leitores.ler_xmls([montar_zip(arqs)])
    fontes['nfse'] = [
        {'numero': '1', 'data': '01/09/2026', 'cod_servico': '0401', 'prestador': 'MEDICINA OCUP', 'cnpj': '1', 'uf': 'PR', 'simples': False,
         'cancelada': False, 'valor': 728.14, 'base': 728.14, 'liquido': 717.22, 'desconto': 0.0, 'iss': 0, 'iss_retido': 0.0,
         'ir': 10.92, 'pis': 0, 'cofins': 0, 'outras': 0, 'descricao': ''},
        {'numero': '2', 'data': '25/09/2026', 'cod_servico': '1401', 'prestador': 'CALIBRACAO', 'cnpj': '2', 'uf': 'SP', 'simples': False,
         'cancelada': False, 'valor': 1835.50, 'base': 1835.50, 'liquido': 1750.15, 'desconto': 0.0, 'iss': 0, 'iss_retido': 0.0,
         'ir': 0, 'pis': 11.93, 'cofins': 55.06, 'outras': 85.35, 'descricao': ''},
        {'numero': '3', 'data': '17/09/2026', 'cod_servico': '1401', 'prestador': 'SIMPLES', 'cnpj': '3', 'uf': 'PR', 'simples': True,
         'cancelada': False, 'valor': 1090.0, 'base': 1090.0, 'liquido': 1090.0, 'desconto': 0.0, 'iss': 0, 'iss_retido': 0.0,
         'ir': 0, 'pis': 0, 'cofins': 0, 'outras': 0, 'descricao': ''},
    ]
    a = Apuracao(copy.deepcopy(CFG), '2026-09', fontes, {}).executar()
    por = lambda tipo: [d for d in a.div if d[1] == tipo]
    dup = [d for d in por('Industrialização x retorno') if d[0] == 'ALTA']
    assert len(dup) == 1 and 'NF 102/' in dup[0][2] and 'NF 103' in dup[0][3] and dup[0][4] == 600.0, dup
    ref = [d for d in por('Industrialização x retorno') if d[0] == 'BAIXA']
    assert len(ref) == 1 and 'NF 104' in ref[0][3] and 'NF 101' in ref[0][3], ref        # retorno cita a NF errada
    sem = por('Industrialização sem insumo')
    assert [d[2][:10] for d in sem] == ['NF 104/1 0'], sem
    fin = por('Contribuinte como consumidor final')
    assert len(fin) == 1 and '106' in fin[0][3] and '107' in fin[0][3] and 'ST de R$ 300.00' in fin[0][3], fin
    assert len(por('Destinatário x UF')) == 1
    assert any('NF 106/' in d[2] for d in por('CFOP revenda x produção'))
    assert not por('NFS-e retenções'), 'NFS-e não pode ter retenção recalculada - segue o documento'
    ext = revisao.analisar_nfse(fontes['nfse'])
    g = {x['numero']: x for x in ext}
    assert g['1']['ir'] == 10.92 and g['1']['descontado'] == 10.92 and 'situacao' not in g['1'], g['1']   # como está na nota
    assert g['2']['outras'] == 85.35 and g['2']['liquido'] == 1750.15
    ck = {l[1]: l[3] for l in revisao.checklist(a)}
    assert ck['NFS-e tomadas: retenções extraídas do documento (sem recálculo)'] == 'OK'
    assert ck['Diferimento nas 5124 internas'] == 'OK'
    print('ok revisão automática (industrialização, consumidor final) e NFS-e extraída como está na nota')


def teste_cte_natureza():
    from creditos import Credito
    ch = lambda n: f'412609{CNPJ}55001{n:09d}3000000{n % 10}'[:44].ljust(44, '0')
    forn = '77777777000191'
    arqs = {'e1.xml': nfe(ch(1), 900, forn, CNPJ, 'PR', [('5102', '96033000', 900.0, 12, 0, 0, '01')], data='2026-09-02'),
            'e2.xml': nfe(ch(2), 901, forn, CNPJ, 'PR', [('5101', '39233090', 900.0, 12, 0, 0, '01')], data='2026-09-02')}
    fontes = leitores.ler_xmls([montar_zip(arqs)])
    for k, nf in ((1, ch(1)), (2, ch(2))):
        fontes['cte'][f'c{k}'] = {'chave': f'c{k}', 'numero': k, 'serie': '1', 'data': dt.date(2026, 9, 3), 'cfop': '6352',
                                  'emit_cnpj': '8', 'emit_nome': 'TRANSP', 'uf_ini': 'RS', 'uf_fim': 'PR', 'toma_cnpj': CNPJ,
                                  'rem_cnpj': forn, 'dest_cnpj': CNPJ, 'vprest': 100.0, 'cst': '00', 'vbc': 100.0, 'picms': 12.0,
                                  'vicms': 12.0, 'autorizada': True, 'chaves_nfe': [nf], 'origem': 'x'}
    a = Apuracao(copy.deepcopy(CFG), '2026-09', fontes, {}).executar()
    cred = {l['cte']['numero']: l['cred_icms'] for l in a.cte_cred}
    assert cred == {1: 0.0, 2: 12.0}, cred             # pincel (uso e consumo) x frasco (insumo)
    assert a.ic['creditos_cte'] == 12.0 and any(d[1] == 'CT-e sem crédito' for d in a.div)
    print('ok CT-e segue a natureza da NF-e')


def teste_registro_erp():
    import revisao
    ch = lambda n: f'412609{CNPJ}55001{n:09d}4000000{n % 10}'[:44].ljust(44, '0')
    arqs = {'a.xml': nfe(ch(1), 201, CNPJ, '22222222000191', 'SP', [('6102', '33059000', 1000.0, 12, 0, 0, '02')], data='2026-09-02')
            .replace('<CST>50</CST>', '<CST>51</CST>').replace('<pIPI>10</pIPI>', '<pIPI>0</pIPI>'),
            'b.xml': nfe(ch(2), 202, CNPJ, '22222222000191', 'PR', [('5102', '33059000', 1000.0, 19.5, 143.0, 0, '02')], data='2026-09-03'),
            'c.xml': nfe(ch(3), 203, CNPJ, '33333333000191', 'PR', [('5101', '33059000', 500.0, 19.5, 0, 0, '02')], data='2026-09-04')}
    txt = ('REGISTRO DE SAIDAS Dia Dest. NF 1 201 2 PR 1.000,00 6102 ICMS 1.000,00 12,00\n\n120,00 0,00 0,00 6102 IPI 1.000,00 0,00 0,00 0,00 0,00 '
           'NF 1 202 3 PR 1.143,00 5102 ICMS 1.000,00 19,50 195,00 0,00 0,00 5102 IPI 1.000,00 10,00 143,00 0,00 143,00 '
           'NF 1 204 4 PR NOTA FISCAL CANCELADA \\! Total Geral')
    reg = leitores.ler_registro_saidas_erp(txt)
    assert [r['numero'] for r in reg] == [201, 202, 204] and reg[1]['ipi'] == 143.0, reg
    fontes = leitores.ler_xmls([montar_zip(arqs)])
    fontes['registro_saidas'] = reg
    a = Apuracao(copy.deepcopy(CFG), '2026-09', fontes, {}).executar()
    d = [x for x in a.div if x[1] == 'XML x Registro ERP']
    assert any('NF 201/' in x[2] and 'UF no XML SP x ERP PR' in x[3] for x in d), d
    assert any('NF 203/' in x[2] and 'não está no Registro' in x[3] for x in d), d
    assert not any('NF 202/' in x[2] for x in d), d
    # mesmo produto (cProd 1) com IPI 0 na 201/203 e 10% na 202
    assert any(x[1] == 'Produto com tratamento diferente' for x in a.div)
    print('ok Registro de Saídas do ERP x XML')


def teste_difal_antecipacao():
    """Casos reais do SPED 07/2026: DIFAL Agatha NF 2230 (Simples, RS, 5.800,00 -> 540,37) e antecipado Emuna NF 2446 (4%, 1.475,00 -> 118,00)."""
    ch = lambda n: f'432607{n:014d}55001{n:09d}5000000{n % 10}'[:44].ljust(44, '0')
    agatha, emuna = '54079636000108', '10625467000180'
    simples = lambda x: (x.replace('<ICMS00><orig>0</orig><CST>00</CST><vBC>5800.0</vBC><pICMS>0</pICMS><vICMS>0.0</vICMS><vICMSST>0</vICMSST></ICMS00>',
                                   '<ICMSSN101><orig>0</orig><CSOSN>101</CSOSN><pCredSN>3.11</pCredSN><vCredICMSSN>0</vCredICMSSN></ICMSSN101>')
                         .replace('<enderEmit><UF>PR</UF></enderEmit>', '<enderEmit><UF>RS</UF></enderEmit><CRT>1</CRT>'))
    imp = lambda x: (x.replace('<orig>0</orig>', '<orig>1</orig>').replace('<enderEmit><UF>PR</UF></enderEmit>', '<enderEmit><UF>SP</UF></enderEmit><CRT>3</CRT>'))
    arqs = {'a.xml': simples(nfe(ch(1), 2230, agatha, CNPJ, 'PR', [('6101', '96033000', 5800.0, 0, 0, 0, '99')], data='2026-09-03')),
            'e.xml': imp(nfe(ch(2), 2446, emuna, CNPJ, 'PR', [('6102', '29322000', 1475.0, 4, 0, 0, '99')], data='2026-09-03'))}
    fontes = leitores.ler_xmls([montar_zip(arqs)])
    forn = {agatha: {'nome': 'AGATHA', 'cfop': '2556', 'icms': False, 'ipi': False, 'ncm': {'96033000': {'cfop': '2556', 'icms': False, 'ipi': False}}},
            emuna: {'nome': 'EMUNA', 'cfop': '2101', 'icms': True, 'ipi': False, 'ncm': {'29322000': {'cfop': '2101', 'icms': True, 'ipi': False}}}}
    a = Apuracao(copy.deepcopy(CFG), '2026-09', fontes, forn).executar()
    d = {l['nota']['numero']: l for l in a.credito.linhas}
    assert d[2230]['natureza'] == 'uso' and d[2230]['difal'] == 540.37, (d[2230]['natureza'], d[2230]['difal'])
    assert d[2446]['antecipado'] == 118.0 and d[2446]['cred_icms'] == 59.0, d[2446]
    # Demonstrativo 07/2026: outros débitos = DIFAL + antecipado; outros créditos = antecipado (saldo zero)
    assert a.ic['ajustes_debito'] == 658.37 and a.ic['ajustes_credito'] == 118.0 and a.antecipado['total'] == 118.0, a.ic
    assert a.ic['recolher'] == round(540.37 - 59.0, 2), a.ic
    print('ok DIFAL (base dupla, Simples) e ICMS antecipado de importado = SPED 07/2026')


def teste_contabil():
    """TXT contábil das entradas: conta do fornecedor (embalagem 600), ISS retido em lote D e IRRF em lote próprio."""
    import lancamentos_contabeis as lc
    from creditos import Credito
    ch = lambda n: f'4126091234567800019055001{n:09d}1{n:08d}'
    pack = '33598527000119'
    fontes = leitores.ler_xmls([montar_zip({'p.xml': nfe(ch(1), 29391, pack, CNPJ, 'PR', [('5101', '48191000', 2931.06, 19.5, 0, 0, '00')], data='2026-09-10')})])
    forn = {pack: {'nome': 'PIZZA PACK', 'cfop': '1101', 'icms': True, 'ipi': False, 'contas_dominio': {'compra': {'600': 6}}}}
    a = Apuracao(copy.deepcopy(CFG), '2026-09', fontes, forn).executar()
    cred = Credito(a, forn, [], {}).analisar()
    serv = {'numero': '15', 'emissao': '2026-09-12', 'prest_nome': 'ESTRE AMBIENTAL', 'prest_cnpj': '27827087000121', 'prest_uf': 'PR',
            'valor': 3000.0, 'iss_retido': 150.0, 'irrf': 45.0, 'csll': 0.0, 'inss': 0.0, 'chave': ''}
    a.cfg['acumuladores_servicos']['27827087000121'] = '806'
    lanc, pend = lc.montar(a, cred, {}, {}, [serv], forn)
    assert not pend, pend
    nf = next(x for x in lanc if x['doc'] == 'NF-e')
    assert nf['blocos'][0][1][0][:2] == ('600', '506'), nf
    destino = os.path.join(tempfile.mkdtemp(), 't.txt')
    lc.gerar_txt(CNPJ, lanc, destino)
    txt = open(destino, encoding='cp1252').read().split('\n')
    assert txt[0].strip() == f'|0000|{CNPJ}|'
    assert '|6100|10/09/2026|600|506|2931,06||COMPRAS DE MERCADORIAS NESTA DATA 29391 EMIT 3359 EMIT 3359 33.598.527/0001-19||||' in '\n'.join(txt)
    i = txt.index('|6000|D||||\r') if '|6000|D||||\r' in txt else txt.index('|6000|D||||')
    assert [t.split('|')[3:6] for t in txt[i + 1:i + 4]] == [['362', '', '3000,00'], ['', '506', '2850,00'], ['', '183', '150,00']], txt[i:i + 4]
    assert '|6100|12/09/2026|506|178|45,00||VALOR A RECOLHER||||' in '\n'.join(txt)
    print('ok TXT contábil das entradas (conta do fornecedor, ISS retido em lote D, IRRF)')


def teste_importacao():
    """SPED de importação: mesma quantidade de campos do arquivo aceito pelo Domínio (Kopp), contadores e totais do XML."""
    import collections as co
    import sped_importacao
    from creditos import Credito
    ch = lambda n, e: f'412609{e}55001{n:09d}1{n:08d}'
    forn_cnpj = '33598527000119'
    arqs = {'e.xml': nfe(ch(1, forn_cnpj), 29391, forn_cnpj, CNPJ, 'PR', [('5101', '48191000', 1000.0, 19.5, 0, 0, '00')], data='2026-09-10'),
            's.xml': nfe(ch(2, CNPJ), 9001, CNPJ, '11222333000181', 'SP', [('6101', '33059000', 500.0, 12, 0, 10, '01')], data='2026-09-11')}
    a = Apuracao(copy.deepcopy(CFG), '2026-09', leitores.ler_xmls([montar_zip(arqs)]), {}).executar()
    cred = Credito(a, {}, [], {}).analisar()
    destino = os.path.join(tempfile.mkdtemp(), 'imp.txt')
    s, L = sped_importacao.gerar(a, cred, {}, {}, {}, CFG['sped_cabecalho'], destino)
    campos = {'C100': 31, 'C170': 40, 'C190': 14, '0150': 15, '0200': 15, '0000': 17}
    for l in L:
        r = l.split('|')[1]
        if r in campos:
            assert len(l.split('|')) == campos[r], (r, len(l.split('|')), l)
    cont = co.Counter(l.split('|')[1] for l in L)
    assert L[-1] == f'|9999|{len(L):09d}|' and cont['C100'] == 2
    c170 = [l.split('|') for l in L if l.startswith('|C170|')]
    assert c170[0][11] == '1101' and c170[0][12] == '100' and c170[0][15] == '195,00', c170[0]   # CFOP de entrada, acumulador, crédito
    assert c170[1][11] == '6101' and c170[1][12] == '201', c170[1]
    docs, n_itens, dif = sped_importacao.auditar(L, a, None, cred)
    assert docs == 2 and n_itens == 2 and not dif, dif
    # compra fora da atividade-fim (uso e consumo): CST do XML, bases de ICMS/IPI/PIS zeradas, PIS/COFINS CST 98 (como a Kopp)
    cred_uso = Credito(a, {}, [], {forn_cnpj: 'uso'}).analisar()
    s2, L2 = sped_importacao.gerar(a, cred_uso, {}, {}, {}, CFG['sped_cabecalho'], destino)
    e = next(l.split('|') for l in L2 if l.startswith('|C170|') and '|1556|' in l)
    assert e[10] == '000' and e[13:16] == ['0,00', '0,00', '0,00'] and e[20] == '' and e[25] == '98' and e[31] == '98', e
    assert not sped_importacao.auditar(L2, a, None, cred_uso)[2]
    print('ok SPED de importação (leiaute aceito pelo Domínio, acumulador no COD_NAT, valores = XML, uso e consumo zerado)')


def teste_servicos():
    """Relatório de serviços tomados: separa com e sem retenção, valores do XML, retenções somadas por tributo."""
    import planilha_nfse
    from openpyxl import load_workbook
    base = {'emissao': '2026-09-10', 'competencia': '2026-09-10', 'prest_uf': 'PR', 'mun_incidencia': '4107652', 'simples': False,
            'cod_trib_nac': '170101', 'servico': 'x', 'descricao': 'servico', 'cstat': '100', 'aliq_iss': 2.0, 'pis': 0.0, 'cofins': 0.0,
            'tp_ret_piscofins': None, 'chave': '', 'iss_retido': 0.0, 'irrf': 0.0, 'csll': 0.0, 'inss': 0.0, 'total_ret': 0.0}
    xs = [dict(base, numero='1', prest_cnpj='11111111000191', prest_nome='SEM RETENCAO', valor=100.0, base_iss=100.0, iss=2.0, liquido=100.0),
          dict(base, numero='2', prest_cnpj='22222222000191', prest_nome='COM IRRF', valor=1000.0, base_iss=1000.0, iss=20.0,
               irrf=15.0, total_ret=15.0, liquido=985.0)]
    a = Apuracao(copy.deepcopy(CFG), '2026-09', leitores.ler_xmls([]), {}).executar()
    destino = os.path.join(tempfile.mkdtemp(), 's.xlsx')
    rows, sem, com = planilha_nfse.gerar(a, xs, destino)
    assert [x['numero'] for x in sem] == [1] and [x['numero'] for x in com] == [2] and com[0]['liquido'] == 985.0
    wb = load_workbook(destino)
    assert wb.sheetnames == ['Resumo', 'Todas as notas', 'Sem retenção', 'Com retenção']
    irrf = next(r for r in wb['Resumo'].iter_rows(values_only=True) if r[0] == 'IRRF')
    assert irrf[1] == 15.0 and 'DARF 1708' in irrf[3], irrf
    print('ok relatório de serviços tomados (todas / sem retenção / com retenção)')


if __name__ == '__main__':
    teste_simples()
    teste_apuracao()
    teste_irpj_trimestre()
    teste_txt()
    teste_sieg_cte()
    teste_revisao()
    teste_cte_natureza()
    teste_registro_erp()
    teste_difal_antecipacao()
    teste_contabil()
    teste_importacao()
    teste_servicos()
