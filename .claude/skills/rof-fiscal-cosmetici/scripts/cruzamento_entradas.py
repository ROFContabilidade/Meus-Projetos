"""Cruzamento das entradas do mês: XML (NF-e, CT-e, NFS-e) x relatórios SIEG da pasta (Detalhamento de Produtos, CT-e, NFS-e ABRASF).
Só compara - nada é alterado. Saída: Cruzamento_Entradas_Relatorios_x_XML_MM-AAAA.xlsx."""
import collections

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

from apuracao import vcontabil

R = lambda x: round((x or 0.0) + 0.0, 2)
H = PatternFill('solid', fgColor='1F3864')
OK = PatternFill('solid', fgColor='E2EFDA')
DIV = PatternFill('solid', fgColor='F8CBAD')
INFO = PatternFill('solid', fgColor='FFF2CC')
CINZA = PatternFill('solid', fgColor='EDEDED')
M = '#,##0.00;-#,##0.00;-'


def _cmp(pares, tol=0.01):
    """[(rótulo, valor XML, valor relatório)] -> lista de diferenças."""
    return [f'{rot}: XML {a:,.2f} x relatório {b:,.2f}' for rot, a, b in pares if a is not None and b is not None and abs(a - b) > tol]


def nfe(a, sieg):
    cnpj = a.cnpj
    xml = {n['chave']: n for n in a.entradas}
    canc = {n['chave']: n for n in a.canceladas if not n.get('propria')}
    rel = {k: n for k, n in (sieg or {}).items() if n['dest_cnpj'] == cnpj and n['emit_cnpj'] != cnpj}
    linhas = []
    for ch in sorted(set(xml) | set(canc) | set(rel), key=lambda c: ((xml.get(c) or canc.get(c) or rel.get(c))['dhemi'], c)):
        n, s = xml.get(ch) or canc.get(ch), rel.get(ch)
        base = n or s
        sit = 'Autorizada' if ch in xml else 'Cancelada' if ch in canc else 'Só no relatório'
        tx = n['tot'] if n else {}
        ts = s['tot'] if s else {}
        if sit == 'Cancelada':
            res = ('Cancelada - confere: evento de cancelamento no XML e status "Cancelamento" no relatório' if s and not s['autorizada']
                   else 'Cancelada no XML (evento)' + (' - relatório mostra autorizada' if s else ' - fora do relatório'))
        elif not s:
            res = 'Não está no relatório SIEG'
        elif not n:
            res = 'Sem XML na pasta'
        else:
            d = _cmp([('valor', tx.get('vNF'), ts.get('vNF')), ('ICMS', tx.get('vICMS'), ts.get('vICMS')),
                      ('IPI', tx.get('vIPI'), ts.get('vIPI')), ('ST', tx.get('vST'), ts.get('vST'))])
            if len(n['itens']) != len(s['itens']):
                d.append(f'itens: XML {len(n["itens"])} x relatório {len(s["itens"])}')
            res = 'OK' if not d else '; '.join(d)
        linhas.append([base['numero'], base['serie'], base['dhemi'], base['emit_nome'], base['emit_cnpj'], base.get('emit_uf') or '', sit,
                       '/'.join(sorted({i['cfop'] for i in n['itens']})) if n else '', len(n['itens']) if n else None,
                       len(s['itens']) if s else None, tx.get('vNF'), ts.get('vNF'), tx.get('vICMS'), ts.get('vICMS'),
                       tx.get('vIPI'), ts.get('vIPI'), tx.get('vST'), ts.get('vST'), res, ch])
    return linhas


def cte(a, xml_cte, sieg_cte):
    cnpj = a.cnpj
    xs = dict(xml_cte or {})
    rs = dict(sieg_cte or {})
    linhas = []
    for ch in sorted(set(xs) | set(rs), key=lambda c: ((xs.get(c) or rs.get(c))['data'] or a.ini, c)):
        x, r = xs.get(ch), rs.get(ch)
        b = x or r
        toma = (x or r)['toma_cnpj']
        papel = 'Cosmetici tomadora' if toma == cnpj else f'tomador {toma or "-"}'
        sit = 'Cancelado' if r and not r['autorizada'] else 'Autorizado'
        if sit == 'Cancelado':
            res = 'CT-e cancelado - status "Cancelamento" no relatório' + (' (XML na pasta)' if x else '')
        elif not x:
            res = 'Sem XML na pasta' + (' (Cosmetici tomadora - pedir o XML)' if toma == cnpj else ' (tomador é outra empresa)')
        elif not r:
            res = 'Não está no relatório SIEG'
        else:
            d = _cmp([('valor', x['vprest'], r['vprest']), ('base ICMS', x['vbc'], r['vbc']), ('ICMS', x['vicms'], r['vicms'])])
            if x['cfop'] != r['cfop']:
                d.append(f'CFOP XML {x["cfop"]} x relatório {r["cfop"]}')
            if x['toma_cnpj'] != r['toma_cnpj']:
                d.append(f'tomador XML {x["toma_cnpj"]} x relatório {r["toma_cnpj"]}')
            res = 'OK' if not d else '; '.join(d)
        nfes = (x or r)['chaves_nfe']
        linhas.append([b['numero'], b['serie'], b['data'], b['emit_nome'], b['emit_cnpj'], f'{b.get("uf_ini") or ""}-{b.get("uf_fim") or ""}',
                       b['cfop'], papel, sit, 'Sim' if x else 'Não', x['vprest'] if x else None, r['vprest'] if r else None,
                       x['vicms'] if x else None, r['vicms'] if r else None, ', '.join(str(int(k[25:34])) + '/' + k[6:20] for k in nfes if len(k) == 44),
                       res, ch])
    return linhas


def nfse(xml_nfse, rel_nfse):
    k = lambda cnpj, num: (cnpj, str(int(num)) if str(num).isdigit() else str(num))
    xs = {k(x['prest_cnpj'], x['numero']): x for x in xml_nfse or []}
    rs = {k(s['cnpj'], s['numero']): s for s in rel_nfse or []}
    linhas = []
    for key in sorted(set(xs) | set(rs), key=lambda kk: ((xs.get(kk) or {}).get('emissao') or '', kk)):
        x, s = xs.get(key), rs.get(key)
        if x and s:
            ret_x = R(x['total_ret'])
            ret_s = R(s['valor'] - s['liquido'] - s['desconto'])
            d = _cmp([('valor', x['valor'], s['valor']), ('ISS', x['iss'], s['iss']), ('IRRF', x['irrf'], s['ir']),
                      ('total retido', ret_x, ret_s), ('líquido', x['liquido'], s['liquido'])])
            if x['simples'] != s['simples']:
                d.append(f'Simples: XML {"sim" if x["simples"] else "não"} x relatório {"sim" if s["simples"] else "não"}')
            res = 'OK' if not d else '; '.join(d)
        else:
            res = 'Sem XML na pasta' if s else 'Não está no relatório ABRASF'
        b = x or {}
        linhas.append([key[1], b.get('emissao') or (s or {}).get('data'), b.get('prest_nome') or (s or {}).get('prestador'), key[0],
                       ('Sim' if b.get('simples') else 'Não') if x else ('Sim' if s['simples'] else 'Não'),
                       b.get('cod_trib_nac') or (s or {}).get('cod_servico'), b.get('servico') or '', b.get('mun_incidencia') or '',
                       b.get('valor'), (s or {}).get('valor'), b.get('base_iss'), b.get('iss'), (s or {}).get('iss'),
                       'Sim' if b.get('iss_retido') else 'Não', b.get('irrf'), (s or {}).get('ir'), b.get('pis'), b.get('cofins'),
                       b.get('csll'), b.get('inss'), b.get('total_ret'), R((s['valor'] - s['liquido'] - s['desconto'])) if s else None,
                       b.get('liquido'), (s or {}).get('liquido'), res, b.get('chave', '')])
    return linhas


def _aba(wb, titulo, sub, cab, larg, linhas, moedas, datas=(), cor=None, primeira=False):
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = titulo[:31]
    ws['A1'] = titulo; ws['A1'].font = Font(bold=True, size=12)
    ws['A2'] = sub; ws['A2'].font = Font(italic=True, size=9, color='595959')
    for c, (t, w) in enumerate(zip(cab, larg), 1):
        x = ws.cell(4, c, t); x.font = Font(bold=True, color='FFFFFF', size=9); x.fill = H
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[L(c)].width = w
    ws.row_dimensions[4].height = 30
    r = 5
    for ln in linhas:
        fl = cor(ln) if cor else None
        for c, v in enumerate(ln, 1):
            x = ws.cell(r, c, v); x.font = Font(size=9)
            if c in moedas: x.number_format = M
            if c in datas: x.number_format = 'DD/MM/YYYY'
            if fl: x.fill = fl
        r += 1
    if r > 5:
        ws.cell(r, 1, 'Total').font = Font(bold=True, size=9)
        for c in sorted(moedas):
            x = ws.cell(r, c, f'=SUM({L(c)}5:{L(c)}{r - 1})'); x.font = Font(bold=True, size=9); x.number_format = M
    ws.freeze_panes = 'D5'
    ws.auto_filter.ref = f'A4:{L(len(cab))}{max(r - 1, 4)}'
    return ws


def _cor(res_col, sit_col=None, sit_cinza=()):
    def f(ln):
        if sit_col is not None and ln[sit_col] in sit_cinza:
            return CINZA
        res = ln[res_col]
        if res == 'OK':
            return OK
        if res.startswith(('Sem XML', 'Não está')) and 'tomador é outra empresa' not in res:
            return INFO
        if 'outra empresa' in res or 'cancel' in res.lower():
            return CINZA
        return DIV
    return f


def gerar(a, sieg_prod, xml_cte, sieg_cte, xml_nfse, rel_nfse, destino):
    wb = Workbook()
    l_nfe, l_cte, l_nfse = nfe(a, sieg_prod), cte(a, xml_cte, sieg_cte), nfse(xml_nfse, rel_nfse)
    conta = lambda ls, col, pref='OK': sum(1 for x in ls if str(x[col]).startswith(pref))
    ws = wb.active; ws.title = 'Resumo'
    ws['A1'] = f'COSMETICI {a.mes[5:]}/{a.mes[:4]} - Entradas: XML x relatórios da pasta (SIEG)'; ws['A1'].font = Font(bold=True, size=12)
    ws['A2'] = 'Só conferência: os valores de cada lado estão exatamente como no XML e no relatório. Nada foi alterado.'
    ws['A2'].font = Font(italic=True, size=9, color='595959')
    cab = ['Documento', 'Relatório da pasta', 'Linhas', 'Conferem (OK)', 'Diferença de valor', 'Sem XML na pasta', 'Fora do relatório', 'Canceladas']
    for c, t in enumerate(cab, 1):
        x = ws.cell(4, c, t); x.font = Font(bold=True, color='FFFFFF', size=9); x.fill = H; x.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[L(c)].width = [12, 40, 8, 12, 14, 14, 14, 11][c - 1]
    def linha(r, doc, rel, ls, col_res):
        vals = [doc, rel, len(ls), conta(ls, col_res), sum(1 for x in ls if str(x[col_res]) != 'OK' and not str(x[col_res]).startswith(('Sem XML', 'Não está', 'Cancelada', 'CT-e cancelado'))),
                conta(ls, col_res, 'Sem XML'), conta(ls, col_res, 'Não está'), sum(1 for x in ls if 'ancel' in str(x[col_res]))]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v).font = Font(size=9)
    linha(5, 'NF-e', 'Relatorio_Detalhamento_Produtos.xlsx (entradas)', l_nfe, 18)
    linha(6, 'CT-e', 'Relatorio_CTe.xlsx', l_cte, 15)
    linha(7, 'NFS-e', 'RelatorioNFS_ABRASF_33039048000162_AAAA-MM.xlsx', l_nfse, 24)
    _aba(wb, 'NF-e entradas', 'XML x Relatorio_Detalhamento_Produtos (SIEG) - por chave. Verde = confere; vermelho = diferença; amarelo = falta de um lado; cinza = cancelada.',
         ['NF', 'Série', 'Emissão', 'Fornecedor', 'CNPJ', 'UF', 'Situação', 'CFOP (XML)', 'Itens XML', 'Itens SIEG', 'Valor XML', 'Valor SIEG',
          'ICMS XML', 'ICMS SIEG', 'IPI XML', 'IPI SIEG', 'ST XML', 'ST SIEG', 'Resultado', 'Chave'],
         [8, 5, 10, 32, 16, 4, 12, 10, 7, 7, 11, 11, 10, 10, 9, 9, 8, 8, 40, 46], l_nfe, set(range(11, 19)), {3},
         _cor(18, 6, ('Cancelada',)))
    _aba(wb, 'CT-e', 'XML x Relatorio_CTe (SIEG) - todos os CT-e em que a Cosmetici aparece. Amarelo = Cosmetici tomadora sem XML; cinza = tomador é outra empresa ou cancelado.',
         ['CT-e', 'Série', 'Emissão', 'Transportadora', 'CNPJ', 'Trajeto', 'CFOP', 'Tomador', 'Situação', 'XML na pasta', 'Valor XML', 'Valor SIEG',
          'ICMS XML', 'ICMS SIEG', 'NF-e transportada', 'Resultado', 'Chave'],
         [10, 5, 10, 32, 16, 7, 6, 26, 10, 8, 10, 10, 9, 9, 40, 46, 46], l_cte, {11, 12, 13, 14}, {3}, _cor(15))
    _aba(wb, 'NFS-e', 'XML (padrão nacional, pasta Doc\'s/Recebidas) x RelatorioNFS_ABRASF (SIEG). Valores como estão em cada documento.',
         ['NFS-e', 'Emissão', 'Prestador', 'CNPJ', 'Simples', 'Cód. serviço', 'Serviço', 'Município incidência', 'Valor XML', 'Valor relatório',
          'Base ISS', 'ISS XML', 'ISS relatório', 'ISS retido', 'IRRF XML', 'IRRF relatório', 'PIS (XML)', 'COFINS (XML)', 'CSLL retida (XML)',
          'INSS retido (XML)', 'Total retido XML', 'Total retido relatório', 'Líquido XML', 'Líquido relatório', 'Resultado', 'Chave'],
         [10, 10, 30, 16, 7, 8, 24, 16, 11, 11, 10, 9, 9, 7, 9, 9, 8, 9, 9, 9, 10, 10, 11, 11, 40, 50], l_nfse,
         {9, 10, 11, 12, 13, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24}, {2}, _cor(24))
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    return l_nfe, l_cte, l_nfse
