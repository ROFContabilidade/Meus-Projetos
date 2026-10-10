"""Leitura das fontes fiscais da COSMETICI.

- XML de NF-e (procNFe), eventos (cancelamento/CC-e/manifestação), inutilização e CT-e,
  soltos, em pastas ou dentro de ZIP (inclusive ZIP dentro de ZIP).
- Relatórios SIEG em xlsx: Relatorio_Detalhamento_Produtos (itens de NF-e) e Relatorio_CTe.
- SPED Fiscal (EFD ICMS/IPI) já transmitido, para conferência e para aprender a classificação
  das entradas por fornecedor.

Tudo devolve dicionários simples; os valores são float com 2 casas.
"""
import datetime as dt
import io
import os
import re
import zipfile
import xml.etree.ElementTree as ET

NS = '{http://www.portalfiscal.inf.br/nfe}'
NS_CTE = '{http://www.portalfiscal.inf.br/cte}'


def _f(x):
    if x is None or x == '':
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).replace('R$', '').strip()
    if ',' in s and '.' in s and s.rfind(',') > s.rfind('.'):
        s = s.replace('.', '').replace(',', '.')      # 1.234,56
    elif ',' in s and '.' not in s:
        s = s.replace(',', '.')                       # 1234,56
    else:
        s = s.replace(',', '')                        # 1,234.56
    try:
        return float(s)
    except ValueError:
        return 0.0


def _t(el, path, ns=NS):
    if el is None:
        return None
    x = el.find(path.replace('n:', ns))
    return x.text if x is not None else None


def _v(el, path, ns=NS):
    return _f(_t(el, path, ns))


def _data(s):
    if not s:
        return None
    return dt.date.fromisoformat(s[:10])


# ---------------------------------------------------------------- arquivos

def iter_xml(caminhos):
    """Gera (nome, bytes) de todo XML encontrado nos caminhos (arquivo, pasta ou ZIP)."""
    def de_zip(nome, dados):
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            for i in z.infolist():
                if i.is_dir():
                    continue
                b = z.read(i)
                n = f'{nome}!{i.filename}'
                if i.filename.lower().endswith('.zip'):
                    yield from de_zip(n, b)
                elif i.filename.lower().endswith('.xml'):
                    yield n, b

    for c in caminhos:
        if os.path.isdir(c):
            for raiz, _, arqs in os.walk(c):
                for a in sorted(arqs):
                    yield from iter_xml([os.path.join(raiz, a)])
        elif c.lower().endswith('.zip'):
            with open(c, 'rb') as fh:
                yield from de_zip(os.path.basename(c), fh.read())
        elif c.lower().endswith('.xml'):
            with open(c, 'rb') as fh:
                yield os.path.basename(c), fh.read()


# ---------------------------------------------------------------- NF-e

def _item_nfe(det, prot_ok):
    p = det.find(NS + 'prod')
    imp = det.find(NS + 'imposto')
    icms = imp.find(NS + 'ICMS') if imp is not None else None
    icms = icms[0] if icms is not None and len(icms) else None
    ipi = imp.find(NS + 'IPI') if imp is not None else None
    ipitrib = ipi.find(NS + 'IPITrib') if ipi is not None else None
    ipint = ipi.find(NS + 'IPINT') if ipi is not None else None
    pis = imp.find(NS + 'PIS') if imp is not None else None
    pis = pis[0] if pis is not None and len(pis) else None
    cof = imp.find(NS + 'COFINS') if imp is not None else None
    cof = cof[0] if cof is not None and len(cof) else None
    it = {
        'n_item': int(det.get('nItem')),
        'cprod': _t(p, 'n:cProd'), 'xprod': _t(p, 'n:xProd'), 'ncm': _t(p, 'n:NCM') or '',
        'cest': _t(p, 'n:CEST'), 'cfop': _t(p, 'n:CFOP'),
        'qtd': _v(p, 'n:qCom'), 'un': _t(p, 'n:uCom'),
        'vprod': _v(p, 'n:vProd'), 'vdesc': _v(p, 'n:vDesc'), 'vfrete': _v(p, 'n:vFrete'),
        'vseg': _v(p, 'n:vSeg'), 'voutro': _v(p, 'n:vOutro'),
        'orig': _t(icms, 'n:orig'), 'cst_icms': _t(icms, 'n:CST') or _t(icms, 'n:CSOSN'),
        'csosn': _t(icms, 'n:CSOSN'), 'pcredsn': _v(icms, 'n:pCredSN'), 'vcredsn': _v(icms, 'n:vCredICMSSN'),
        'vbc': _v(icms, 'n:vBC'), 'picms': _v(icms, 'n:pICMS'), 'vicms': _v(icms, 'n:vICMS'),
        'predbc': _v(icms, 'n:pRedBC'),
        'vbcst': _v(icms, 'n:vBCST'), 'pmvast': _v(icms, 'n:pMVAST'), 'picmsst': _v(icms, 'n:pICMSST'),
        'vst': _v(icms, 'n:vICMSST'), 'vfcpst': _v(icms, 'n:vFCPST'),
        'cst_ipi': _t(ipitrib, 'n:CST') or _t(ipint, 'n:CST'),
        'vbc_ipi': _v(ipitrib, 'n:vBC'), 'pipi': _v(ipitrib, 'n:pIPI'), 'vipi': _v(ipitrib, 'n:vIPI'),
        'vipidevol': _v(det, 'n:impostoDevol/n:IPI/n:vIPIDevol'),
        'cst_pis': _t(pis, 'n:CST'), 'vbc_pis': _v(pis, 'n:vBC'), 'ppis': _v(pis, 'n:pPIS'), 'vpis': _v(pis, 'n:vPIS'),
        'cst_cof': _t(cof, 'n:CST'), 'vbc_cof': _v(cof, 'n:vBC'), 'pcof': _v(cof, 'n:pCOFINS'), 'vcof': _v(cof, 'n:vCOFINS'),
    }
    return it


def ler_nfe(root, origem):
    inf = root.find(f'.//{NS}infNFe')
    if inf is None:
        return None
    ide = inf.find(NS + 'ide')
    emit = inf.find(NS + 'emit')
    dest = inf.find(NS + 'dest')
    prot = root.find(f'.//{NS}protNFe/{NS}infProt')
    cstat = _t(prot, 'n:cStat')
    nota = {
        'chave': inf.get('Id')[3:], 'origem': origem, 'modelo': _t(ide, 'n:mod'),
        'numero': int(_t(ide, 'n:nNF')), 'serie': _t(ide, 'n:serie'),
        'dhemi': _data(_t(ide, 'n:dhEmi') or _t(ide, 'n:dEmi')),
        'dhsaient': _data(_t(ide, 'n:dhSaiEnt') or _t(ide, 'n:dSaiEnt')),
        'tpnf': _t(ide, 'n:tpNF'), 'finnfe': _t(ide, 'n:finNFe'), 'natop': _t(ide, 'n:natOp'),
        'ind_final': _t(ide, 'n:indFinal'), 'id_dest': _t(ide, 'n:idDest'),
        'emit_cnpj': _t(emit, 'n:CNPJ') or _t(emit, 'n:CPF'), 'emit_nome': _t(emit, 'n:xNome'),
        'emit_uf': _t(emit, 'n:enderEmit/n:UF'), 'emit_crt': _t(emit, 'n:CRT'),
        'dest_cnpj': _t(dest, 'n:CNPJ') or _t(dest, 'n:CPF') or _t(dest, 'n:idEstrangeiro'),
        'dest_nome': _t(dest, 'n:xNome'), 'dest_uf': _t(dest, 'n:enderDest/n:UF'),
        'dest_ie': _t(dest, 'n:IE'), 'ind_ie_dest': _t(dest, 'n:indIEDest'),
        'cstat': cstat, 'autorizada': cstat in ('100', '150', None),
        'vnf': _v(inf, 'n:total/n:ICMSTot/n:vNF'),
        'tot': {k: _v(inf, f'n:total/n:ICMSTot/n:{k}') for k in
                ('vProd', 'vDesc', 'vFrete', 'vSeg', 'vOutro', 'vBC', 'vICMS', 'vST', 'vFCPST', 'vIPI', 'vIPIDevol', 'vPIS', 'vCOFINS', 'vNF')},
        'refnfe': [r.text for r in ide.findall(f'{NS}NFref/{NS}refNFe')],
        'infcpl': _t(inf, 'n:infAdic/n:infCpl') or '',
        'itens': [_item_nfe(d, True) for d in inf.findall(NS + 'det')],
    }
    return nota


def ler_evento(root, origem):
    out = []
    for ev in root.iter(NS + 'infEvento'):
        tp = _t(ev, 'n:tpEvento')
        ch = _t(ev, 'n:chNFe')
        if not tp or not ch:
            continue
        ret = root.find(f'.//{NS}retEvento/{NS}infEvento')
        cst = _t(ret, 'n:cStat') if ret is not None else None
        out.append({'chave': ch, 'tp': tp, 'desc': _t(ev, 'n:detEvento/n:descEvento'),
                    'data': _data(_t(ev, 'n:dhEvento')), 'cstat': cst, 'origem': origem,
                    'correcao': _t(ev, 'n:detEvento/n:xCorrecao')})
    return out


def ler_inut(root, origem):
    inf = root.find(f'.//{NS}inutNFe/{NS}infInut')
    if inf is None:
        inf = root.find(f'.//{NS}infInut')
    if inf is None:
        return None
    return {'serie': _t(inf, 'n:serie'), 'ini': int(_t(inf, 'n:nNFIni')), 'fim': int(_t(inf, 'n:nNFFin')),
            'just': _t(inf, 'n:xJust'), 'origem': origem}


def ler_cte_xml(root, origem):
    inf = root.find(f'.//{NS_CTE}infCte')
    if inf is None:
        return None
    t = lambda p: _t(inf, p, NS_CTE)
    v = lambda p: _v(inf, p, NS_CTE)
    toma = t('n:ide/n:toma3/n:toma') or t('n:ide/n:toma4/n:toma')
    papeis = {'0': 'n:rem', '1': 'n:exped', '2': 'n:receb', '3': 'n:dest'}
    if toma in papeis:
        toma_cnpj = t(papeis[toma] + '/n:CNPJ')
    else:
        toma_cnpj = t('n:ide/n:toma4/n:CNPJ')
    icms = inf.find(f'{NS_CTE}imp/{NS_CTE}ICMS')
    icms = icms[0] if icms is not None and len(icms) else None
    prot = root.find(f'.//{NS_CTE}protCTe/{NS_CTE}infProt')
    return {
        'chave': inf.get('Id')[3:], 'origem': origem, 'numero': int(t('n:ide/n:nCT')), 'serie': t('n:ide/n:serie'),
        'data': _data(t('n:ide/n:dhEmi')), 'cfop': t('n:ide/n:CFOP'),
        'emit_cnpj': t('n:emit/n:CNPJ'), 'emit_nome': t('n:emit/n:xNome'),
        'uf_ini': t('n:ide/n:UFIni'), 'uf_fim': t('n:ide/n:UFFim'),
        'toma_cnpj': toma_cnpj, 'rem_cnpj': t('n:rem/n:CNPJ'), 'dest_cnpj': t('n:dest/n:CNPJ'),
        'vprest': v('n:vPrest/n:vTPrest'), 'cst': _t(icms, 'n:CST', NS_CTE),
        'vbc': _v(icms, 'n:vBC', NS_CTE), 'picms': _v(icms, 'n:pICMS', NS_CTE), 'vicms': _v(icms, 'n:vICMS', NS_CTE),
        'autorizada': _t(prot, 'n:cStat', NS_CTE) in ('100', None),
        'chaves_nfe': [x.text for x in inf.iter(NS_CTE + 'chave')],
    }


def ler_xmls(caminhos):
    """Lê todos os XML e devolve {'nfe': {chave: nota}, 'eventos': [...], 'inut': [...], 'cte': {...}, 'outros': [...]}."""
    r = {'nfe': {}, 'eventos': [], 'inut': [], 'cte': {}, 'outros': [], 'erros': []}
    for nome, b in iter_xml(caminhos):
        try:
            root = ET.fromstring(b)
        except ET.ParseError as e:
            r['erros'].append((nome, str(e)))
            continue
        tag = root.tag.split('}')[-1]
        if tag in ('nfeProc', 'NFe'):
            n = ler_nfe(root, nome)
            if n:
                antigo = r['nfe'].get(n['chave'])
                if antigo is None or (antigo['cstat'] is None and n['cstat']):
                    r['nfe'][n['chave']] = n
        elif tag in ('procEventoNFe', 'evento', 'envEvento', 'retEnvEvento'):
            r['eventos'].extend(ler_evento(root, nome))
        elif tag in ('procInutNFe', 'inutNFe', 'retInutNFe'):
            i = ler_inut(root, nome)
            if i:
                r['inut'].append(i)
        elif tag in ('cteProc', 'CTe'):
            c = ler_cte_xml(root, nome)
            if c:
                r['cte'][c['chave']] = c
        else:
            r['outros'].append((nome, tag))
    # eventos repetidos (mesmo XML em mais de um ZIP)
    vistos, ev = set(), []
    for e in r['eventos']:
        k = (e['chave'], e['tp'], e['data'], e['correcao'])
        if k not in vistos:
            vistos.add(k)
            ev.append(e)
    r['eventos'] = ev
    inut = {}
    for i in r['inut']:
        inut.setdefault((str(int(i['serie'])), i['ini'], i['fim']), i)
    r['inut'] = list(inut.values())
    return r


# ---------------------------------------------------------------- relatórios SIEG (xlsx)

def _linhas_xlsx(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    for i, l in enumerate(linhas[:10]):
        if l and any(isinstance(c, str) and c.strip() in ('Chave', 'Numero') for c in l):
            cab = [str(c).strip() if c is not None else '' for c in l]
            return [dict(zip(cab, x)) for x in linhas[i + 1:] if x and any(c not in (None, '') for c in x)]
    raise ValueError(f'{caminho}: cabeçalho do relatório SIEG não encontrado')


def ler_sieg_produtos(caminho):
    """Relatorio_Detalhamento_Produtos.xlsx (SIEG) -> notas no mesmo formato de ler_nfe."""
    notas = {}
    for l in _linhas_xlsx(caminho):
        ch = str(l.get('Chave') or '').strip()
        if len(ch) != 44:
            continue
        d = l.get('Dt_Emissao')
        d = d.date() if isinstance(d, dt.datetime) else (dt.datetime.strptime(str(d)[:10], '%d/%m/%Y').date() if d else None)
        status = str(l.get('Status') or '')
        n = notas.get(ch)
        if n is None:
            fin = {'Normal': '1', 'Complementar': '2', 'Ajuste': '3', 'Devolução': '4', 'Devolucao': '4'}.get(str(l.get('Finalidade')), '1')
            n = notas[ch] = {
                'chave': ch, 'origem': os.path.basename(caminho), 'modelo': ch[20:22],
                'numero': int(float(l.get('Numero'))), 'serie': str(int(float(l.get('Serie') or 0))),
                'dhemi': d, 'dhsaient': d, 'tpnf': '1' if str(l.get('Tp_NF', '')).startswith('Sa') else '0',
                'finnfe': fin, 'natop': None,
                'emit_cnpj': str(l.get('CNPJ_CPF_Emit') or ''), 'emit_nome': l.get('Rz_Emit'), 'emit_uf': l.get('UF_Emit'),
                'emit_crt': None, 'dest_cnpj': str(l.get('CNPJ_CPF_Dest') or ''), 'dest_nome': l.get('Rz_Dest'),
                'dest_uf': l.get('UF_Dest'), 'dest_ie': l.get('IE_Dest'), 'ind_ie_dest': None,
                'cstat': '101' if 'ancel' in status else '100', 'autorizada': 'ancel' not in status,
                'vnf': _f(l.get('Valor_Total_Nota')), 'tot': {}, 'refnfe': [], 'itens': []}
        g = lambda k: _f(l.get(k))
        s = lambda k: (str(l.get(k)).strip() if l.get(k) not in (None, '') else None)
        n['itens'].append({
            'n_item': len(n['itens']) + 1, 'cprod': None, 'xprod': s('Produto'), 'ncm': str(l.get('NCM') or ''),
            'cest': s('CEST'), 'cfop': str(l.get('CFOP') or ''), 'qtd': g('Quantidade'), 'un': s('Unidade_Comercial'),
            'vprod': g('Valor_Produto'), 'vdesc': g('Desconto'), 'vfrete': g('Valor_Frete'), 'vseg': 0.0, 'voutro': g('Valor_Outro'),
            'orig': s('Origem'), 'cst_icms': s('ICMS_CST') or s('CSOSN'), 'csosn': s('CSOSN'), 'pcredsn': 0.0, 'vcredsn': 0.0, 'vbc': g('ICMS_Base_Calculo'),
            'picms': g('ICMS_Percentual'), 'vicms': g('Valor_ICMS'), 'predbc': 0.0,
            'vbcst': g('Base_Calculo_ST'), 'pmvast': g('MVA_ST_Percentual'), 'picmsst': g('ICMS_ST_Percentual'),
            'vst': g('ICMS_ST_Valor'), 'vfcpst': 0.0,
            'cst_ipi': s('IPI_CST') or s('IPI_CST_Nao_Trib'), 'vbc_ipi': g('IPI_Base_Calculo'),
            'pipi': g('IPI_Aliquota_Percentual'), 'vipi': g('Valor_IPI'), 'vipidevol': 0.0,
            'cst_pis': s('PIS_CST'), 'vbc_pis': g('PIS_Base_Calculo'), 'ppis': g('PIS_Percentual'), 'vpis': g('PIS_Valor'),
            'cst_cof': s('COFINS_CST'), 'vbc_cof': g('COFINS_Base_Calculo'), 'pcof': g('COFINS_Percentual'), 'vcof': g('COFINS_Valor'),
        })
    for n in notas.values():
        n['tot'] = {'vICMS': round(sum(i['vicms'] for i in n['itens']), 2), 'vST': round(sum(i['vst'] for i in n['itens']), 2),
                    'vIPI': round(sum(i['vipi'] for i in n['itens']), 2), 'vNF': n['vnf']}
    return notas


def ler_sieg_cte(caminho):
    """Relatorio_CTe.xlsx (SIEG) -> CT-e no mesmo formato de ler_cte_xml."""
    out = {}
    for l in _linhas_xlsx(caminho):
        ch = str(l.get('Chave') or '').strip()
        if len(ch) != 44:
            continue
        d = l.get('Dt_Emissao')
        d = d.date() if isinstance(d, dt.datetime) else (dt.datetime.strptime(str(d)[:10], '%d/%m/%Y').date() if d else None)
        status = str(l.get('Status') or '')
        toma = str(l.get('CNPJ_CPF_Tomador') or '') or str(l.get('CNPJ_CPF_Outro_Tomador') or '')
        out[ch] = {
            'chave': ch, 'origem': os.path.basename(caminho), 'numero': int(float(l.get('Numero'))), 'serie': str(l.get('Serie')),
            'data': d, 'cfop': str(l.get('CFOP') or ''), 'emit_cnpj': str(l.get('CNPJ_CPF_Emitente') or ''),
            'emit_nome': l.get('Rz_Emit'), 'uf_ini': l.get('Uf_Inicio'), 'uf_fim': l.get('Uf_Final'),
            'toma_cnpj': toma, 'rem_cnpj': str(l.get('CNPJ_CPF_Rem') or ''), 'dest_cnpj': str(l.get('CNPJ_CPF_Dest') or ''),
            'vprest': _f(l.get('Valor_Transporte')), 'cst': str(l.get('CST') or ''), 'vbc': _f(l.get('ICMS_Base_Calculo')),
            'picms': _f(l.get('ICMS_Aliq')), 'vicms': _f(l.get('ICMS')),
            'autorizada': 'ancel' not in status, 'chaves_nfe': [str(l.get('Chave_NFe') or '')],
        }
    return out


# ---------------------------------------------------------------- SPED Fiscal

def ler_sped(caminho):
    """EFD ICMS/IPI -> dicionário com participantes, documentos (C100/C190, D100/D190) e apuração (E110, E210, E520)."""
    with open(caminho, 'rb') as fh:
        bruto = fh.read()
    texto = bruto.decode('latin-1')
    r = {'part': {}, 'itens': {}, 'docs': [], 'cte': [], 'e110': None, 'e111': [], 'e116': [], 'e210': {}, 'e250': [],
         'e510': [], 'e520': None, 'periodo': None}
    doc = None
    uf_st = None
    for linha in texto.splitlines():
        if not linha.startswith('|'):
            continue
        c = linha.split('|')
        reg = c[1]
        if reg == '9999':
            break
        if reg == '0000':
            r['periodo'] = (c[4], c[5])
            r['cnpj'] = c[7]
        elif reg == '0150':
            r['part'][c[2]] = {'nome': c[3], 'cnpj': c[5] or c[6]}
        elif reg == '0200':
            r['itens'][c[2]] = {'descr': c[3], 'ncm': c[8]}
        elif reg == 'C100':
            doc = {'ind_oper': c[2], 'ind_emit': c[3], 'cod_part': c[4], 'mod': c[5], 'cod_sit': c[6], 'serie': c[7],
                   'numero': c[8], 'chave': c[9], 'dt_doc': c[10], 'dt_es': c[11], 'vl_doc': _f(c[12]), 'c170': [], 'c190': []}
            r['docs'].append(doc)
        elif reg == 'C170' and doc is not None:
            doc['c170'].append({'cod_item': c[3], 'vl_item': _f(c[7]), 'cst_icms': c[10], 'cfop': c[11],
                                'vl_icms': _f(c[15]), 'vl_ipi': _f(c[24]) if len(c) > 24 else 0.0})
        elif reg == 'C190' and doc is not None:
            doc['c190'].append({'cst': c[2], 'cfop': c[3], 'aliq': _f(c[4]), 'vl_opr': _f(c[5]), 'vl_bc': _f(c[6]),
                                'vl_icms': _f(c[7]), 'vl_bc_st': _f(c[8]), 'vl_st': _f(c[9]), 'vl_red_bc': _f(c[10]),
                                'vl_ipi': _f(c[11])})
        elif reg == 'D100':
            doc = {'ind_oper': c[2], 'ind_emit': c[3], 'cod_part': c[4], 'mod': c[5], 'cod_sit': c[6], 'serie': c[7],
                   'numero': c[9], 'chave': c[10], 'dt_doc': c[11], 'dt_es': c[12], 'vl_doc': _f(c[15]), 'd190': []}
            r['cte'].append(doc)
        elif reg == 'D190' and doc is not None and 'd190' in doc:
            doc['d190'].append({'cst': c[2], 'cfop': c[3], 'aliq': _f(c[4]), 'vl_opr': _f(c[5]), 'vl_bc': _f(c[6]),
                                'vl_icms': _f(c[7])})
        elif reg == 'E110':
            k = ['tot_debitos', 'aj_debitos', 'tot_aj_debitos', 'estornos_cred', 'tot_creditos', 'aj_creditos',
                 'tot_aj_creditos', 'estornos_deb', 'sld_credor_ant', 'sld_apurado', 'tot_ded', 'icms_recolher',
                 'sld_credor_transportar', 'deb_esp']
            r['e110'] = dict(zip(k, map(_f, c[2:16])))
        elif reg == 'E111':
            r['e111'].append({'cod': c[2], 'descr': c[3], 'valor': _f(c[4])})
        elif reg == 'E116':
            r['e116'].append({'cod_or': c[2], 'valor': _f(c[3]), 'venc': c[4], 'cod_rec': c[5]})
        elif reg == 'E200':
            uf_st = c[2]
        elif reg == 'E210':
            k = ['ind_mov', 'sld_cred_ant', 'devol_st', 'ressarc_st', 'outros_cred', 'aj_creditos', 'retencao_st',
                 'outros_deb', 'aj_debitos', 'sld_dev_ant_ded', 'deducoes', 'icms_recol_st', 'sld_cred_transportar', 'deb_esp']
            r['e210'][uf_st] = dict(zip(k, [c[2]] + list(map(_f, c[3:16]))))
        elif reg == 'E250':
            r['e250'].append({'cod_or': c[2], 'valor': _f(c[3]), 'venc': c[4], 'cod_rec': c[5]})
        elif reg == 'E510':
            r['e510'].append({'cfop': c[2], 'cst_ipi': c[3], 'vl_cont': _f(c[4]), 'vl_bc_ipi': _f(c[5]), 'vl_ipi': _f(c[6])})
        elif reg == 'E520':
            k = ['sd_ant', 'debitos', 'creditos', 'od', 'oc', 'sc', 'sd']
            r['e520'] = dict(zip(k, map(_f, c[2:9])))
    for d in r['docs'] + r['cte']:
        p = r['part'].get(d['cod_part'], {})
        d['cnpj'] = p.get('cnpj')
        d['nome'] = p.get('nome')
    return r


def cfop_entrada(cfop_fornecedor):
    """CFOP de saída do fornecedor (5xxx/6xxx/7xxx) -> CFOP de entrada correspondente (1xxx/2xxx/3xxx)."""
    c = str(cfop_fornecedor)
    return {'5': '1', '6': '2', '7': '3'}.get(c[:1], c[:1]) + c[1:]


def ler_relatorio_faturamento_erp(caminho):
    """Relatório de Faturamento do ERP do cliente (Cosmos, PDF ou texto do pdftotext -layout) -> {numero: dados}."""
    if caminho.lower().endswith('.pdf'):
        import subprocess
        texto = subprocess.run(['pdftotext', '-layout', caminho, '-'], capture_output=True, text=True, check=True).stdout
    else:
        texto = open(caminho, encoding='utf-8', errors='ignore').read()
    num = r'(-?[\d.]+,\d{2})'
    rx = re.compile(r'^\s*\d+\s+(\d+)\s+(\d+)\s+(\d{2}/\d{2}/\d{2})\s+(\w+)\s+(.*?)\s+([A-Z])\s+' + r'\s+'.join([num] * 10) + r'\s*$')
    out = {}
    for linha in texto.splitlines():
        m = rx.match(linha)
        if not m:
            continue
        g = m.groups()
        campos = ['vprod', 'vfrete', 'vseg', 'vdesc', 'voutro', 'bc', 'icms', 'st', 'ipi', 'vnf']
        d = {'serie': g[0], 'numero': int(g[1]), 'data': dt.datetime.strptime(g[2], '%d/%m/%y').date(), 'tipo': g[3],
             'pessoa': g[4].strip(), 'situacao': g[5]}
        d.update({k: _f(v) for k, v in zip(campos, g[6:6 + len(campos)])})
        out[d['numero']] = d
    return out


def _moeda(x):
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    x = str(x).replace('R$', '').strip()
    if x in ('', '-') or x.startswith('#'):
        return 0.0
    if ',' in x and '.' in x:
        x = x.replace(',', '') if x.rfind('.') > x.rfind(',') else x.replace('.', '').replace(',', '.')
    elif ',' in x:
        x = x.replace(',', '.')
    try:
        return float(x)
    except ValueError:
        return 0.0


def ler_relatorio_nfse_abrasf(caminho):
    """RelatorioNFS_ABRASF_<cnpj>_<AAAA-MM>.xlsx (SIEG): NFS-e tomadas, com valores e retenções."""
    import openpyxl
    ws = openpyxl.load_workbook(caminho, read_only=True, data_only=True).worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    cab = None
    out = []
    for row in linhas:
        vals = ['' if v is None else v for v in row]
        if cab is None:
            if 'Numero' in [str(v).strip() for v in vals] and 'Valor_Servico' in [str(v).strip() for v in vals]:
                cab = {str(v).strip(): k for k, v in enumerate(vals)}
            continue
        g = lambda c: vals[cab[c]] if c in cab and cab[c] < len(vals) else ''
        if not str(g('Numero')).strip() or not str(g('Prestador')).strip() or str(g('Prestador')).startswith('#'):
            continue
        base = g('Base_Calculo')
        iss_ret = g('ISSQN')
        out.append({
            'numero': str(g('Numero')).strip(), 'data': str(g('Dt_Emissao'))[:10], 'cod_servico': re.sub(r'\D', '', str(g('Cod_Servico'))).zfill(4)[:4],
            'prestador': str(g('RzPrestador')).strip(), 'cnpj': re.sub(r'\D', '', str(g('Prestador'))),
            'uf': str(g('UF_Prest')), 'simples': str(g('Optante_SN')).strip() == '1',
            'cancelada': bool(str(g('Dt_Cancelamento')).strip()) or 'cancel' in str(g('Status')).lower(),
            'valor': _moeda(g('Valor_Servico')), 'base': None if base in ('', None) else _moeda(base),
            'liquido': _moeda(g('Valor_Liquido')), 'desconto': _moeda(g('Desconto_Incondic')),
            'iss': _moeda(g('ISS')), 'iss_retido': _moeda(iss_ret),
            'ir': _moeda(g('IR')), 'pis': _moeda(g('PIS')), 'cofins': _moeda(g('COFINS')), 'outras': _moeda(g('OutRetencoes')),
            'descricao': str(g('Descriminacao'))[:200],
        })
    return out


def ler_registro_saidas_erp(caminho_ou_texto):
    """Registro de Saídas do ERP Cosmos (PDF "Saidas Cosmetici.pdf" ou o texto dele).
    Cada nota: NF <série> <número> <dia> <UF> <vlr contábil> <CFOP> ICMS <base> <alíq> <imposto> <isentas> <outras>
    <CFOP> IPI <base> <alíq> <imposto> <isentas> <outras>."""
    txt = caminho_ou_texto
    if os.path.exists(str(caminho_ou_texto)):
        if str(caminho_ou_texto).lower().endswith('.pdf'):
            import subprocess
            txt = subprocess.run(['pdftotext', '-raw', caminho_ou_texto, '-'], capture_output=True, text=True).stdout
        else:
            with open(caminho_ou_texto, encoding='utf-8', errors='replace') as fh:
                txt = fh.read()
    txt = re.sub(r'\s+', ' ', txt)
    num = r'(-?[\d\.]+,\d{2})'
    pat = re.compile(r'NF \d+ (\d+) (\d{1,2}) ([A-Z]{2}) ' + num + r' (\d{4}) ICMS ' + ' '.join([num] * 5)
                     + r' (\d{4}) IPI ' + ' '.join([num] * 5))
    f = lambda x: float(x.replace('.', '').replace(',', '.'))
    out = []
    for m in pat.finditer(txt):
        g = m.groups()
        out.append({'numero': int(g[0]), 'dia': int(g[1]), 'uf': g[2], 'vc': f(g[3]), 'cfop': g[4],
                    'bc_icms': f(g[5]), 'aliq': f(g[6]), 'icms': f(g[7]), 'isentas_icms': f(g[8]), 'outras_icms': f(g[9]),
                    'bc_ipi': f(g[11]), 'aliq_ipi': f(g[12]), 'ipi': f(g[13]), 'isentas_ipi': f(g[14]), 'outras_ipi': f(g[15])})
    for m in re.finditer(r'NF \d+ (\d+) (\d{1,2}) ([A-Z]{2}) NOTA FISCAL CANCELADA', txt):
        out.append({'numero': int(m.group(1)), 'dia': int(m.group(2)), 'uf': m.group(3), 'cancelada': True, 'vc': 0.0, 'cfop': None})
    return out
