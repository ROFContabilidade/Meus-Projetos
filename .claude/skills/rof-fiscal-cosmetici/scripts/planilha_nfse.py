"""Relatório das NFS-e de serviços tomados do mês (Relacao_Servicos_Tomados_Cosmetici_MM-AAAA.xlsx) e a base de dados do mês
(dados/nfse_tomadas_MM-AAAA.json). Abas: Resumo (totais e retenções a recolher), Todas as notas, Sem retenção, Com retenção.
Valores e retenções EXATAMENTE como estão no XML da NFS-e (padrão nacional) - nada é recalculado."""
import datetime as dt
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import planilha_entradas as pe

F = pe.F
H = pe.HEAD
RET = pe.RET
TOT = pe.TOT
M = pe.MOEDA

# guia de cada retenção (informativo, para conferência do recolhimento)
GUIAS = [('ISS retido', 'iss_retido', 'Guia do ISS retido do município de incidência (tomador recolhe)', 'conforme o calendário do município'),
         ('IRRF', 'irrf', 'DARF 1708 - IRRF sobre serviços de PJ', 'até o dia 20 do mês seguinte (último dia útil do 2º decêndio)'),
         ('PIS/COFINS/CSLL (CSRF)', 'csll', 'DARF 5952 - PIS/COFINS/CSLL retidos (4,65%)', 'até o dia 20 do mês seguinte (último dia útil do 2º decêndio)'),
         ('INSS retido', 'inss', 'DCTFWeb / DARF previdenciário - retenção de 11% (cessão de mão de obra)', 'até o dia 20 do mês seguinte')]

COLS = [('Emissão', 10, 'data', 'd'), ('Competência', 10, 'competencia', 'd'), ('NFS-e', 9, 'numero', ''), ('Situação', 10, 'situacao', ''), ('Prestador', 32, 'emitente', ''),
        ('CNPJ', 15, 'cnpj', ''), ('Simples', 7, 'simples', ''), ('Município incidência', 12, 'municipio', ''), ('Cód. serviço', 8, 'cod', ''),
        ('Descrição do serviço', 40, 'descricao', ''), ('Valor do serviço', 12, 'vNF', 'm'), ('Base ISS', 10, 'base_iss', 'm'),
        ('% ISS', 6, 'aliq_iss', ''), ('ISS do prestador', 10, 'iss', 'm'), ('ISS retido', 9, 'iss_retido', 'm'), ('IRRF retido', 9, 'irrf', 'm'),
        ('PIS/COFINS/CSLL retidos (vRetCSLL)', 12, 'csll', 'm'), ('INSS retido', 9, 'inss', 'm'), ('Total retido (vTotalRet)', 11, 'total_ret', 'm'),
        ('Líquido a pagar ao prestador', 12, 'liquido', 'm'), ('Retenção', 20, 'tem_ret', ''), ('Retenção PIS/COFINS (XML)', 22, 'tp_pc', ''),
        ('Acumulador (Domínio)', 10, 'acumulador', ''), ('Conta contábil (Domínio)', 14, 'conta', ''),
        ('Conferência XML x relatório ABRASF', 26, 'conferencia', ''), ('Chave', 50, 'chave', '')]


def linhas(a, xml_nfse, rel_nfse=None, forn=None):
    import cruzamento_entradas as ce
    import lancamentos_contabeis as lc
    from creditos import Credito
    rows = pe.nfse(a.cfg, xml_nfse)
    por = {(x['prest_cnpj'], str(int(x['numero'])) if str(x['numero']).isdigit() else str(x['numero'])): x for x in xml_nfse or []}
    conf = {(ln[3], str(ln[0])): ln[24] for ln in ce.nfse(xml_nfse, rel_nfse)} if rel_nfse else {}
    cred = Credito(a, forn or {}, [], a.cfg.get('natureza_entradas', {})).analisar()
    lanc, pend = lc.montar(a, cred, {}, {}, xml_nfse, forn or {})
    conta = {}
    for x in lanc:
        if x['doc'] == 'NFS-e':
            d = next(ln for _, b in x['blocos'] for ln in b if ln[0])
            conta[(x['cnpj'], str(x['numero']))] = f'D {d[0]} / C {d[1] or "506"}'
    for x in pend:
        if x['doc'] == 'NFS-e':
            conta[(x['cnpj'], str(x['numero']))] = 'definir: ' + x['motivo']
    for r in rows:
        k = (r['cnpj'], str(r['numero']))
        x = por.get(k) or {}
        comp = x.get('competencia')
        r['competencia'] = dt.date.fromisoformat(str(comp)[:10]) if comp else None
        r['data'] = dt.date.fromisoformat(str(r['data'])[:10]) if r['data'] and not isinstance(r['data'], dt.date) else r['data']
        r['descricao'] = (x.get('descricao') or r.get('servico') or '')[:250]
        r['conta'] = conta.get(k, '')
        r['conferencia'] = conf.get(k, 'sem relatório ABRASF' if not rel_nfse else 'não está no relatório')
    return rows


def gerar(a, xml_nfse, destino, rel_nfse=None, forn=None, base_json=None):
    rows = linhas(a, xml_nfse, rel_nfse, forn)
    com = [r for r in rows if r['tem_ret'].startswith('SIM')]
    sem = [r for r in rows if not r['tem_ret'].startswith('SIM')]
    mm = f'{a.mes[5:]}/{a.mes[:4]}'
    emp = a.cfg['empresa']['nome']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumo'
    ws['A1'] = f'{emp} - Serviços tomados {mm}'; ws['A1'].font = F(bold=True, size=12)
    ws['A2'] = 'Valores e retenções exatamente como estão no XML de cada NFS-e (padrão nacional). Nada foi recalculado.'
    ws['A2'].font = F(italic=True, color='595959')
    r = 4
    for rot, v in (('Notas no mês', len(rows)), ('Sem retenção', len(sem)), ('Com retenção', len(com))):
        ws.cell(r, 1, rot).font = F(bold=True); ws.cell(r, 2, v).font = F(); r += 1
    soma = lambda ls, k: round(sum(x.get(k) or 0.0 for x in ls), 2)
    r += 1
    for i, c in enumerate(['', 'Todas', 'Sem retenção', 'Com retenção'], 1):
        x = ws.cell(r, i, c); x.font = F(bold=True, color='FFFFFF'); x.fill = H
    r += 1
    for rot, k in (('Valor dos serviços', 'vNF'), ('ISS do prestador', 'iss'), ('Total retido', 'total_ret'), ('Líquido a pagar', 'liquido')):
        ws.cell(r, 1, rot).font = F(bold=True)
        for i, ls in enumerate((rows, sem, com), 2):
            x = ws.cell(r, i, soma(ls, k)); x.font = F(); x.number_format = M
        r += 1
    r += 1
    for i, c in enumerate(['Retenção a recolher', 'Valor', 'Notas', 'Guia', 'Vencimento (informativo)'], 1):
        x = ws.cell(r, i, c); x.font = F(bold=True, color='FFFFFF'); x.fill = H
    r += 1
    for rot, k, guia, venc in GUIAS:
        notas = [x for x in com if x.get(k)]
        vals = [rot, soma(notas, k), ', '.join(f'{x["numero"]} {x["emitente"][:20]} ({x[k]:,.2f})' for x in notas) or '-', guia, venc]
        for i, v in enumerate(vals, 1):
            x = ws.cell(r, i, v); x.font = F(); x.alignment = Alignment(wrap_text=True, vertical='top')
            if i == 2: x.number_format = M
            if notas: x.fill = RET
        r += 1
    ws.cell(r + 1, 1, 'A soma das retenções por tributo deve bater com o "Total retido" das notas com retenção.').font = F(italic=True, color='595959')
    for col, w in zip('ABCDE', [26, 14, 60, 48, 40]):
        ws.column_dimensions[col].width = w
    cor = lambda ln: RET if ln['tem_ret'].startswith('SIM') else None
    pe._aba(wb, 'Todas as notas', f'NFS-e de serviços tomados {mm}, na ordem da emissão. Amarelo = nota com retenção.', COLS, rows, cor,
            contar=lambda i, u: [('Com retenção', f'=COUNTIF(U{i}:U{u},"SIM*")'), ('Sem retenção', f'=COUNTIF(U{i}:U{u},"Não")')])
    pe._aba(wb, 'Sem retenção', f'NFS-e sem nenhuma retenção no XML - {mm}. O valor integral é pago ao prestador.', COLS, sem)
    pe._aba(wb, 'Com retenção', f'NFS-e com retenção no XML - {mm}. O tomador (Cosmetici) recolhe o valor retido; paga ao prestador o líquido.',
            COLS, com, lambda ln: RET)
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    if base_json:
        with open(base_json, 'w', encoding='utf-8') as fh:
            json.dump({'mes': a.mes, 'gerado_de': 'XML das NFS-e (padrão nacional) da pasta Doc\'s/Recebidas',
                       'notas': [{k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in x.items() if k != 'tipo'} for x in rows]},
                      fh, ensure_ascii=False, indent=1)
    return rows, sem, com
