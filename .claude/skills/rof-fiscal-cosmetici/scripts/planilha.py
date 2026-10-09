"""Planilha da apuração fiscal mensal da COSMETICI (Apuracao_Fiscal_Cosmetici_MM-AAAA.xlsx)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apuracao import vcontabil

F = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k})
thin = Side(style='thin', color='BFBFBF')
B = Border(top=thin, bottom=thin, left=thin, right=thin)
HEAD = PatternFill('solid', fgColor='D9E1F2')
EDIT = PatternFill('solid', fgColor='FFF2CC')
ALTA = PatternFill('solid', fgColor='F8CBAD')
MEDIA = PatternFill('solid', fgColor='FCE4D6')
OK = PatternFill('solid', fgColor='E2EFDA')
MOEDA = '#,##0.00;-#,##0.00;-'
PCT = '0.00%'


def _cab(ws, row, headers, col=1):
    for c, h in enumerate(headers, col):
        x = ws.cell(row, c, h)
        x.font = F(bold=True); x.fill = HEAD; x.border = B
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _linha(ws, row, valores, col=1, moeda=(), bold=False, fill=None):
    for c, v in enumerate(valores, col):
        x = ws.cell(row, c, v)
        x.font = F(bold=bold); x.border = B
        if fill:
            x.fill = fill
        if (c - col) in moeda or (isinstance(v, float) and not moeda):
            x.number_format = MOEDA
        if hasattr(v, 'year'):
            x.number_format = 'DD/MM/YYYY'


def _larguras(ws, larguras):
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _titulo(ws, texto, sub=None):
    ws['A1'] = texto; ws['A1'].font = F(bold=True, size=12)
    if sub:
        ws['A2'] = sub; ws['A2'].font = F(italic=True, color='666666')


def _bloco(ws, r, titulo):
    ws.cell(r, 1, titulo).font = F(bold=True, size=11, color='1F3864')
    return r + 1


def resumo(wb, a):
    ws = wb.active; ws.title = 'Resumo'
    emp = a.cfg['empresa']
    _titulo(ws, f'{emp["nome"]} - CNPJ {emp["cnpj"]} - Apuração fiscal {a.mes[5:]}/{a.mes[:4]}',
            f'Regime: {emp["regime"]} - Empresa {emp["codigo_dominio"]} no Domínio. Gerado a partir dos XML/relatórios do mês; '
            'conferir as abas Divergências e Observações antes de emitir as guias.')
    r = _bloco(ws, 4, 'Guias do mês')
    _cab(ws, r, ['Imposto', 'Guia / código', 'Valor', 'Vencimento']); r += 1
    for nome, guia, valor, venc in a.guias():
        _linha(ws, r, [nome, guia, valor, venc]); r += 1

    r = _bloco(ws, r + 1, 'Faturamento do mês (base IRPJ/CSLL)')
    f = a.fat
    for rot, v in [('Vendas (acumuladores de venda)', f['vendas']), ('Industrialização por encomenda (5124/6124)', f['industrializacao']),
                   ('(-) Devoluções de venda', -f['devolucoes']), ('(-) ICMS-ST', -f['st']), ('(-) IPI', -f['ipi'])]:
        _linha(ws, r, [rot, None, v]); r += 1
    _linha(ws, r, ['Faturamento líquido', None, f'=SUM(C{r-5}:C{r-1})'], bold=True); ws.cell(r, 3).number_format = MOEDA; r += 1

    r = _bloco(ws, r + 1, 'PIS / COFINS (fórmula da planilha FAT da Cosmetici)')
    _cab(ws, r, ['Base', 'Valor da base', 'PIS %', 'PIS', 'COFINS %', 'COFINS']); r += 1
    pc, cp = a.pc, a.cfg['pis_cofins']
    l1 = r
    _linha(ws, r, ['Industrialização (vlr contábil - ICMS)', pc['base_industrializacao'], cp['aliquotas_industrializacao']['pis'],
                   f'=ROUND(B{r}*C{r},2)', cp['aliquotas_industrializacao']['cofins'], f'=ROUND(B{r}*E{r},2)'], moeda=(1, 3, 5)); r += 1
    _linha(ws, r, ['Vendas monofásicas (vlr contábil - ICMS - IPI - ST)', pc['base_vendas'], cp['aliquotas_vendas']['pis'],
                   f'=ROUND(B{r}*C{r},2)', cp['aliquotas_vendas']['cofins'], f'=ROUND(B{r}*E{r},2)'], moeda=(1, 3, 5)); r += 1
    for rr in (l1, l1 + 1):
        ws.cell(rr, 3).number_format = PCT; ws.cell(rr, 5).number_format = PCT
    _linha(ws, r, ['Total', f'=B{l1}+B{l1+1}', None, f'=D{l1}+D{l1+1}', None, f'=F{l1}+F{l1+1}'], moeda=(1, 3, 5), bold=True); r += 1
    if pc['devolucoes_liquidas']:
        ws.cell(r, 1, f'Devoluções líquidas no mês: R$ {pc["devolucoes_liquidas"]:,.2f} - '
                      + ('deduzidas da base de vendas.' if pc['deduziu_devolucoes'] else 'NÃO deduzidas (ver Observações).')).font = F(italic=True)
        r += 1

    ir = a.ir
    r = _bloco(ws, r + 1, f'IRPJ / CSLL - {ir["trimestre"]}' + ('' if ir['fechamento'] else ' (parcial - trimestre ainda aberto ou faltando mês)'))
    _cab(ws, r, ['Mês', 'Faturamento', 'Origem']); r += 1
    m1 = r
    for mm, v, origem in ir['meses']:
        _linha(ws, r, [f'{mm[5:]}/{mm[:4]}', v, origem], moeda=(1,), fill=ALTA if v is None else None); r += 1
    _linha(ws, r, ['Receita do trimestre', f'=SUM(B{m1}:B{r-1})'], moeda=(1,), bold=True); rr = r; r += 1
    c = a.cfg['irpj_csll']
    linhas = [('CSLL - base presumida', f'=ROUND(B{rr}*{c["presuncao_csll"]},2)', f'{c["presuncao_csll"]:.0%}'),
              ('CSLL', f'=ROUND(B{rr+1}*{c["aliquota_csll"]},2)', f'{c["aliquota_csll"]:.0%}'),
              ('IRPJ - base presumida', f'=ROUND(B{rr}*{c["presuncao_irpj"]},2)', f'{c["presuncao_irpj"]:.0%}'),
              ('IRPJ normal', f'=ROUND(B{rr+3}*{c["aliquota_irpj"]},2)', f'{c["aliquota_irpj"]:.0%}'),
              ('Base do adicional', f'=MAX(0,B{rr+3}-{c["limite_adicional_trimestre"]})', f'excede R$ {c["limite_adicional_trimestre"]:,.0f}'),
              ('IRPJ adicional', f'=ROUND(B{rr+5}*{c["adicional_irpj"]},2)', f'{c["adicional_irpj"]:.0%}'),
              ('IRPJ total', f'=B{rr+4}+B{rr+6}', '')]
    for rot, fml, info in linhas:
        _linha(ws, r, [rot, fml, info], moeda=(1,), bold=rot in ('CSLL', 'IRPJ total')); r += 1

    r = _bloco(ws, r + 1, 'ICMS próprio')
    ic = a.ic
    itens = [('Débitos (saídas)', ic['debitos']), ('Ajustes a débito', ic['ajustes_debito']),
             ('(-) Créditos NF-e de entrada', -ic['creditos_nfe']), ('(-) Créditos CT-e', -ic['creditos_cte']),
             ('(-) Ajustes a crédito', -ic['ajustes_credito']), ('(-) Saldo credor anterior', -ic['saldo_credor_anterior'])]
    i0 = r
    for rot, v in itens:
        _linha(ws, r, [rot, v]); r += 1
    _linha(ws, r, ['ICMS a recolher (negativo = saldo credor)', f'=SUM(B{i0}:B{r-1})'], moeda=(1,), bold=True); r += 1
    for x in ic['ajustes']:
        ws.cell(r, 1, f'Ajuste {x["tipo"]} {x.get("cod", "")}: {x.get("descr", "")} - R$ {x["valor"]:,.2f}').font = F(italic=True); r += 1

    r = _bloco(ws, r + 1, 'ICMS-ST (substituto) por UF de destino')
    for uf_, v in a.st['por_uf'].items():
        _linha(ws, r, [f'ICMS-ST retido - {uf_}', v]); r += 1
    _linha(ws, r, ['(-) ST de devoluções', -a.st['devolucoes']]); r += 1
    _linha(ws, r, ['ICMS-ST a recolher PR', a.st['recolher_pr']], bold=True); r += 1

    r = _bloco(ws, r + 1, 'IPI')
    _linha(ws, r, ['Débitos (saídas)', a.ip['debitos']]); r += 1
    _linha(ws, r, ['(-) Créditos (entradas)', -a.ip['creditos']]); r += 1
    _linha(ws, r, ['IPI a recolher (negativo = saldo credor)', f'=B{r-2}+B{r-1}'], moeda=(1,), bold=True); r += 1

    if a.conf:
        r = _bloco(ws, r + 1, 'Conferência com o SPED Fiscal do mês')
        _cab(ws, r, ['Item', 'Apuração (XML)', 'SPED', 'Diferença']); r += 1
        for rot, v1, v2 in a.conf:
            dif = None if v2 is None else round(v1 - v2, 2)
            _linha(ws, r, [rot, v1, v2, dif], moeda=(1, 2, 3), fill=OK if dif == 0 else (MEDIA if dif is not None else None)); r += 1

    n_alta = sum(1 for d in a.div if d[0] == 'ALTA')
    r = _bloco(ws, r + 1, 'Conferência dos documentos')
    for rot, v in [('Notas de saída autorizadas', len(a.saidas)), ('Notas canceladas (saída e entrada)', len(a.canceladas)),
                   ('Números inutilizados', len(a.inutilizadas)), ('Notas de entrada', len(a.entradas)),
                   ('CT-e tomados', len(a.ctes)), ('Divergências ALTA', n_alta), ('Divergências (total)', len(a.div)),
                   ('Observações para a contadora', len(a.obs))]:
        _linha(ws, r, [rot, v], fill=ALTA if rot == 'Divergências ALTA' and v else None); ws.cell(r, 2).number_format = '0'; r += 1
    _larguras(ws, [52, 22, 34, 16, 12, 16])
    ws.freeze_panes = 'A4'


def acumuladores(wb, a):
    ws = wb.create_sheet('Saídas por acumulador')
    _titulo(ws, 'Resumo das saídas por acumulador (mesmo formato do relatório do Domínio)')
    _cab(ws, 3, ['Acumulador', 'Descrição', 'Grupo', 'Vlr contábil', 'Base ICMS', 'ICMS', 'Outras', 'IPI', 'BC ICMS-ST', 'ICMS-ST'])
    r = 4
    for x in a.acumuladores:
        _linha(ws, r, [x['acumulador'], x['descricao'], x['grupo'], x['vc'], x['bc'], x['icms'], x['outras'], x['ipi'], x['bcst'], x['st']],
               moeda=tuple(range(3, 10)), fill=ALTA if x['acumulador'] == '???' else None); r += 1
    _linha(ws, r, ['Total', '', ''] + [f'=SUM({get_column_letter(c)}4:{get_column_letter(c)}{r-1})' for c in range(4, 11)],
           moeda=tuple(range(3, 10)), bold=True)
    _larguras(ws, [11, 44, 16, 15, 15, 13, 14, 13, 14, 13])


def notas_saida(wb, a):
    ws = wb.create_sheet('Notas de saída')
    _titulo(ws, 'Notas de saída autorizadas do mês')
    H = ['Número', 'Série', 'Emissão', 'Cliente', 'CNPJ/CPF', 'UF', 'CFOPs', 'Vlr contábil', 'Base ICMS', 'ICMS', 'IPI', 'ICMS-ST', 'Chave']
    _cab(ws, 3, H)
    r = 4
    for n in a.saidas:
        its = n['itens']
        s = lambda k: round(sum(i[k] for i in its), 2)
        _linha(ws, r, [n['numero'], n['serie'], n['dhemi'], n['dest_nome'], n['dest_cnpj'], n['dest_uf'],
                       ', '.join(sorted({i['cfop'] for i in its})), round(sum(vcontabil(i) for i in its), 2),
                       s('vbc'), s('vicms'), s('vipi'), round(s('vst') + s('vfcpst'), 2), n['chave']], moeda=(7, 8, 9, 10, 11)); r += 1
    _linha(ws, r, ['Total'] + [''] * 6 + [f'=SUM({get_column_letter(c)}4:{get_column_letter(c)}{r-1})' for c in range(8, 13)] + [''],
           moeda=(7, 8, 9, 10, 11), bold=True)
    _larguras(ws, [9, 6, 11, 38, 16, 5, 14, 14, 14, 12, 12, 12, 47])
    ws.freeze_panes = 'A4'; ws.auto_filter.ref = f'A3:M{r-1}'


def itens_saida(wb, a):
    ws = wb.create_sheet('Itens de saída')
    _titulo(ws, 'Itens das notas de saída (base para a conferência)')
    H = ['Nota', 'Item', 'Produto', 'NCM', 'CFOP', 'Vlr contábil', 'CST ICMS', 'Base ICMS', '% ICMS', 'ICMS', 'CST IPI', '% IPI', 'IPI',
         'ICMS-ST', 'CST PIS', '% PIS', '% COFINS']
    _cab(ws, 3, H)
    r = 4
    for n in a.saidas:
        for i in n['itens']:
            _linha(ws, r, [n['numero'], i['n_item'], i['xprod'], i['ncm'], i['cfop'], round(vcontabil(i), 2), i['cst_icms'], i['vbc'],
                           i['picms'], i['vicms'], i['cst_ipi'], i['pipi'], i['vipi'], i['vst'], i['cst_pis'], i['ppis'], i['pcof']],
                   moeda=(5, 7, 9, 12, 13)); r += 1
    _larguras(ws, [8, 5, 40, 10, 6, 13, 7, 13, 7, 11, 7, 7, 11, 11, 7, 7, 8])
    ws.freeze_panes = 'A4'; ws.auto_filter.ref = f'A3:Q{r-1}'


def entradas(wb, a):
    ws = wb.create_sheet('Entradas e créditos')
    _titulo(ws, 'Entradas do mês e créditos de ICMS/IPI',
            'Coluna "Regra": natureza do item e de onde veio (SPED anterior, NCM ou CFOP). Laranja = sem histórico, conferir. Detalhe completo na planilha de créditos.')
    H = ['Emissão', 'Fornecedor', 'CNPJ', 'NF', 'Item', 'Produto', 'NCM', 'CFOP forn.', 'CFOP entrada', 'Vlr contábil',
         'ICMS destacado', 'Crédito ICMS', 'IPI destacado', 'Crédito IPI', 'Regra', 'Confere?', 'Observação']
    _cab(ws, 4, H)
    r = 5
    for l in a.ent:
        n, i = l['nota'], l['item']
        _linha(ws, r, [n['dhemi'], n['emit_nome'], n['emit_cnpj'], n['numero'], i['n_item'], i['xprod'], i['ncm'], i['cfop'], l['cfop'],
                       l['vc'], l['icms_destacado'], l['cred_icms'], l['ipi_destacado'], l['cred_ipi'], l['regra'], '', ''],
               moeda=(9, 10, 11, 12, 13), fill=MEDIA if not l.get('aprendido') and (l['icms_destacado'] or l['ipi_destacado']) else None)
        ws.cell(r, 16).fill = EDIT; ws.cell(r, 17).fill = EDIT; r += 1
    _linha(ws, r, ['Total'] + [''] * 8 + [f'=SUM({get_column_letter(c)}5:{get_column_letter(c)}{r-1})' for c in range(10, 15)] + ['', '', ''],
           moeda=(9, 10, 11, 12, 13), bold=True)
    if r > 5:
        dv = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'P5:P{r-1}')
    _larguras(ws, [11, 30, 16, 9, 5, 34, 10, 8, 8, 13, 12, 12, 11, 11, 24, 9, 30])
    ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:Q{r-1}'

    ws = wb.create_sheet('CT-e')
    _titulo(ws, 'CT-e em que a COSMETICI é tomadora (crédito de ICMS do frete)')
    _cab(ws, 3, ['Emissão', 'Transportadora', 'CNPJ', 'CT-e', 'CFOP', 'CFOP entrada', 'Origem-Destino', 'Valor', 'Base', '%', 'ICMS', 'Crédito'])
    r = 4
    for l in a.cte_cred:
        c = l['cte']
        _linha(ws, r, [c['data'], c['emit_nome'], c['emit_cnpj'], c['numero'], c['cfop'], l['cfop'], f'{c["uf_ini"]}-{c["uf_fim"]}',
                       c['vprest'], c['vbc'], c['picms'], c['vicms'], l['cred_icms']], moeda=(7, 8, 10, 11)); r += 1
    _linha(ws, r, ['Total'] + [''] * 6 + [f'=SUM({get_column_letter(c)}4:{get_column_letter(c)}{r-1})' for c in range(8, 13)],
           moeda=(7, 8, 9, 10, 11), bold=True)
    _larguras(ws, [11, 36, 16, 10, 6, 8, 10, 11, 11, 6, 10, 10])


def divergencias(wb, a):
    ws = wb.create_sheet('Divergências')
    _titulo(ws, 'Divergências encontradas na conferência', 'ALTA = corrigir antes de fechar; MÉDIA = conferir; BAIXA = informativo.')
    _cab(ws, 4, ['Gravidade', 'Tipo', 'Documento', 'Descrição', 'Valor', 'Resolvido?', 'Observação'])
    r = 5
    for g, t, d, desc, v in a.div:
        _linha(ws, r, [g, t, d, desc, v, '', ''], moeda=(4,), fill={'ALTA': ALTA, 'MÉDIA': MEDIA}.get(g))
        ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(r, 6).fill = EDIT; ws.cell(r, 7).fill = EDIT; r += 1
    if r == 5:
        ws.cell(5, 1, 'Nenhuma divergência.').font = F(italic=True)
    else:
        dv = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'F5:F{r-1}')
    _larguras(ws, [10, 22, 42, 80, 12, 11, 30])
    ws.freeze_panes = 'A5'

    ws = wb.create_sheet('Observações')
    _titulo(ws, 'Pontos para a contadora decidir (não alteram a apuração sem confirmação)')
    for k, o in enumerate(a.obs, 3):
        x = ws.cell(k, 1, f'{k-2}. {o}'); x.font = F(); x.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[k].height = 15 * (len(o) // 120 + 1)
    if not a.obs:
        ws.cell(3, 1, 'Nenhuma.').font = F(italic=True)
    ws.column_dimensions['A'].width = 140


def canceladas(wb, a):
    ws = wb.create_sheet('Canceladas e inutilizadas')
    _titulo(ws, 'Notas canceladas e números inutilizados no mês')
    _cab(ws, 3, ['Tipo', 'Número', 'Série', 'Emissão', 'Destinatário/Emitente', 'Valor', 'Chave'])
    r = 4
    for n in a.canceladas:
        _linha(ws, r, ['Cancelada' if n['propria'] else 'Cancelada (fornecedor)', n['numero'], n['serie'], n['dhemi'],
                       n['dest_nome'] if n['propria'] else n['emit_nome'], n['vnf'], n['chave']], moeda=(5,)); r += 1
    for s, x in a.inutilizadas:
        _linha(ws, r, ['Inutilizada', x, s, None, '', None, '']); r += 1
    _larguras(ws, [22, 9, 6, 11, 40, 13, 47])


def gerar(a, destino):
    wb = Workbook()
    resumo(wb, a)
    acumuladores(wb, a)
    notas_saida(wb, a)
    itens_saida(wb, a)
    entradas(wb, a)
    divergencias(wb, a)
    canceladas(wb, a)
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
