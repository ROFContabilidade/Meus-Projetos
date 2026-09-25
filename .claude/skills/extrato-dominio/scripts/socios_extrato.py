#!/usr/bin/env python3
"""Planilha de movimentos com os sócios, mês a mês, conforme o extrato (entradas e saídas).

Lê os classificados finais de cada mês (o CSV que gerou o TXT) e separa tudo o que envolve os
sócios cadastrados no JSON da empresa (campo "socios": nome e CPF):
  - pagamentos e recebimentos em que o favorecido é o sócio (nome completo ou CPF, mesmo mascarado);
  - lançamentos nas contas com o nome do sócio no plano de contas (ex.: conta do sócio, mútuo do
    sócio) e nas contas extras de "contas_socios" no JSON (ex.: {"620": "Sócios (mútuo)"}).
Padrão ROF: só o primeiro nome não basta. Linha com o primeiro nome do sócio que não confere
pelo nome completo nem pelo CPF vai para a aba "A conferir" (amarelo) e vira pergunta.

Uso:
  python socios_extrato.py -e empresas/<cod>-<empresa>.json -o <COD>_<ANO>_socios.xlsx \
      [--historico empresas/<cod>-<empresa>-socios.csv] <classificado final do mês>...
Com --historico, os movimentos dos meses informados substituem os mesmos meses no histórico e a
planilha sai com o ano todo: numa conversa nova basta o classificado do mês (o histórico vem no
pacote da skill da empresa).
Abas: "Resumo" (mês × sócio: entradas, saídas, líquido e acumulado), "Movimentos" (cada linha,
verde = confirmado) e "A conferir" (amarelo), se houver.
"""
import argparse, csv, json, re, sys, unicodedata
from collections import defaultdict
from decimal import Decimal

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    sys.exit("Instale openpyxl: pip install openpyxl")

VERDE = PatternFill("solid", fgColor="C6EFCE")
AMARELO = PatternFill("solid", fgColor="FFEB9C")
CINZA = PatternFill("solid", fgColor="D9D9D9")
MESES = ["", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto",
         "Setembro", "Outubro", "Novembro", "Dezembro"]


def norm(s):
    s = unicodedata.normalize("NFKD", s or "")
    return re.sub(r"\s+", " ", "".join(c for c in s if not unicodedata.combining(c))).strip().upper()


def digitos_visiveis(doc):
    return [t for t in re.split(r"\*+", re.sub(r"[^\d*]", "", doc or "")) if len(t) >= 3]


def mesmo_cpf(doc, cpf):
    cheio = re.sub(r"\D", "", cpf or "")
    partes = digitos_visiveis(doc)
    return bool(cheio and partes) and all(t in cheio for t in partes)


def carregar_socios(emp):
    socios, contas = [], {}
    for s in emp.get("socios", []):
        nome = norm(s.get("nome"))
        p = nome.split()
        # o extrato corta o nome ("JOAO SILV"): primeiro nome + 4 letras do segundo
        chave = f"{p[0]} {p[1][:4]}" if len(p) > 1 else nome
        socios.append({"nome": s.get("nome", ""), "cpf": s.get("cpf", ""), "chave": chave, "primeiro": p[0] if p else ""})
        # contas do plano com o nome do sócio (ex.: conta "JOAO SILVA", "MUTUO JOAO SILVA")
        for c, nome_conta in (emp.get("contas") or {}).items():
            if f" {chave}" in f" {norm(nome_conta)}":
                contas[str(c)] = s.get("nome", "")
    for c, quem in (emp.get("contas_socios") or {}).items():
        contas[str(c)] = quem
    return socios, contas


def identificar(linha, socios, contas):
    """(sócio, confere) — confere=False quando só o primeiro nome bate."""
    desc = f" {norm(linha['descricao'])} "
    doc = linha.get("documento") or ""
    for s in socios:
        if (s["cpf"] and mesmo_cpf(doc, s["cpf"])) or f" {s['chave']}" in desc:
            return s["nome"], True
    if linha.get("conta") in contas:
        return contas[linha["conta"]], True
    for s in socios:
        if s["primeiro"] and re.search(rf"(?<![A-Z]){re.escape(s['primeiro'])}(?![A-Z])", desc):
            if digitos_visiveis(doc) and s["cpf"] and not mesmo_cpf(doc, s["cpf"]):
                continue  # mesmo primeiro nome, outro CPF/CNPJ: não é o sócio
            return s["nome"], False
    return None, None


def tipo_mov(conta, nome_conta, valor, contas_socio):
    n = norm(nome_conta)
    if conta and conta not in contas_socio and not any(k in n for k in ("PRO-LABORE", "PRO LABORE", "MUTUO", "LUCRO")):
        return "Pago ao sócio como despesa da empresa (reembolso etc.)" if valor < 0 else "Recebido do sócio (conferir conta)"
    if "PRO-LABORE" in n or "PRO LABORE" in n:
        return "Pró-labore"
    if "MUTUO" in n:
        return "Mútuo"
    if "LUCRO" in n:
        return "Distribuição de lucros"
    return "Retirada" if valor < 0 else "Entrada / devolução"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("classificados", nargs="+", help="CSV final de cada mês (separador ;)")
    ap.add_argument("-e", "--empresa", required=True)
    ap.add_argument("-o", "--saida", required=True)
    ap.add_argument("--historico", help="CSV acumulado dos movimentos com sócios (lido e atualizado)")
    a = ap.parse_args()

    emp = json.load(open(a.empresa, encoding="utf-8"))
    plano = emp.get("contas", {})
    socios, contas = carregar_socios(emp)
    if not socios:
        sys.exit("O JSON da empresa não tem 'socios'.")

    movs, conferir = [], []
    for arq in a.classificados:
        for l in csv.DictReader(open(arq, encoding="utf-8-sig"), delimiter=";"):
            quem, ok = identificar(l, socios, contas)
            if not quem:
                continue
            v = Decimal(l["valor"])
            nome_conta = plano.get(l.get("conta", ""), "")
            r = {"mes": l["data"][6:10] + "-" + l["data"][3:5], "data": l["data"], "socio": quem,
                 "descricao": l["descricao"], "documento": l.get("documento", ""),
                 "entrada": v if v > 0 else Decimal(0), "saida": -v if v < 0 else Decimal(0),
                 "conta": l.get("conta", ""), "nome_conta": nome_conta[13:] if nome_conta[:1].isdigit() else nome_conta,
                 "tipo": tipo_mov(l.get("conta", ""), nome_conta, v, contas), "status": l.get("status", ""),
                 "origem": l.get("regra", "")}
            (movs if ok and r["status"] == "CONFIRMADO" else conferir).append(r)
    if a.historico:
        campos = ["mes", "data", "socio", "descricao", "documento", "entrada", "saida", "conta", "nome_conta",
                  "tipo", "status", "origem", "conferir"]
        novos_meses = {r["mes"] for r in movs + conferir}
        antigos = []
        try:
            antigos = [x for x in csv.DictReader(open(a.historico, encoding="utf-8-sig"), delimiter=";")
                       if x["mes"] not in novos_meses]
        except FileNotFoundError:
            pass
        for x in antigos:
            x["entrada"], x["saida"] = Decimal(x["entrada"]), Decimal(x["saida"])
            (conferir if x.pop("conferir") == "1" else movs).append(x)
        with open(a.historico, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=campos, delimiter=";")
            w.writeheader()
            for lista, flag in ((movs, "0"), (conferir, "1")):
                for r in lista:
                    w.writerow({**r, "conferir": flag})
        print(f"Histórico: {a.historico} ({len(antigos)} movimento(s) de meses anteriores mantidos)")
    movs.sort(key=lambda r: (r["data"][6:] + r["data"][3:5] + r["data"][:2], r["socio"]))
    conferir.sort(key=lambda r: (r["data"][6:] + r["data"][3:5] + r["data"][:2], r["socio"]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append([f"{emp.get('razao_social', '')} - movimentos com sócios conforme o extrato"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Mês", "Sócio", "Entradas", "Saídas (retiradas)", "Líquido (saídas − entradas)", "Acumulado no ano",
               "Pagos ao sócio como despesa (fora do líquido)"])
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = CINZA
    tot = defaultdict(lambda: [Decimal(0), Decimal(0), Decimal(0)])
    for r in movs:
        t = tot[(r["mes"], r["socio"])]
        if r["tipo"].startswith("Pago ao sócio como despesa"):
            t[2] += r["saida"] - r["entrada"]
        else:
            t[0] += r["entrada"]
            t[1] += r["saida"]
    acum = defaultdict(Decimal)
    meses = sorted({m for m, _ in tot})
    for m in meses:
        for s in sorted({s for mm, s in tot if mm == m}):
            e, sa, ou = tot[(m, s)]
            acum[s] += sa - e
            ws.append([f"{MESES[int(m[5:])]}/{m[:4]}", s, float(e), float(sa), float(sa - e), float(acum[s]),
                       float(ou) or None])
    ws.append([])
    for s in sorted(acum):
        e = sum(v[0] for (m, ss), v in tot.items() if ss == s)
        sa = sum(v[1] for (m, ss), v in tot.items() if ss == s)
        ou = sum(v[2] for (m, ss), v in tot.items() if ss == s)
        ws.append(["TOTAL", s, float(e), float(sa), float(sa - e), "", float(ou) or None])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, min_col=3, max_col=7):
        for c in row:
            c.number_format = "#,##0.00"
    for col, w in zip("ABCDEFG", [16, 36, 14, 18, 24, 18, 22]):
        ws.column_dimensions[col].width = w

    cab = ["Mês", "Data", "Sócio", "Descrição no extrato", "CPF/CNPJ", "Entrada", "Saída", "Conta",
           "Nome da conta", "Tipo", "Status", "Origem da classificação"]

    def aba(titulo, linhas, cor):
        w = wb.create_sheet(titulo)
        w.append(cab)
        for c in w[1]:
            c.font = Font(bold=True)
            c.fill = CINZA
        for r in linhas:
            w.append([f"{MESES[int(r['mes'][5:])]}/{r['mes'][:4]}", r["data"], r["socio"], r["descricao"],
                      r["documento"], float(r["entrada"]) or None, float(r["saida"]) or None, r["conta"],
                      r["nome_conta"], r["tipo"], r["status"], r["origem"]])
            for c in w[w.max_row]:
                c.fill = cor
            w.cell(w.max_row, 6).number_format = w.cell(w.max_row, 7).number_format = "#,##0.00"
        for col, wd in zip("ABCDEFGHIJKL", [14, 11, 34, 48, 18, 12, 12, 7, 34, 20, 12, 60]):
            w.column_dimensions[col].width = wd
        w.freeze_panes = "A2"
        w.auto_filter.ref = w.dimensions

    aba("Movimentos", movs, VERDE)
    if conferir:
        aba("A conferir", conferir, AMARELO)
    wb.save(a.saida)
    print(f"{len(movs)} movimento(s) com sócios | {len(conferir)} a conferir -> {a.saida}")
    for s in sorted(acum):
        e = sum(v[0] for (m, ss), v in tot.items() if ss == s)
        sa = sum(v[1] for (m, ss), v in tot.items() if ss == s)
        ou = sum(v[2] for (m, ss), v in tot.items() if ss == s)
        print(f"  {s}: entradas {e:,.2f} | saídas {sa:,.2f} | líquido {sa - e:,.2f} | pagos como despesa {ou:,.2f}")
    for r in conferir:
        print(f"  A CONFERIR {r['data']} {r['descricao'][:50]} {r['documento']} {r['entrada'] or -r['saida']} conta {r['conta']}")


if __name__ == "__main__":
    main()
