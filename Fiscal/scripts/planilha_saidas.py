"""Relação das NF-e de SAÍDA do mês em uma única aba, por data de saída e número:
notas a importar no Domínio, canceladas (destacadas) e numerações a inutilizar.
Valores do XML; colunas que vêm zeradas em todas as notas do mês ficam de fora.
"""
import os
import re
from collections import Counter
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = PatternFill('solid', fgColor='1F3864')
CINZA = PatternFill('solid', fgColor='D9E1F2')
VERDE = PatternFill('solid', fgColor='E2EFDA')
AMARELO = PatternFill('solid', fgColor='FFF2CC')
CANC = PatternFill('solid', fgColor='FFC7CE')
INUT = PatternFill('solid', fgColor='D9D9D9')
FINO = Side(style='thin', color='BFBFBF')
NUM = '#,##0.00'
Z = Decimal('0')

# (coluna, campo do total da NF-e)
VALORES = [('Valor produtos', 't_vProd'), ('Frete', 't_vFrete'), ('Seguro', 't_vSeg'), ('Desconto', 't_vDesc'),
           ('Outras despesas', 't_vOutro'), ('Base ICMS', 't_vBC'), ('ICMS', 't_vICMS'),
           ('ICMS desonerado', 't_vICMSDeson'), ('FCP', 't_vFCP'), ('Base ICMS ST', 't_vBCST'), ('ICMS ST', 't_vST'),
           ('FCP ST', 't_vFCPST'), ('IPI', 't_vIPI'), ('IPI devolvido', 't_vIPIDevol'), ('PIS', 't_vPIS'),
           ('COFINS', 't_vCOFINS'), ('DIFAL UF destino', 't_vICMSUFDest'), ('FCP UF destino', 't_vFCPUFDest'),
           ('Valor da nota', 't_vNF')]


def _br(d):
    d = str(d or '')
    return f'{d[8:10]}/{d[5:7]}/{d[:4]}' if len(d) >= 10 and d[4] == '-' else d


def _iso(d):
    d = str(d or '')
    return f'{d[6:10]}-{d[3:5]}-{d[:2]}' if re.match(r'\d\d/\d\d/\d{4}', d) else d[:10]


def _lista(c):
    return ', '.join(k for k, _ in Counter(c).most_common())


def _so_na_sefaz(xml_dir):
    """Chaves que só existem no download da SEFAZ (DownloadDFe), fora dos zips que a empresa mandou.
    Na importação pelo zip do cliente essas notas ficariam de fora."""
    if not xml_dir or not os.path.isdir(xml_dir):
        return set()
    sefaz, cliente = set(), set()
    for raiz, _, arqs in os.walk(xml_dir):
        alvo = sefaz if 'downloaddfe' in raiz.lower() else cliente
        for a in arqs:
            m = re.search(r'\d{44}', a)
            if m:
                alvo.add(m.group(0))
    return sefaz - cliente


def gerar(res, saida):
    notas = res['notas']
    so_sefaz = _so_na_sefaz(res.get('xml_dir'))
    comp = res['competencia']
    emp = res['regras'].r['empresa']
    linhas = []
    for s in res['saidas']:
        n = notas.get(s['Chave'])
        if not n:
            continue
        cancel = s['Status'].startswith('CANCELADA')
        its = n['itens']
        ln = {'Situação': 'CANCELADA - lançar zerada' if cancel else
              ('IMPORTAR' if s['Status'] == 'OK' else 'CONFERIR: ' + s['Status']),
              'Data saída': n.get('dt_saida_ent') or n['dt_emissao'], 'Data emissão': n['dt_emissao'],
              'Nota': int(n['numero']), 'Série': n['serie'], 'Cliente': n['dest_nome'], 'CNPJ/CPF': n['dest_cnpj'],
              'UF': n['dest_uf'], 'Natureza da operação': n['natOp'], 'CFOP': _lista(i['CFOP'] for i in its),
              'CST ICMS': _lista(i['orig'] + i['CST_ICMS'] for i in its),
              'cBenef': _lista(i['cBenef'] for i in its if i['cBenef']),
              'Acumulador': '' if cancel else re.sub(r'\d{4}→', '', s.get('Acum. sugerido', '')),
              'Chave': n['chave'], '_tipo': 'C' if cancel else 'N'}
        for col, k in VALORES:
            ln[col] = n.get(k) or Z
        if n['chave'] in so_sefaz and not cancel:
            ln['Observação'] = ('Nota ativa, autorizada na SEFAZ. O XML não veio no zip de vendas da empresa, só no '
                                'download da SEFAZ: importar este XML junto (pasta XML_fora_do_zip_do_cliente).')
        if cancel:
            ln['Observação'] = (f'Cancelada. Valores do XML só para referência: no Domínio lançar com situação '
                                f'CANCELADA, só cliente e chave, tudo zerado.')
        linhas.append(ln)
    for l in res['lacunas']:
        sit = l['Situação']
        if not sit.startswith(('Inutilizar', 'Inutilizada')):
            continue
        m = re.search(r'(\d{4}-\d\d-\d\d|\d\d/\d\d/\d{4})', sit)
        d = _iso(m.group(1)) if m else ''
        cli = str(l.get('Cliente (ERP/Domínio)') or '').strip()
        ln = {'Situação': 'INUTILIZADA', 'Data saída': d, 'Data emissão': d, 'Nota': int(l['Número']),
              'Série': str(l['Série']).lstrip('0') or '0', 'Cliente': '', '_tipo': 'I',
              'Observação': 'Número sem nota emitida: vai no TXT de inutilizadas (C100 COD_SIT 05).'}
        if cli:      # o ERP tem pedido com cliente e valor, mas a NF-e não existe na SEFAZ
            v = l.get('Valor no ERP')
            ln.update({'Situação': 'INUTILIZADA - CONFIRMAR COM A KOPP', 'Cliente': cli + ' (ERP)', '_tipo': 'A',
                       'Observação': f'No relatório do ERP com cliente e valor R$ {v:,.2f}'.replace(',', 'X')
                       .replace('.', ',').replace('X', '.') + ', sem chave e sem situação. A NF-e não existe na '
                       'SEFAZ (fora do download DFe). Confirmar se a mercadoria saiu: se saiu, falta emitir nota; '
                       'o número precisa ser inutilizado na SEFAZ.'})
        linhas.append(ln)
    linhas.sort(key=lambda l: (l['Data saída'], l['Nota']))

    # colunas de valor zeradas em todas as notas do mês saem da planilha
    ativos = [c for c, _ in VALORES if c == 'Valor da nota' or any(l.get(c) for l in linhas)]
    mostra_emissao = any(l['Data emissão'] != l['Data saída'] for l in linhas)
    cols = (['Situação', 'Data saída'] + (['Data emissão'] if mostra_emissao else []) +
            ['Nota', 'Série', 'Cliente', 'CNPJ/CPF', 'UF', 'Natureza da operação', 'CFOP', 'CST ICMS']
            + (['cBenef'] if any(l.get('cBenef') for l in linhas) else []) + ['Acumulador'] + ativos +
            ['Observação', 'Chave'])

    wb = Workbook()
    ws = wb.active
    ws.title = f'Saídas {comp[5:]}-{comp[:4]}'
    ws['A1'] = f'Notas fiscais de saída {comp[5:]}/{comp[:4]} — {emp["nome"]}'
    ws['A1'].font = Font(bold=True, size=13, color='1F3864')
    nI = sum(1 for l in linhas if l['_tipo'] == 'N')
    nC = sum(1 for l in linhas if l['_tipo'] == 'C')
    nU = sum(1 for l in linhas if l['_tipo'] in 'IA')
    nA = sum(1 for l in linhas if l['_tipo'] == 'A')
    ws['A2'] = (f'{nI} para importar no Domínio · {nC} canceladas (vermelho) · {nU} inutilizadas (cinza; {nA} em amarelo '
                f'para confirmar com a empresa). '
                'Ordem: data de saída e número. Valores do XML.')
    ws['A2'].font = Font(italic=True, color='595959')
    r0 = 4
    for j, c in enumerate(cols, 1):
        cel = ws.cell(r0, j, c)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = AZUL
        cel.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for i, ln in enumerate(linhas, r0 + 1):
        fill = {'C': CANC, 'I': INUT, 'A': AMARELO}.get(ln['_tipo'])
        for j, c in enumerate(cols, 1):
            v = ln.get(c)
            if c.startswith('Data'):
                v = _br(v)
            if isinstance(v, Decimal):
                v = float(v)
            cel = ws.cell(i, j, v)
            cel.border = Border(bottom=FINO)
            if isinstance(v, float):
                cel.number_format = NUM
            if fill:
                cel.fill = fill
                if ln['_tipo'] == 'C':
                    cel.font = Font(color='9C0006')
        if not fill:
            ws.cell(i, 1).fill = VERDE if ln['Situação'] == 'IMPORTAR' else AMARELO
    # totais só das notas a importar (cancelada e inutilizada entram zeradas)
    tot = r0 + len(linhas) + 1
    sit_col = f'$A${r0 + 1}:$A${tot - 1}'
    ws.cell(tot, 1, 'TOTAL A IMPORTAR').font = Font(bold=True)
    for c in ativos:
        j = cols.index(c) + 1
        col = get_column_letter(j)
        cel = ws.cell(tot, j, f'=SUMIF({sit_col},"IMPORTAR",{col}{r0 + 1}:{col}{tot - 1})')
        cel.font = Font(bold=True)
        cel.number_format = NUM
    for j in range(1, len(cols) + 1):
        ws.cell(tot, j).fill = CINZA
    ws.freeze_panes = ws.cell(r0 + 1, cols.index('Cliente') + 2)
    ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(cols))}{tot - 1}'
    ws.row_dimensions[r0].height = 34
    larg = {'Situação': 24, 'Cliente': 36, 'Natureza da operação': 32, 'CNPJ/CPF': 16, 'Observação': 48,
            'Chave': 46, 'Nota': 8, 'Série': 6, 'UF': 5, 'Data saída': 11, 'Data emissão': 11, 'CFOP': 8,
            'CST ICMS': 9, 'cBenef': 10, 'Acumulador': 11}
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = larg.get(c, 13)
    wb.save(saida)
    return nI, nC, nU
