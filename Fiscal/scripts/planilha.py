"""Grava o resultado da conferência em uma planilha Excel formatada."""
from collections import Counter
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = PatternFill('solid', fgColor='1F3864')
VERDE = PatternFill('solid', fgColor='E2EFDA')
AMARELO = PatternFill('solid', fgColor='FFF2CC')
VERMELHO = PatternFill('solid', fgColor='F8CBAD')
CINZA = PatternFill('solid', fgColor='EDEDED')
FINO = Side(style='thin', color='BFBFBF')
NUM = '#,##0.00'


def _cor_status(v):
    v = str(v or '')
    if v.startswith('OK') or v.startswith('CANCELADA - OK') or v == 'INFO':
        return VERDE
    if 'NÃO' in v or v == 'ALTA' or 'LANÇADA COM VALOR' in v or 'SEM COD_SIT' in v:
        return VERMELHO
    return AMARELO


def _aba(wb, titulo, linhas, larguras=None, status_col=None, nota=None):
    ws = wb.create_sheet(titulo)
    r0 = 1
    if nota:
        ws.cell(1, 1, nota).font = Font(italic=True, color='595959')
        r0 = 3
    if not linhas:
        ws.cell(r0, 1, 'Nada a mostrar.')
        return ws
    cols = list(linhas[0].keys())
    for extra in linhas[1:]:
        for k in extra:
            if k not in cols:
                cols.append(k)
    for j, c in enumerate(cols, 1):
        cel = ws.cell(r0, j, c)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = AZUL
        cel.alignment = Alignment(wrap_text=True, vertical='center')
    for i, ln in enumerate(linhas, r0 + 1):
        for j, c in enumerate(cols, 1):
            v = ln.get(c)
            if isinstance(v, Decimal):
                v = float(v)
            cel = ws.cell(i, j, v)
            if isinstance(v, float):
                cel.number_format = NUM
            cel.border = Border(bottom=FINO)
        if status_col and status_col in cols:
            ws.cell(i, cols.index(status_col) + 1).fill = _cor_status(ln.get(status_col))
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(cols))}{r0 + len(linhas)}'
    larguras = larguras or {}
    for j, c in enumerate(cols, 1):
        amostra = [len(str(ln.get(c) or '')) for ln in linhas[:200]]
        w = larguras.get(c, min(max([len(c) * 0.9] + amostra) + 2, 45))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[r0].height = 32
    return ws


POS_ESCRITURACAO = ('SPED', 'importação', 'Domínio', 'Corrigido', 'E113')


def _so_pre(linhas):
    """No modo pré-importação tira as colunas que só fazem sentido depois do lançamento."""
    return [{k: v for k, v in ln.items()
             if not any(t in k for t in POS_ESCRITURACAO) or k == 'Histórico do fornecedor no Domínio'}
            for ln in linhas]


CSRF_PARTES = (('PIS', Decimal('0.65')), ('COFINS', Decimal('3')), ('CSLL', Decimal('1')))


def _retencoes(nfse):
    """Separa as NFS-e tomadas com e sem retenção (base da EFD-Reinf)."""
    com, sem = [], []
    tot = Counter()
    for x in nfse:
        if str(x.get('Status', '')).startswith(('OK (cancelada', 'JÁ LANÇADA', 'CANCELADA')):
            continue
        v = {k: Decimal(str(x.get(f'{k} retido') or 0)) for k in ('IRRF', 'CSRF', 'INSS', 'ISS')}
        base = {'Número': x['Número'], 'Emissão': x.get('Data emissão', ''), 'Competência': x['Competência'],
                'Prestador': x['Prestador'], 'CNPJ': x['CNPJ'], 'Serviço (LC 116)': x.get('Serviço (LC 116)', ''),
                'Natureza do serviço': x.get('Natureza do serviço', ''), 'Valor bruto': x['Valor'],
                'Acum. sugerido': x.get('Acum. sugerido', ''), 'Acumulador (nome)': x.get('Acumulador sugerido (nome)', '')}
        if not any(v.values()):
            sem.append({**base, 'Simples Nacional': x.get('Simples Nacional', ''), 'Alertas': x.get('Alertas', '')})
            continue
        ln = dict(base)
        ln['IRRF'] = v['IRRF']
        ln['Cód. DARF IRRF'] = ('8045' if x.get('Serviço (LC 116)', '').startswith('17.25') else '1708') if v['IRRF'] else ''
        for nome, aliq in CSRF_PARTES:
            ln[nome] = (v['CSRF'] * aliq / Decimal('4.65')).quantize(Decimal('0.01')) if v['CSRF'] else Decimal('0')
        ln['CSRF total'] = v['CSRF']
        ln['Cód. DARF CSRF'] = '5952' if v['CSRF'] else ''
        ln['INSS retido'] = v['INSS']
        ln['ISS retido'] = v['ISS']
        ln['Evento EFD-Reinf'] = ', '.join(e for e, ok in (('R-4020', v['IRRF'] or v['CSRF']), ('R-2010', v['INSS'])) if ok)
        ln['Observação'] = ('ISS retido: guia municipal de ' + str(x.get('Município do prestador', '')) if v['ISS'] else '')
        com.append(ln)
        for k in ('IRRF', 'PIS', 'COFINS', 'CSLL', 'CSRF total', 'INSS retido', 'ISS retido', 'Valor bruto'):
            tot[k] += Decimal(str(ln[k] or 0))
    if com:
        com.append({'Número': 'TOTAL', 'Prestador': f'{len(com)} notas com retenção',
                    **{k: tot[k] for k in ('Valor bruto', 'IRRF', 'PIS', 'COFINS', 'CSLL', 'CSRF total', 'INSS retido', 'ISS retido')}})
    return com, sem, tot


NOTA_REINF = ('NFS-e tomadas COM retenção: base do R-4020 (IRRF e PIS/COFINS/CSLL, DARF 1708/8045 e 5952) e do R-2010 (INSS). '
              'Atenção: no R-4020 o período é o mês do PAGAMENTO (ou crédito) ao prestador, não o da emissão da nota.')


def _abas_servicos(wb, nfse):
    com, sem, tot = _retencoes(nfse)
    _aba(wb, 'Retenções (EFD-Reinf)', com, nota=NOTA_REINF,
         larguras={'Prestador': 34, 'Natureza do serviço': 34, 'Observação': 40})
    _aba(wb, 'NFS-e sem retenção', sem, nota='NFS-e tomadas sem nenhuma retenção. "Alertas" aponta onde a retenção '
         'era esperada (prestador fora do Simples, serviço da lista de retenção).',
         larguras={'Prestador': 34, 'Natureza do serviço': 34, 'Alertas': 80})
    return com, sem, tot


def _linhas_reinf(nfse):
    com, sem, tot = _retencoes(nfse)
    if not com:
        return [('   Nenhuma NFS-e com retenção.', None)]
    f = lambda v: f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return [(f'   Com retenção: {len(com) - 1} notas (aba "Retenções (EFD-Reinf)") · IRRF R$ {f(tot["IRRF"])} · '
             f'PIS/COFINS/CSLL R$ {f(tot["CSRF total"])} · INSS R$ {f(tot["INSS retido"])} · ISS R$ {f(tot["ISS retido"])}', 'destaque'),
            (f'   Sem retenção: {len(sem)} notas (aba "NFS-e sem retenção"), '
             f'{sum(1 for x in sem if "falta retenção" in (x["Alertas"] or ""))} com retenção que parecia devida', None)]


def _escreve_resumo(ws, linhas):
    for i, (txt, estilo) in enumerate(linhas, 1):
        c = ws.cell(i, 1, txt)
        if estilo == 'titulo':
            c.font = Font(bold=True, size=15, color='1F3864')
        elif estilo == 'sub':
            c.font = Font(color='595959')
        elif estilo == 'secao':
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = AZUL
        elif estilo == 'destaque':
            c.font = Font(bold=True)
            c.fill = VERDE
    ws.column_dimensions['A'].width = 130


def _regras(R):
    regras = [{'Tipo': 'NCM', 'Chave': r['prefixo'], 'Finalidade': r['finalidade'], 'Motivo': r['motivo']}
              for r in R.r['finalidade_por_ncm']]
    regras += [{'Tipo': 'Fornecedor', 'Chave': k, 'Finalidade': v['finalidade'], 'Motivo': v['motivo']}
               for k, v in R.r['finalidade_por_fornecedor'].items()]
    regras += [{'Tipo': 'Item', 'Chave': f"{r['cnpj']} NCM {r.get('ncm', '')}", 'Finalidade': r['finalidade'],
                'Motivo': r['motivo']} for r in R.r.get('finalidade_por_item', [])]
    regras += [{'Tipo': 'Compra → CFOP/acumulador', 'Chave': fin,
                'Finalidade': f"interna {v['interna']['cfop']}/ac {v['interna']['ac']} · interestadual "
                              f"{v['interestadual']['cfop']}/ac {v['interestadual']['ac']}",
                'Motivo': f"crédito ICMS: {v.get('credito_icms')} · crédito IPI: {v.get('credito_ipi')}"}
               for fin, v in R.r['compras'].items()]
    return regras


def gravar_pre(res, saida):
    R = res['regras']
    emp = R.r['empresa']
    comp = res['competencia']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumo'
    itens, notas_t, dev, sai, pend = res['itens'], res['notas_terc'], res['devol'], res['saidas'], res['pendencias']
    sai_ativas = [s for s in sai if not s['Status'].startswith('CANCELADA')]
    dev_ativas = [d for d in dev if not d['Status'].startswith('CANCELADA')]
    fin = Counter(i['Finalidade'] for i in itens)
    com_conta = sum(1 for i in itens if i['Conta contábil sugerida'])
    conf_hist = sum(1 for n in notas_t if n['Status'] == 'CONFERIR HISTÓRICO')
    sem_hist = sum(1 for n in notas_t if 'sem lançamento no histórico' in (n['Observações'] or ''))
    per = ', '.join(f'{a} a {b}' for a, b in res['periodos_hist'] if a) or 'não informado'
    nfse = res['nfse']
    linhas = [
        (f'Pré-conferência {comp[5:]}/{comp[:4]} — {emp["nome"]} (antes de importar no Domínio)', 'titulo'),
        (f'CNPJ {emp["cnpj"]} · Domínio {emp["codigo_dominio"]} · {emp["regime"]} · {emp["uf"]}', 'sub'),
        ('', None),
        ('O QUE FOI LIDO', 'secao'),
        (f'{len(res["notas"])} NF-e lidas do XML, item por item · lista da SEFAZ com {len(res["sefaz"])} chaves', None),
        (f'Histórico do Domínio usado como referência: {per} '
         f'({sum(sum(c.values()) for c in res["hist"].values())} lançamentos, {len(res["hist"])} participantes)', None),
        ('Não há SPED nem relatório do Domínio deste mês: as colunas de comparação com o lançamento final não aparecem.', None),
        ('', None),
        ('ENTRADAS DE FORNECEDORES — o que lançar e como', 'secao'),
        (f'{len(notas_t)} notas · {len(itens)} itens com CFOP, acumulador e conta contábil sugeridos', 'destaque'),
        ('   Finalidades: ' + ', '.join(f'{k.lower()} {v}' for k, v in fin.most_common()), None),
        (f'   Conta contábil sugerida em {com_conta} de {len(itens)} itens', None),
        (f'Notas em que a sugestão difere de como o fornecedor foi lançado antes: {conf_hist} (status CONFERIR HISTÓRICO)', None),
        (f'Fornecedores sem nenhum lançamento no histórico: {sem_hist} notas', None),
        ('', None),
        ('NOTAS EMITIDAS PELA KOPP', 'secao'),
        (f'Saídas: {len(sai_ativas)} notas · prontas {sum(1 for s in sai_ativas if s["Status"] == "OK")} · '
         f'canceladas {len(sai) - len(sai_ativas)}', None),
        (f'Entradas próprias (devoluções/trocas/retornos): {len(dev_ativas)} · prontas '
         f'{sum(1 for d in dev_ativas if d["Status"] == "OK")}', None),
        ('', None),
        ('SERVIÇOS TOMADOS (NFS-e)', 'secao'),
        (f'{len(nfse)} NFS-e na lista nacional · prontas para lançar {sum(1 for x in nfse if x["Status"] == "A LANÇAR")} · '
         f'conferir antes {sum(1 for x in nfse if x["Status"] == "CONFERIR")} · '
         f'já lançadas antes {sum(1 for x in nfse if x["Status"].startswith("JÁ LANÇADA"))}', None),
        ('   Acumulador de serviço escolhido pelo item da LC 116 da nota + retenções (catálogo de acumuladores do Domínio); '
         'conta pela natureza do serviço ou pelo razão.', None),
        *_linhas_reinf(nfse),
        ('', None),
        ('PENDÊNCIAS PARA RESOLVER ANTES DE IMPORTAR', 'secao'),
        *[(f'   {g}: {c}', None) for g, c in sorted(Counter(p['Gravidade'] for p in pend).items(),
                                                   key=lambda x: ['ALTA', 'VERIFICAR', 'INFO'].index(x[0]))],
        ('   Detalhe na aba "Pendências".', None),
    ]
    _escreve_resumo(ws, linhas)
    ordem = {'ALTA': 0, 'VERIFICAR': 1, 'INFO': 2}
    _aba(wb, 'Pendências', sorted(pend, key=lambda p: ordem.get(p['Gravidade'], 9)), status_col='Gravidade',
         larguras={'O que fazer / por quê': 110, 'Chave': 46})
    _aba(wb, 'Entradas - itens', _so_pre(itens), status_col='Status',
         larguras={'Descrição do item': 42, 'Por quê': 40, 'Alertas': 50, 'Fornecedor': 30,
                   'Histórico do fornecedor no Domínio': 40},
         nota='Um item por linha, com o CFOP, o acumulador e a conta contábil a usar no lançamento.')
    _aba(wb, 'Entradas - notas', _so_pre(notas_t), status_col='Status',
         larguras={'Observações': 70, 'Chave': 46, 'Histórico do fornecedor no Domínio': 40})
    _aba(wb, 'Devoluções e trocas', _so_pre(dev), status_col='Status', larguras={'Observações': 70, 'Chave': 46},
         nota='NF-e de entrada emitidas pela própria empresa (devolução de venda, remessa para troca, consignação).')
    _aba(wb, 'Saídas', _so_pre(sai), status_col='Status', larguras={'Observações': 70, 'Chave': 46})
    _aba(wb, 'Numeração', _so_pre(res['lacunas']), status_col='Gravidade',
         nota='Números da série própria que não têm XML autorizado na pasta do mês.')
    _aba(wb, 'Completude NF-e e CT-e', res['completude'], status_col='Situação',
         nota='Documentos destinados à empresa: lista da SEFAZ x XML. Só aparecem os que faltam em algum lugar.')
    _aba(wb, 'NFS-e tomadas', _so_pre(nfse), status_col='Status',
         nota='Lista nacional de NFS-e recebidas, com o acumulador usado para o prestador nos meses anteriores.')
    _abas_servicos(wb, nfse)
    _aba(wb, 'Regras', _regras(R), larguras={'Motivo': 90, 'Finalidade': 50})
    wb.save(saida)


def gravar_planilha(res, saida):
    if res.get('pre'):
        return gravar_pre(res, saida)
    R = res['regras']
    emp = R.r['empresa']
    comp = res['competencia']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumo'

    itens, notas_t, dev, sai = res['itens'], res['notas_terc'], res['devol'], res['saidas']
    corr = [i for i in itens if i['Corrigido à mão?']]
    igual = lambda i, cf, ac: (i['CFOP sugerido'], str(i['Acum. sugerido'])) == (i[cf], str(i[ac]))
    acerto_corr = sum(igual(i, 'CFOP final (SPED)', 'Acum. final (SPED)') for i in corr)
    acerto_geral = sum(igual(i, 'CFOP final (SPED)', 'Acum. final (SPED)') for i in itens)
    pend_fin = sum(1 for i in itens if i['Finalidade'] == 'PENDENTE')
    divergentes = [i for i in itens if i['Status'] not in ('OK', 'FINALIDADE A DEFINIR')]
    difal = [n for n in notas_t if n['DIFAL estimado'] and n['DIFAL lançado (E113)']]
    difal_ok = sum(1 for n in difal if abs(n['DIFAL estimado'] - n['DIFAL lançado (E113)']) <= Decimal('0.10'))
    sai_ativas = [s for s in sai if not s['Status'].startswith('CANCELADA')]
    dev_ativas = [d for d in dev if not d['Status'].startswith('CANCELADA')]
    tipos = Counter()
    for n in res['notas'].values():
        own = n['emit_cnpj'] == R.cnpj
        tipos['Saídas emitidas pela empresa' if own and n['tpNF'] == '1' else
              'Entradas emitidas pela empresa (devoluções/trocas)' if own else 'Notas de fornecedores'] += 1
    pend = res['pendencias']

    linhas = [
        (f'Conferência fiscal {comp[5:]}/{comp[:4]} — {emp["nome"]}', 'titulo'),
        (f'CNPJ {emp["cnpj"]} · Domínio {emp["codigo_dominio"]} · {emp["regime"]} · {emp["uf"]}', 'sub'),
        ('', None),
        ('O QUE FOI LIDO', 'secao'),
        (f'{len(res["notas"])} NF-e lidas do XML, item por item (sem erro de leitura: {len(res["erros"]) == 0})', None),
        *[(f'   · {k}: {v}', None) for k, v in tipos.most_common()],
        (f'SPED Fiscal: {len(res["sped"]["docs"])} documentos (C100/D100) · relatórios do Domínio: '
         f'{len(res["dom_e"])} entradas e {len(res["dom_s"])} saídas', None),
        ('', None),
        ('ENTRADAS DE FORNECEDORES — identificação de cada item', 'secao'),
        (f'{len(notas_t)} notas · {len(itens)} itens classificados por finalidade (insumo, embalagem, ferramenta, uso e consumo...)', None),
        (f'Sugestão do robô igual ao lançamento final do SPED: {acerto_geral} de {len(itens)} itens', 'destaque'),
        (f'Itens que o escritório corrigiu à mão depois de importar: {len(corr)} — o robô já sugeria a correção certa em {acerto_corr}', 'destaque'),
        (f'Itens sem regra (finalidade a definir com vocês): {pend_fin}', None),
        (f'Itens em que discordo do lançamento final: {len(divergentes)} (ver aba "Entradas - itens", filtro Status)', None),
        (f'DIFAL recalculado (base dupla PR) igual ao lançado no E113 em {difal_ok} de {len(difal)} notas de uso e consumo', None),
        ('', None),
        ('SAÍDAS E DEVOLUÇÕES', 'secao'),
        (f'Saídas: {len(sai_ativas)} notas conferidas · OK {sum(1 for s in sai_ativas if s["Status"] == "OK")} · '
         f'canceladas {len(sai) - len(sai_ativas)} (todas com COD_SIT 02 no SPED)', None),
        (f'Entradas emitidas pela Kopp (devoluções/trocas): {len(dev_ativas)} · lançamento OK {sum(1 for d in dev_ativas if d["Status"] == "OK")}'
         f' · {sum(1 for d in dev_ativas if "tributado) e ICMS zero" in (d.get("Observações") or ""))} saem do ERP com CST tributado e ICMS zero', None),
        ('', None),
        ('PENDÊNCIAS PARA DECIDIR / CORRIGIR', 'secao'),
        *[(f'   {g}: {c}', None) for g, c in sorted(Counter(p['Gravidade'] for p in pend).items(),
                                                   key=lambda x: ['ALTA', 'VERIFICAR', 'INFO'].index(x[0]))],
        ('   Detalhe na aba "Pendências".', None),
        ('', None),
        ('COMO LER AS ABAS', 'secao'),
        ('Entradas - itens: cada item de cada nota de fornecedor. "Sugerido" = robô; "final (SPED)" = como ficou no SPED entregue;', None),
        ('   "na importação" = como o Domínio importou (relatório antes das correções). "Corrigido à mão?" = mudou entre os dois.', None),
        ('Entradas - notas: a mesma análise por nota, com créditos de ICMS/IPI e DIFAL.', None),
        ('Devoluções e trocas / Saídas: XML x Domínio x SPED nota a nota (CFOP, acumulador, base, ICMS, isentas, outras, IPI, cBenef).', None),
        ('Numeração / Completude / NFS-e: notas faltando, inutilizadas, canceladas, CT-e e serviços tomados.', None),
        ('Regras: o que o robô usou para decidir. Cada pendência resolvida vira uma linha nova aqui.', None),
    ]
    _escreve_resumo(ws, linhas)

    ordem = {'ALTA': 0, 'VERIFICAR': 1, 'INFO': 2}
    _aba(wb, 'Pendências', sorted(pend, key=lambda p: ordem.get(p['Gravidade'], 9)), status_col='Gravidade',
         larguras={'O que fazer / por quê': 110, 'Chave': 46})
    _aba(wb, 'Entradas - itens', itens, status_col='Status',
         larguras={'Descrição do item': 42, 'Por quê': 40, 'Alertas': 50, 'Fornecedor': 30},
         nota='Um item por linha. Filtre "Status" diferente de OK para ver só o que diverge do lançamento final.')
    _aba(wb, 'Entradas - notas', notas_t, status_col='Status', larguras={'Observações': 70, 'Chave': 46})
    _aba(wb, 'Devoluções e trocas', dev, status_col='Status', larguras={'Observações': 70, 'Chave': 46},
         nota='NF-e de entrada emitidas pela própria empresa (devolução de venda, remessa para troca, consignação).')
    _aba(wb, 'Saídas', sai, status_col='Status', larguras={'Observações': 70, 'Chave': 46})
    _aba(wb, 'Numeração', res['lacunas'], status_col='Gravidade',
         nota='Números da série própria que não têm XML autorizado na pasta do mês.')
    _aba(wb, 'Completude NF-e e CT-e', res['completude'], status_col='Situação',
         nota='Documentos destinados à empresa: lista da SEFAZ x XML x SPED x Domínio. Só aparecem os que faltam em algum lugar.')
    _aba(wb, 'NFS-e tomadas', res['nfse'], status_col='Status',
         nota='Lista nacional de NFS-e recebidas x lançamentos de serviços tomados no Domínio.')
    _abas_servicos(wb, res['nfse'])
    regras = _regras(R)
    _aba(wb, 'Regras', regras, larguras={'Motivo': 90, 'Finalidade': 50})
    wb.save(saida)
