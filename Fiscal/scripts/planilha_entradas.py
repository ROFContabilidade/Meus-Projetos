"""Planilha de conferência das ENTRADAS do mês, uma linha por documento.
Aba "Todas as notas" (NF-e 55, CT-e 57 e NFS-e juntas) e depois separadas: NF-e dentro do estado,
NF-e fora do estado, CT-e e Serviços tomados. NFS-e com qualquer retenção fica em vermelho.
"""
from collections import Counter, defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = PatternFill('solid', fgColor='1F3864')
CINZA = PatternFill('solid', fgColor='D9E1F2')
VERDE = PatternFill('solid', fgColor='E2EFDA')
AMARELO = PatternFill('solid', fgColor='FFF2CC')
VERMELHO = PatternFill('solid', fgColor='F8CBAD')
FINO = Side(style='thin', color='BFBFBF')
NUM = '#,##0.00'
Z = Decimal('0')

ORDEM_STATUS = {'A LANÇAR': 0, 'CONFERIR': 1}


def _data(v):
    v = str(v or '')
    return f'{v[8:10]}/{v[5:7]}/{v[:4]}' if len(v) >= 10 and v[4] == '-' else v


def _chave_data(v):
    v = _data(v)
    return v[6:10] + v[3:5] + v[:2]


def _cor(status):
    s = str(status or '').upper()
    if s.startswith('A LANÇAR') or s == 'OK':
        return VERDE
    if 'CANCEL' in s or 'NÃO LANÇAR' in s or 'JÁ LANÇADA' in s or 'DOCUMENTO DO FORNECEDOR' in s:
        return CINZA
    if 'CONFERIR' in s or 'PENDENTE' in s or 'DEFINIR' in s:
        return AMARELO
    return VERMELHO


RET_FILL = PatternFill('solid', fgColor='FFC7CE')
RET_FONT = Font(bold=True, color='9C0006')
RETENCOES = ['IRRF retido', 'PIS/COFINS/CSLL retido', 'INSS retido', 'ISS retido']


def _tem_retencao(ln):
    return any(ln.get(c) for c in RETENCOES)


def _aba(wb, titulo, descricao, linhas, colunas, somar, larguras, legenda=None):
    ws = wb.create_sheet(titulo)
    ws.cell(1, 1, titulo).font = Font(bold=True, size=13, color='1F3864')
    ws.cell(2, 1, descricao).font = Font(italic=True, color='595959')
    if legenda:
        ws.cell(3, 1, legenda).font = RET_FONT
    r0 = 4
    for j, c in enumerate(colunas, 1):
        cel = ws.cell(r0, j, c)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = AZUL
        cel.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for i, ln in enumerate(linhas, r0 + 1):
        ret = _tem_retencao(ln)
        for j, c in enumerate(colunas, 1):
            v = ln.get(c)
            if c == 'Emissão':
                v = _data(v)
            if isinstance(v, Decimal):
                v = float(v)
            cel = ws.cell(i, j, v)
            cel.border = Border(bottom=FINO)
            if isinstance(v, float):
                cel.number_format = NUM
            if ret and j > 1:          # nota com retenção: linha inteira em vermelho
                cel.fill = RET_FILL
                if c in RETENCOES or c == 'Retenção':
                    cel.font = RET_FONT
        ws.cell(i, 1).fill = _cor(ln.get('Situação'))
    tot = r0 + len(linhas) + 1
    ws.cell(tot, 1, f'TOTAL ({len(linhas)} documentos)').font = Font(bold=True)
    for c in somar:
        j = colunas.index(c) + 1
        col = get_column_letter(j)
        cel = ws.cell(tot, j, f'=SUBTOTAL(9,{col}{r0 + 1}:{col}{tot - 1})' if linhas else 0)
        cel.font = Font(bold=True)
        cel.number_format = NUM
    for j in range(1, len(colunas) + 1):
        ws.cell(tot, j).fill = CINZA
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(colunas))}{max(r0, tot - 1)}'
    ws.row_dimensions[r0].height = 34
    for j, c in enumerate(colunas, 1):
        ws.column_dimensions[get_column_letter(j)].width = larguras.get(c, max(10, min(len(c) + 2, 16)))
    return ws


def _itens_por_nota(res):
    por = defaultdict(list)
    for it in res['itens']:
        por[(it['Nota'], it['CNPJ'])].append(it)
    return por


def _sem_credito(its, cancel):
    if cancel:
        return ''
    v = sum(((i['ICMS destacado'] or Z) for i in its if i['ICMS destacado'] and not i['Crédito ICMS sugerido']), Z)
    if not v:
        return ''
    return (f'ICMS destacado na nota R$ {v:.2f} sem crédito (uso e consumo/fora da atividade): lançar base e ICMS '
            'zerados').replace('.', ',')


def _resumo_itens(its):
    nomes = [str(i['Descrição do item'])[:40] for i in its]
    txt = '; '.join(nomes[:3])
    return txt + (f' (+{len(nomes) - 3} itens)' if len(nomes) > 3 else '')




def _sit(status):
    s = str(status or '')
    if s in ('OK', 'A LANÇAR'):
        return 'A LANÇAR'
    if s.upper().startswith('CANCELAD'):
        return 'CANCELADA - lançar zerada' if 'ZERAD' in s.upper() or s.upper().startswith('CANCELADA') else s
    if s.startswith(('NÃO LANÇAR', 'JÁ LANÇADA', 'OK (')):
        return s
    return 'CONFERIR'


def _txt_ret(ln):
    nomes = {'IRRF retido': 'IRRF', 'PIS/COFINS/CSLL retido': 'PIS/COFINS/CSLL', 'INSS retido': 'INSS',
             'ISS retido': 'ISS'}
    tem = [nomes[c] for c in RETENCOES if ln.get(c)]
    return 'SIM - ' + ', '.join(tem) if tem else ''


def gerar(res, saida):
    R = res['regras']
    uf = R.uf
    notas = res['notas']
    itens = _itens_por_nota(res)
    docs = []          # todas as entradas, já no formato comum

    # ---------------- NF-e de fornecedores (modelo 55)
    for n in res['notas_terc']:
        x = notas.get(n['Chave'], {})
        its = itens.get((n['Nota'], x.get('emit_cnpj', '')), [])
        cancel = 'CANCELADA' in n['Status']
        fins = Counter(i['Finalidade'] for i in its)
        docs.append({
            'Situação': 'CANCELADA - lançar zerada' if cancel else _sit(n['Status']),
            'Tipo': f'NF-e {"dentro do estado" if n["UF"] == uf else "fora do estado"}', 'Modelo': '55',
            'Emissão': n['Emissão'], 'Número': n['Nota'], 'Série': x.get('serie', ''),
            'Participante': n['Fornecedor'], 'CNPJ': x.get('emit_cnpj', ''), 'UF': n['UF'],
            'CFOP do fornecedor': ', '.join(sorted({i['CFOP forn.'] for i in its})),
            'CFOP': n['CFOP sugerido'], 'Acumulador': n['Acum. sugerido'],
            'Valor contábil': Z if cancel else x.get('t_vNF', n['Valor NF']),
            # base e ICMS só dos itens que dão crédito: uso e consumo (1556/1407) entra zerado
            'Base ICMS': Z if cancel else sum(((i['BC ICMS'] or Z) for i in its if i['Crédito ICMS sugerido']
                                               and i['ICMS destacado']), Z),
            'ICMS': Z if cancel else sum(((i['Crédito ICMS sugerido'] or Z) for i in its if i['ICMS destacado']), Z),
            'IPI': Z if cancel else n['IPI destacado'],
            'Crédito ICMS': n['Crédito ICMS sugerido'], 'Crédito IPI': n['Crédito IPI sugerido'],
            'DIFAL a recolher': n['DIFAL estimado'], 'ICMS antecipado 8%': n.get('ICMS antecipado', Z),
            'Finalidade': ', '.join(k.replace('_', ' ').lower() for k in fins),
            'Conta contábil': n.get('Contas contábeis sugeridas', ''), 'Itens': _resumo_itens(its),
            'Observação': '; '.join(o for o in (n['Observações'], _sem_credito(its, cancel)) if o),
            'Chave': n['Chave'], 'Conferido': ''})
    for p in res['pendencias']:          # documentos do fornecedor que não geram entrada
        if p['Tipo'] == 'Documento do fornecedor':
            x = notas.get(p['Chave'], {})
            docs.append({'Situação': 'NÃO LANÇAR - documento do fornecedor',
                         'Tipo': f'NF-e {"dentro do estado" if x.get("emit_uf") == uf else "fora do estado"}',
                         'Modelo': '55', 'Emissão': x.get('dt_emissao', ''), 'Número': x.get('numero', ''),
                         'Série': x.get('serie', ''), 'Participante': p['Participante'], 'CNPJ': x.get('emit_cnpj', ''),
                         'UF': x.get('emit_uf', ''), 'Valor contábil': Z, 'Observação': p['O que fazer / por quê'],
                         'Chave': p['Chave'], 'Conferido': ''})

    # ---------------- devoluções / retornos emitidos pela própria empresa (modelo 55)
    for d in res['devol']:
        x = notas.get(d['Chave'], {})
        cancel = d['Status'].startswith('CANCELADA')
        u = d.get('UF') or x.get('dest_uf', '')
        docs.append({'Situação': 'CANCELADA - lançar zerada' if cancel else _sit(d['Status']),
                     'Tipo': f'NF-e {"dentro do estado" if u == uf else "fora do estado"}', 'Modelo': '55',
                     'Emissão': d['Data'], 'Número': d['Nota'], 'Série': x.get('serie', ''),
                     'Participante': d['Cliente (remetente)'] + ' (emitida pela empresa)',
                     'CNPJ': x.get('dest_cnpj', ''), 'UF': u, 'CFOP': d['CFOP XML'],
                     'Acumulador': d.get('Acum. esperado', ''),
                     'Valor contábil': Z if cancel else d['Valor NF'], 'Base ICMS': Z if cancel else d.get('BC ICMS XML'),
                     'ICMS': Z if cancel else d.get('ICMS XML'), 'IPI': Z if cancel else d.get('IPI XML'),
                     'Finalidade': 'devolução / retorno', 'Itens': d['Natureza'],
                     'Observação': '; '.join(o for o in (d.get('Observações', ''),
                                                          f'NF de venda: {d["NF de venda referenciada"]}'
                                                          if d.get('NF de venda referenciada') else '') if o),
                     'Chave': d['Chave'], 'Conferido': ''})

    # ---------------- CT-e (modelo 57)
    for c in res.get('cte', []):
        st = c['Status']
        docs.append({'Situação': 'A LANÇAR' if st == 'A LANÇAR' else
                     ('CANCELADO - lançar zerado' if 'ZERADO' in st else st),
                     'Tipo': 'CT-e', 'Modelo': '57', 'Emissão': c['Data'], 'Número': c['CT-e'], 'Série': '',
                     'Participante': c['Transportador'], 'CNPJ': c['CNPJ'], 'UF': c['Trajeto'][:2],
                     'Trajeto': c['Trajeto'], 'Tomador (quem paga)': c['Tomador (quem paga)'],
                     'Remetente': c['Remetente'], 'Destinatário': c['Destinatário'],
                     'CFOP': c['CFOP de entrada sugerido'], 'Acumulador': c['Acum. sugerido'],
                     'Valor contábil': c['Valor'], 'Base ICMS': c['BC ICMS'], 'ICMS': c['ICMS'],
                     'Crédito ICMS': c['Crédito ICMS sugerido'], 'Observação': c['Observações'],
                     'Chave': c['Chave'], 'Conferido': ''})

    # ---------------- NFS-e (serviços tomados)
    for x in res['nfse']:
        ln = {'Situação': _sit(x['Status']), 'Tipo': 'NFS-e (serviço tomado)', 'Modelo': 'NFS-e',
              'Emissão': x.get('Data emissão', ''), 'Competência': x['Competência'], 'Número': x['Número'],
              'Série': '', 'Participante': x['Prestador'], 'CNPJ': x['CNPJ'],
              'UF': '', 'Município': x.get('Município do prestador', x.get('Município ISS', '')),
              'Serviço (LC 116)': x.get('Serviço (LC 116)', ''), 'Natureza do serviço': x.get('Natureza do serviço', ''),
              'Acumulador': x.get('Acum. sugerido', ''), 'Valor contábil': x['Valor'],
              'IRRF retido': x.get('IRRF retido') or None, 'PIS/COFINS/CSLL retido': x.get('CSRF retido') or None,
              'INSS retido': x.get('INSS retido') or None, 'ISS retido': x.get('ISS retido') or None,
              'Simples Nacional': x.get('Simples Nacional', ''),
              'Conta contábil': x.get('Conta contábil sugerida', ''), 'Observação': x.get('Alertas', ''),
              'Conferido': ''}
        ln['Retenção'] = _txt_ret(ln)
        ln['Total retido'] = sum((ln[c] or Z for c in RETENCOES), Z) or None
        docs.append(ln)

    ordem_tipo = {f'NF-e dentro do estado': 0, 'NF-e fora do estado': 1, 'CT-e': 2, 'NFS-e (serviço tomado)': 3}
    ordem = lambda l: (ORDEM_STATUS.get(l['Situação'], 5), _chave_data(l['Emissão']), str(l['Participante']))
    larg = {'Situação': 26, 'Tipo': 20, 'Participante': 34, 'Finalidade': 18, 'Conta contábil': 34, 'Itens': 44,
            'Observação': 50, 'Chave': 46, 'CNPJ': 16, 'Emissão': 11, 'Número': 10, 'Série': 6, 'UF': 5,
            'Modelo': 7, 'Retenção': 26, 'Natureza do serviço': 30, 'Município': 18, 'Simples Nacional': 20,
            'Tomador (quem paga)': 30, 'Remetente': 28, 'Destinatário': 28}
    legenda = ('Linha em VERMELHO = NFS-e com retenção (IRRF, PIS/COFINS/CSLL, INSS ou ISS): gerar guia e informar '
               'na EFD-Reinf / ISS retido.')
    comp = res['competencia']
    emp = R.r['empresa']

    wb = Workbook()
    wb.remove(wb.active)
    # 1) todas as notas juntas
    todas = sorted(docs, key=lambda l: (ordem_tipo.get(l['Tipo'], 9), ORDEM_STATUS.get(l['Situação'], 5),
                                        _chave_data(l['Emissão']), str(l['Participante'])))
    _aba(wb, 'Todas as notas', f'Todas as entradas {comp[5:]}/{comp[:4]} — {emp["nome"]}: NF-e (mod. 55), CT-e (mod. 57) '
         'e NFS-e. Verde = lançar · amarelo = conferir · cinza = cancelada (lançar zerada) / não lançar.',
         todas,
         ['Situação', 'Tipo', 'Modelo', 'Emissão', 'Número', 'Série', 'Participante', 'CNPJ', 'UF', 'CFOP', 'Acumulador',
          'Valor contábil', 'Base ICMS', 'ICMS', 'IPI', 'Crédito ICMS', 'DIFAL a recolher', 'ICMS antecipado 8%',
          'Retenção', 'IRRF retido', 'PIS/COFINS/CSLL retido', 'INSS retido', 'ISS retido', 'Conta contábil',
          'Observação', 'Chave', 'Conferido'],
         ['Valor contábil', 'Base ICMS', 'ICMS', 'IPI', 'Crédito ICMS', 'DIFAL a recolher', 'ICMS antecipado 8%',
          'IRRF retido', 'PIS/COFINS/CSLL retido', 'INSS retido', 'ISS retido'], larg, legenda)

    # 2) separadas
    col_nfe = ['Situação', 'Emissão', 'Número', 'Série', 'Participante', 'CNPJ', 'UF', 'CFOP do fornecedor', 'CFOP',
               'Acumulador', 'Valor contábil', 'Base ICMS', 'ICMS', 'IPI', 'Crédito ICMS', 'Crédito IPI',
               'DIFAL a recolher', 'ICMS antecipado 8%', 'Finalidade', 'Conta contábil', 'Itens', 'Observação',
               'Chave', 'Conferido']
    som_nfe = ['Valor contábil', 'Base ICMS', 'ICMS', 'IPI', 'Crédito ICMS', 'Crédito IPI', 'DIFAL a recolher',
               'ICMS antecipado 8%']
    sel = lambda t: sorted((d for d in docs if d['Tipo'] == t), key=ordem)
    _aba(wb, f'NF-e dentro do estado ({uf})', f'NF-e modelo 55 do {uf} (CFOP 1xxx), inclusive devoluções emitidas pela '
         'empresa. Uso e consumo entra com base e ICMS zerados.', sel('NF-e dentro do estado'), col_nfe, som_nfe, larg)
    _aba(wb, 'NF-e fora do estado', 'NF-e modelo 55 de outros estados (CFOP 2xxx). DIFAL só em uso e consumo; '
         'ICMS antecipado 8% nas notas com crédito a 4%.', sel('NF-e fora do estado'), col_nfe, som_nfe, larg)
    _aba(wb, 'CT-e', 'Fretes (modelo 57). Só lança quando a empresa é a TOMADORA (quem paga o frete).', sel('CT-e'),
         ['Situação', 'Emissão', 'Número', 'Participante', 'CNPJ', 'Trajeto', 'Tomador (quem paga)', 'Remetente',
          'Destinatário', 'CFOP', 'Acumulador', 'Valor contábil', 'Base ICMS', 'ICMS', 'Crédito ICMS', 'Observação',
          'Chave', 'Conferido'],
         ['Valor contábil', 'Base ICMS', 'ICMS', 'Crédito ICMS'], larg)
    sv = sel('NFS-e (serviço tomado)')
    sv.sort(key=lambda l: (not _tem_retencao(l), ORDEM_STATUS.get(l['Situação'], 5), _chave_data(l['Emissão'])))
    _aba(wb, 'Serviços tomados', 'NFS-e recebidas: acumulador pelo item da LC 116. As notas COM RETENÇÃO vêm primeiro.',
         sv,
         ['Situação', 'Emissão', 'Competência', 'Número', 'Participante', 'CNPJ', 'Município', 'Serviço (LC 116)',
          'Natureza do serviço', 'Acumulador', 'Valor contábil', 'Retenção', 'IRRF retido', 'PIS/COFINS/CSLL retido',
          'INSS retido', 'ISS retido', 'Total retido', 'Simples Nacional', 'Conta contábil', 'Observação', 'Conferido'],
         ['Valor contábil', 'IRRF retido', 'PIS/COFINS/CSLL retido', 'INSS retido', 'ISS retido', 'Total retido'],
         larg, legenda)

    # 3) resumo no fim
    ws0 = wb.create_sheet('Resumo')
    ws0['A1'] = f'Entradas {comp[5:]}/{comp[:4]} — {emp["nome"]}'
    ws0['A1'].font = Font(bold=True, size=14, color='1F3864')
    cab = ['Grupo', 'Documentos', 'A lançar', 'Conferir', 'Cancelados / não lançar', 'Valor a lançar',
           'Crédito ICMS', 'Com retenção', 'Total retido']
    for j, c in enumerate(cab, 1):
        cel = ws0.cell(3, j, c)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = AZUL
    grupos = [(t, [d for d in docs if d['Tipo'] == t]) for t in ordem_tipo] + [('TOTAL', docs)]
    for i, (nome, ls) in enumerate(grupos, 4):
        vals = [nome, len(ls), sum(1 for l in ls if l['Situação'] == 'A LANÇAR'),
                sum(1 for l in ls if l['Situação'] == 'CONFERIR'),
                sum(1 for l in ls if l['Situação'] not in ('A LANÇAR', 'CONFERIR')),
                float(sum((Decimal(str(l.get('Valor contábil') or 0)) for l in ls
                           if l['Situação'] in ('A LANÇAR', 'CONFERIR')), Z)),
                float(sum((Decimal(str(l.get('Crédito ICMS') or 0)) for l in ls), Z)),
                sum(1 for l in ls if _tem_retencao(l)),
                float(sum((Decimal(str(l.get(c) or 0)) for l in ls for c in RETENCOES), Z))]
        for j, v in enumerate(vals, 1):
            cel = ws0.cell(i, j, v)
            if isinstance(v, float):
                cel.number_format = NUM
            if nome == 'TOTAL':
                cel.font = Font(bold=True)
                cel.fill = CINZA
    for j, w in enumerate([26, 12, 10, 10, 22, 16, 14, 13, 14], 1):
        ws0.column_dimensions[get_column_letter(j)].width = w
    wb.save(saida)
